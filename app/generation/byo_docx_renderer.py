"""Render a recruiter-supplied DOCX template with Jinja2 placeholders."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docxtpl import DocxTemplate

from app.extraction.candidate_schema import CandidateProfile


def docx_has_placeholders(path: Path) -> bool:
    """Return True if the document contains any {{ or {% Jinja2 markers."""
    doc = Document(path)
    texts: list[str] = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                texts.append(cell.text)
    combined = "".join(texts)
    return "{{" in combined or "{%" in combined


def build_docxtpl_context(profile: CandidateProfile, *, blind: bool) -> dict:
    """Map a CandidateProfile to the flat/nested dict exposed to the template."""

    def s(val: object) -> str:
        return str(val) if val is not None else ""

    ctx: dict = {
        "full_name": s(profile.full_name),
        "email": s(profile.email),
        "phone": s(profile.phone),
        "location": s(profile.location),
        "linkedin": str(profile.linkedin_url) if profile.linkedin_url else "",
        "portfolio": str(profile.portfolio_url) if profile.portfolio_url else "",
        "summary": s(profile.professional_summary),
        "skills": list(profile.skills),
        "languages": [
            {"name": s(lang.name), "proficiency": s(lang.proficiency)}
            for lang in profile.languages
        ],
        "experience": [
            {
                "title": s(job.title),
                "company": s(job.company),
                "location": s(job.location),
                "start_date": s(job.start_date),
                "end_date": s(job.end_date),
                "bullets": list(job.description),
            }
            for job in profile.work_experience
        ],
        "education": [
            {
                "institution": s(edu.institution),
                "degree": s(edu.degree),
                "field_of_study": s(edu.field_of_study),
                "start_date": s(edu.start_date),
                "end_date": s(edu.end_date),
            }
            for edu in profile.education
        ],
        "certifications": [
            {
                "name": s(cert.name),
                "issuer": s(cert.issuer),
                "date": s(cert.date),
            }
            for cert in profile.certifications
        ],
        "salary_expectation": s(profile.salary_expectation),
        "notice_period": s(profile.notice_period),
        "work_authorization": s(profile.work_authorization),
        "interview_availability": s(profile.interview_availability),
    }

    if blind:
        for key in ("full_name", "email", "phone", "linkedin", "portfolio"):
            ctx[key] = ""

    return ctx


def render_byo_docx(
    template_path: Path,
    profile: CandidateProfile,
    output_path: Path,
    *,
    blind: bool,
) -> None:
    """Fill *template_path* with candidate data and write the result to *output_path*."""
    tpl = DocxTemplate(template_path)
    tpl.render(build_docxtpl_context(profile, blind=blind))
    tpl.save(output_path)
