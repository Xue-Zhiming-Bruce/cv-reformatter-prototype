from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml.ns import qn

from app.template_analysis.artifacts import LAYOUT_SPEC_FILENAME
from app.template_analysis.schemas import (
    ColumnStyle,
    DecorationStyle,
    HeaderLayoutStyle,
    LayoutTemplateSpec,
    PageStyle,
    SectionLayoutSpec,
    SpacingStyle,
    TextStyle,
)


class TargetDocxAnalysisError(RuntimeError):
    """Raised when a DOCX target cannot provide a safe renderer style."""


def analyze_target_docx(
    docx_path: str | Path,
    output_dir: str | Path,
) -> LayoutTemplateSpec:
    source_path = Path(docx_path)
    if not source_path.exists() or not source_path.is_file():
        raise FileNotFoundError(f"Target DOCX not found: {source_path}")
    if source_path.suffix.lower() != ".docx":
        raise TargetDocxAnalysisError("Target DOCX analysis requires a .docx file.")

    try:
        document = Document(source_path)
    except Exception as exc:
        raise TargetDocxAnalysisError(
            "The target DOCX is corrupt or cannot be read as a Word document."
        ) from exc

    section = document.sections[0]
    normal = document.styles["Normal"]
    body = _text_style(normal, fallback=TextStyle())
    title = _text_style(
        document.styles["Title"],
        fallback=TextStyle(
            font_family=body.font_family,
            font_size_pt=max(body.font_size_pt * 1.6, body.font_size_pt),
            bold=True,
            color_hex=body.color_hex,
        ),
    )
    heading = _text_style(
        document.styles["Heading 2"],
        fallback=TextStyle(
            font_family=body.font_family,
            font_size_pt=max(body.font_size_pt * 1.15, body.font_size_pt),
            bold=True,
            color_hex=body.color_hex,
        ),
    )

    page_width = _points(section.page_width, 595.28)
    page_height = _points(section.page_height, 841.89)
    paragraph_format = normal.paragraph_format
    heading_format = document.styles["Heading 2"].paragraph_format
    columns = _column_style(section)
    header_border = _header_border(document)
    heading_border, heading_rules_uniform = _uniform_heading_border(document)
    warnings = [
        "DOCX target typography and page geometry are applied through the controlled renderer; "
        "arbitrary body content and unrecognized template placeholders are not copied."
    ]
    if columns.count == 2:
        warnings.append(
            "Two DOCX section columns were detected; content is mapped into controlled main and sidebar regions."
        )
    if heading_border is not None and not heading_rules_uniform:
        warnings.append(
            "Non-uniform Heading 2 borders detected; the global heading_rule "
            "decoration remains disabled."
        )
    rule_border = heading_border if heading_rules_uniform else header_border
    rule_color, rule_width = _border_style(rule_border, heading.color_hex)

    measured_sections = _measured_sections(document, heading, body)
    header = _measured_header(document)
    limited_features = []
    if not measured_sections:
        limited_features.append("limited_capability:docx_section_structure_not_measured")
    if columns.count == 2:
        limited_features.append("limited_capability:docx_column_membership_not_measured")
    spec = LayoutTemplateSpec(
        source_type="docx_analysis",
        template_name="analyzed_docx",
        template_version="docx-measured-1",
        structure_contract=(
            "limited_capability" if limited_features else "measured"
        ),
        page=PageStyle(
            width_pt=page_width,
            height_pt=page_height,
            orientation="portrait" if page_height >= page_width else "landscape",
            margin_top_pt=_points(section.top_margin, 54.0),
            margin_right_pt=_points(section.right_margin, 54.0),
            margin_bottom_pt=_points(section.bottom_margin, 54.0),
            margin_left_pt=_points(section.left_margin, 54.0),
        ),
        columns=columns,
        body=body,
        title=title,
        heading=heading,
        spacing=SpacingStyle(
            line_spacing=_line_spacing(paragraph_format.line_spacing),
            paragraph_after_pt=_points(paragraph_format.space_after, 4.0),
            section_before_pt=_points(heading_format.space_before, 10.0),
            section_after_pt=_points(heading_format.space_after, 4.0),
        ),
        decoration=DecorationStyle(
            primary_color_hex=heading.color_hex,
            accent_color_hex=heading.color_hex,
            rule_color_hex=rule_color,
            rule_width_pt=rule_width,
            header_rule=header_border is not None,
            heading_rule=heading_rules_uniform,
        ),
        header=header,
        sections=measured_sections,
        unsupported_features=limited_features,
        show_document_title=False,
        confidence=0.9,
        warnings=warnings,
    )
    resolved_output_dir = Path(output_dir)
    resolved_output_dir.mkdir(parents=True, exist_ok=True)
    (resolved_output_dir / LAYOUT_SPEC_FILENAME).write_text(
        json.dumps(spec.model_dump(mode="json"), indent=2, ensure_ascii=False),
        encoding="utf-8",
    )
    return spec


def _text_style(style: Any, *, fallback: TextStyle) -> TextStyle:
    font = style.font
    color = fallback.color_hex
    if font.color is not None and font.color.rgb is not None:
        color = f"#{font.color.rgb}"
    return TextStyle(
        font_family=font.name or fallback.font_family,
        font_size_pt=_points(font.size, fallback.font_size_pt),
        bold=fallback.bold if font.bold is None else font.bold,
        italic=fallback.italic if font.italic is None else font.italic,
        color_hex=color,
        character_spacing_pt=_character_spacing(style),
    )


def _character_spacing(style: Any) -> float:
    run_properties = style.element.rPr
    if run_properties is None:
        return 0.0
    spacing = run_properties.find(qn("w:spacing"))
    if spacing is None:
        return 0.0
    try:
        return max(-5.0, min(20.0, float(spacing.get(qn("w:val"), "0")) / 20))
    except ValueError:
        return 0.0


def _measured_sections(
    document: Any,
    heading_style: TextStyle,
    body_style: TextStyle,
) -> list[SectionLayoutSpec]:
    sections: list[SectionLayoutSpec] = []
    entry_title_style = _text_style(
        document.styles["Heading 3"], fallback=body_style
    )
    counts: dict[str, int] = {}
    aliases = {
        "summary": "summary",
        "professional summary": "summary",
        "skills": "skills",
        "languages": "languages",
        "experience": "work_experience",
        "work experience": "work_experience",
        "education": "education",
        "certifications": "certifications",
        "contact": "contact",
        "additional details": "additional_details",
    }
    for paragraph in document.paragraphs:
        if paragraph.style.name != "Heading 2" or not paragraph.text.strip():
            continue
        label = paragraph.text.strip()
        source_key = " ".join(label.casefold().split())
        source = aliases.get(source_key, "additional_details")
        counts[source] = counts.get(source, 0) + 1
        sections.append(
            SectionLayoutSpec(
                section_id=f"docx-{source}-{counts[source]}",
                source=source,
                label=label,
                placement="full_width",
                heading_style=_text_style(paragraph.style, fallback=heading_style),
                entry_title_style=(
                    entry_title_style if source == "work_experience" else None
                ),
                additional_section_heading=(
                    label if source_key not in aliases else None
                ),
            )
        )
    return sections


def _measured_header(document: Any) -> HeaderLayoutStyle:
    fields: list[str] = []
    separator = " "
    for paragraph in document.paragraphs:
        if paragraph.style.name == "Heading 2":
            break
        text = paragraph.text.strip()
        if not text:
            continue
        if "|" in text:
            separator = " | "
        for part in text.split("|"):
            value = part.strip()
            lowered = value.casefold()
            digits = sum(character.isdigit() for character in value)
            field = None
            if "@" in value:
                field = "email"
            elif "linkedin." in lowered:
                field = "linkedin_url"
            elif lowered.startswith(("http://", "https://", "www.")):
                field = "portfolio_url"
            elif digits >= 7:
                field = "phone"
            elif "," in value:
                field = "location"
            if field and field not in fields:
                fields.append(field)
    return HeaderLayoutStyle(
        show_candidate_subheading=True,
        contact_fields=fields,
        contact_separator=separator,
    )


def _column_style(section: Any) -> ColumnStyle:
    columns_nodes = section._sectPr.xpath("./w:cols")
    if not columns_nodes:
        return ColumnStyle()
    count = int(columns_nodes[0].get(qn("w:num"), "1"))
    if count != 2:
        return ColumnStyle()
    gutter_twips = float(columns_nodes[0].get(qn("w:space"), "360"))
    return ColumnStyle(
        count=2,
        gutter_pt=max(gutter_twips / 20, 0),
        left_column_ratio=0.5,
        sidebar_side="left",
    )


def _bottom_border(value: Any) -> Any | None:
    paragraph_properties = (
        value._p.pPr if hasattr(value, "_p") else value.element.pPr
    )
    if paragraph_properties is not None:
        borders = paragraph_properties.find(qn("w:pBdr"))
        if borders is not None and (bottom := borders.find(qn("w:bottom"))) is not None:
            return bottom
    if hasattr(value, "_p"):
        return _bottom_border(value.style)
    base_style = getattr(value, "base_style", None)
    return _bottom_border(base_style) if base_style is not None else None


def _header_border(document: Any) -> Any | None:
    title_border = _bottom_border(document.styles["Title"])
    if title_border is not None:
        return title_border
    for paragraph in document.paragraphs:
        if paragraph.style.name == "Heading 2":
            break
        if (border := _bottom_border(paragraph)) is not None:
            return border
    return None


def _uniform_heading_border(document: Any) -> tuple[Any | None, bool]:
    style_border = _bottom_border(document.styles["Heading 2"])
    if style_border is not None:
        return style_border, True
    paragraphs = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style.name == "Heading 2"
    ]
    if not paragraphs:
        return None, False
    borders = [_bottom_border(paragraph) for paragraph in paragraphs]
    present = [border for border in borders if border is not None]
    return (present[0] if present else None), bool(present) and len(present) == len(borders)


def _border_style(border: Any | None, fallback_color: str) -> tuple[str, float]:
    if border is None:
        return fallback_color, 0.75
    color = border.get(qn("w:color"), fallback_color.removeprefix("#"))
    if color.lower() == "auto" or len(color) != 6:
        color = fallback_color.removeprefix("#")
    try:
        width = float(border.get(qn("w:sz"), "6")) / 8
    except ValueError:
        width = 0.75
    return f"#{color.upper()}", max(0.25, min(6.0, width))


def _line_spacing(value: Any) -> float:
    if value is None:
        return 1.1
    if isinstance(value, (int, float)):
        return max(0.8, min(2.0, float(value)))
    return 1.1


def _points(value: Any, default: float) -> float:
    if value is None:
        return default
    points = getattr(value, "pt", None)
    return round(float(points if points is not None else value), 3)
