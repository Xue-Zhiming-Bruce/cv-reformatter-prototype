"""Native OOXML DOCX generation for the compiled C2 render plan.

Phase-1 pure-refactor split from ``c2_docx_renderer``: the authored-DOCX
constants, written-font resolution, fit models, document construction
(``build_document``), OOXML inspection, expected-structure derivation, and
content accounting/reading-order gates live here. ``c2_docx_renderer``
re-exports every name, so all existing imports keep working unchanged.
"""

# Phase-1 split (pure refactor, commit "c2: split C2 modules by responsibility
# without behavior change"): this code was MOVED verbatim from the original
# module named in the function references; all public entry points and import
# paths are preserved by re-exports in the original modules. No output,
# criterion, or data meaning was changed.

from __future__ import annotations

import re
import zipfile
from functools import lru_cache
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph
from pydantic import Field

from tests.experiments.a_pipeline import _median
from tests.experiments.c2_renderer import _norm
from tests.experiments.c2_state import (  # phase-1 split
    C2LayoutState,
    StateModel,
    StyleToken,
)





# Word paragraph borders carry sizes in eighths of a point and spacing in
# points (clamped by Word to 0..31); character spacing uses twentieths.
BORDER_SIZE_QUANTUM = 8.0
SPACING_QUANTUM = 20.0

# Documented renderer rules (the smallest rules that make built-in Word
# style defaults unable to leak uncontrolled visual behavior):
# - unmeasured run color renders black (Word's built-in Heading-1 blue is
#   never inherited);
# - every rendered paragraph carries explicit spacing/line-spacing/indent
#   control (built-in Normal/List/Heading spacing is overridden);
# - entry rows with metadata render as borderless two-column tables whose
#   left column takes 70% of the measured entry text width (the evidence
#   measures the entry-column x0 and the right edge only; metadata is
#   right-aligned inside its column so the rendered topology matches).
UNMEASURED_COLOR_HEX = "000000"
ENTRY_LEFT_COLUMN_FRACTION = 0.70

# ---------------------------------------------------------------------------
# Rendered-geometry measurement and bounded fitting (owner work order)
# ---------------------------------------------------------------------------
# Comparison unit: points. Word units are converted only at the compiler
# boundary (1 pt = 20 twips, 1 pt = 12,700 EMU, half-point font sizes,
# eighth-point border widths). Tolerances documented BEFORE evaluating
# results — the smallest defensible values (mirroring the C2-0b Chrome
# vector-quantization tolerance where the same quantization applies):
TOLERANCE_PT = {
    "page_geometry": 0.5,
    "rule_x_extent": 1.0,
    "rule_stroke": 0.5,
    "local_position": 1.5,
    "local_gap": 1.5,
    "font_size": 0.5,
    "column_right_edge": 1.5,
}
MAX_FITTING_ITERATIONS = 3
# A trailing page is classified sparse when its content extends less than this
# fraction of the writable height below the top margin (documented before
# evaluation; the frozen C1 D->E second page measures ~0.44, so appended
# overflow sections never trip it).
SPARSE_TRAILING_PAGE_FRACTION = 0.30


# ---------------------------------------------------------------------------
# Source presentation markers vs candidate content (work order Part 2.3)
# ---------------------------------------------------------------------------

# CONFIRMED presentation markers: a leading bullet glyph or dash followed by
# whitespace is visually redundant with a native Word list bullet and is
# converted into it (the substantive text is never rewritten). Arrows and
# every other leading glyph are NOT confirmed markers — the state carries no
# evidence they are presentation — so they are always preserved as content.
BULLET_GLYPH_MARKERS = ("•", "◦", "·", "▪")
DASH_MARKERS = ("-", "–", "—")


# ---------------------------------------------------------------------------
# Truthful typography: requested vs written vs rendered fonts (work order Part 3)
# ---------------------------------------------------------------------------

_FONT_SEARCH_DIRS = (
    Path("/System/Library/Fonts"),
    Path("/System/Library/Fonts/Supplemental"),
    Path("/Library/Fonts"),
    Path.home() / "Library/Fonts",
    Path("/usr/share/fonts"),
    Path("/usr/local/share/fonts"),
)
# Documented family alias tried before the portable fallback ("Charter BT"
# measures the installed "Charter" family).
FAMILY_ALIASES = {"charter bt": "charter"}


@lru_cache(maxsize=1)
def installed_font_families() -> frozenset[str]:
    """Lowercase stems of locally installed font files (no network access)."""
    families: set[str] = set()
    for directory in _FONT_SEARCH_DIRS:
        if not directory.is_dir():
            continue
        for path in directory.rglob("*"):
            if path.suffix.casefold() in {".ttf", ".otf", ".ttc"}:
                families.add(path.stem.casefold())
    return frozenset(families)


def _font_is_installed(family: str) -> bool:
    key = family.casefold().replace(" ", "")
    installed = installed_font_families()
    if key in installed:
        return True
    alias = FAMILY_ALIASES.get(key)
    if alias and alias in installed:
        return True
    first = key.split("_")[0] if "_" in key else None
    return bool(first) and first in installed


def resolve_written_fonts(state: C2LayoutState) -> dict[str, dict[str, Any]]:
    """The documented portable font policy, per measured style token.

    Inspects installed fonts first (no network). An installed measured family
    is written verbatim; an unavailable one is written as the portable
    sans-serif substitute ("Arial" — available on the pinned owner-review
    renderer and metric-compatible with Liberation Sans elsewhere) and
    classified ``substituted``. The requested measured family is always
    recorded separately, so substituted typography can never pass as exact.
    """
    resolved: dict[str, dict[str, Any]] = {}
    for token in state.styles:
        requested = token.font_family
        substituted = not _font_is_installed(requested)
        resolved[token.style_id] = {
            "requested": requested,
            "written": "Arial" if substituted else requested,
            "substituted": substituted,
            "detail": (
                f"requested {requested!r} is not installed in the pinned preview "
                "environment; documented portable sans fallback 'Arial' written"
                if substituted
                else f"{requested!r} installed"
            ),
        }
    return resolved


def _half_point_round(size_pt: float) -> float:
    """OOXML stores font sizes in half-points; round to the NEAREST half
    point at the compiler boundary (Word quantization, documented)."""
    return round(float(size_pt) * 2.0) / 2.0


def normalize_font_family(fontname: str | None) -> str | None:
    """Rendered PDF font name -> comparable family name.

    Strips the embedded-subset prefix (``CAAAAA+``), the style suffix
    (``-Bold``), and PostScript-only decorations (``MT``), lowercased and
    space-free, so a rendered name can be compared with the written family
    deterministically."""
    if not fontname:
        return None
    name = fontname.split("+")[-1]
    base = name.split("-")[0]
    folded = base.casefold().replace(" ", "")
    for suffix in ("mt", "ps"):
        if folded.endswith(suffix):
            folded = folded[: -len(suffix)]
    return folded


def strip_presentation_marker(text: str, native_bullet: bool) -> str:
    """Return the rendered text for a line converted to a native Word bullet.

    Only a confirmed leading presentation marker (bullet glyph or dash plus
    whitespace) is removed, and ONLY when the line actually becomes a native
    Word bullet. Zero-bullet designs and plain paragraphs keep every glyph
    verbatim; arrows are never stripped.
    """
    if not native_bullet:
        return text
    for marker in (*BULLET_GLYPH_MARKERS, *DASH_MARKERS):
        if text.startswith(marker) and (
            len(text) == len(marker) or text[len(marker)].isspace()
        ):
            return text[len(marker):].lstrip()
    return text


# ---------------------------------------------------------------------------
# Bounded deterministic fitting (work order Part 5)
# ---------------------------------------------------------------------------


class BulletTiers(StateModel):
    """Target-measured absolute bullet anchors (page x, points) supplied by the
    fitter when the state declares no bullet tiers."""

    marker_x0_pt: float
    text_x0_pt: float


class SectionFit(StateModel):
    """Per-section fitting corrections (documented translation rules only).

    Section rhythm is a per-section declared measurement, so every spacing/
    indent/rule control is section-scoped: a failing gap moves ONLY its own
    section's control (a shared control would triple-apply each section's
    delta and oscillate)."""

    node_id: str
    entry_right_edge_pt: float = 0.0  # right-column width correction
    entry_table_indent_pt: float = 0.0  # table indent correction
    # Entry child detail/bullet lines (verbatim textline leaves): fitter
    # correction relative to the declared entry-column alignment (C2-0cR).
    entry_child_text_indent_pt: float = 0.0
    item_left_indent_pt: float | None = None  # absolute item indent (items sections)
    bullet_marker_correction_pt: float = 0.0
    bullet_text_correction_pt: float = 0.0
    inter_entry_pt: float = 0.0
    heading_space_before_pt: float = 0.0
    heading_space_after_pt: float = 0.0
    heading_border_space_pt: float = 0.0
    rule_left_indent_pt: float = 0.0
    rule_right_indent_pt: float = 0.0
    bullet_tiers: BulletTiers | None = None


class FitAdjustments(StateModel):
    """Bounded deterministic DOCX fitting corrections.

    Every value comes from a measured target/state value minus the measured
    rendered delta, applied through a documented renderer-neutral/DOCX
    translation rule; zero corrections leave the declared compiler behavior
    untouched. The fitter never invents controls outside this model.
    """

    sections: dict[str, SectionFit] = Field(default_factory=dict)

    def section(self, node_id: str) -> SectionFit:
        return self.sections.setdefault(node_id, SectionFit(node_id=node_id))


# ---------------------------------------------------------------------------
# Compatibility contract models (deterministic, evidence-backed)
# ---------------------------------------------------------------------------


class ExactClaim(StateModel):
    """An ``exact`` compatibility claim, output-verified against the written
    OOXML. A claim without recorded inspection evidence fails the report."""

    claim: str
    verified: bool
    evidence: str


class AdjustedFeature(StateModel):
    """A content-preserving, approved visual degradation (proposal §16.3).

    ``adjusted`` NEVER drops content: the feature renders with a visible,
    non-blocking warning recorded here.
    """

    feature: str
    detail: str
    content_preserved: bool = True
    evidence: list[str] = Field(default_factory=list)


class UnsupportedFeature(StateModel):
    """A feature the DOCX spike cannot honestly reproduce.

    Every entry requires explicit owner confirmation before the output could
    ever be treated as acceptable; the canonical run reports it fail-closed.
    """

    feature: str
    detail: str
    evidence: list[str] = Field(default_factory=list)
    content_at_risk: str


class PaginationCompatibility(StateModel):
    """Explicit pagination evidence — a page-count change can never disappear
    from the compatibility report."""

    target_page_count: int
    frozen_c1_page_count: int | None
    docx_preview_page_count: int | None
    classification: str  # exact | adjusted | unsupported
    detail: str = ""


class ConversionCompatibilityReport(StateModel):
    """Deterministic cross-format compatibility contract (proposal §16.3)."""

    schema_version: str = "c2-conversion-compatibility/1"
    source_format: str = "layout-state/1 (C2LayoutState JSON) + candidate render context"
    output_format: str = "docx (WordprocessingML OOXML)"
    exact: list[ExactClaim] = Field(default_factory=list)
    adjusted: list[AdjustedFeature] = Field(default_factory=list)
    unsupported: list[UnsupportedFeature] = Field(default_factory=list)
    pagination: PaginationCompatibility | None = None
    typography: dict[str, Any] | None = None
    content_loss_risk: bool = False
    fallback_applied: list[str] = Field(default_factory=list)
    owner_confirmation_required: bool = False


class DocumentReviewResult(StateModel):
    """C2-0eB: document-level POST-RENDER review result (additive, separate
    from any pre-render adaptation action). The sparse-page density is a
    measured property of the RENDERED document; this result classifies it
    for owner review. It never mutates the render plan, never triggers
    re-rendering, and never converts an output into a passing one-shot
    result."""

    schema_version: str = "c2-docx-review-result/1"
    status: str  # ready | review_required | unsupported
    reason_code: str  # no_sparse_trailing_page | sparse_trailing_page | trailing_page_density_unmeasurable
    page_count: int | None = None
    trailing_page_density: float | None = None
    density_threshold: float | None = None
    evidence_ref: str = "docx_rendered_geometry.json#sparse_trailing_page"


# ---------------------------------------------------------------------------
# Low-level OOXML helpers (python-docx has no typed API for these)
# ---------------------------------------------------------------------------


def _apply_token(run: Any, token: StyleToken | None, written_font: str | None = None) -> None:
    """Apply one measured style token to a run as direct formatting.

    Documented rules: an unmeasured color renders black (a built-in Word
    style color such as Heading-1 blue is never inherited), an unmeasured
    tier (None token) leaves the run at the paragraph's own controlled
    defaults, and the run carries the WRITTEN font family (the documented
    portable fallback when the measured family is not installed) — never a
    bare family name the pinned renderer would substitute unpredictably."""
    if token is None:
        return
    if written_font:
        run.font.name = written_font
    run.font.size = Pt(_half_point_round(token.font_size_pt))
    if token.bold:
        run.font.bold = True
    if token.italic:
        run.font.italic = True
    run.font.color.rgb = RGBColor.from_string(
        (token.color_hex or UNMEASURED_COLOR_HEX).lstrip("#").upper()
    )
    if token.character_spacing_pt:
        # Native Word character spacing in twentieths of a point.
        r_pr = run._element.get_or_add_rPr()  # noqa: SLF001
        spacing = r_pr.find(qn("w:spacing"))
        if spacing is None:
            spacing = OxmlElement("w:spacing")
            r_pr.append(spacing)
        spacing.set(qn("w:val"), str(int(round(float(token.character_spacing_pt) * SPACING_QUANTUM))))


def _paragraph_border(paragraph: Any, edge: str, stroke_pt: float, color_hex: str, space_pt: float | None) -> None:
    """Native Word paragraph border (an editable paragraph property, never a
    floating shape): the honest DOCX equivalent of a measured section rule."""
    p_pr = paragraph._p.get_or_add_pPr()  # noqa: SLF001
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    element = OxmlElement(f"w:{edge}")
    element.set(qn("w:val"), "single")
    element.set(qn("w:sz"), str(max(1, int(round(float(stroke_pt) * BORDER_SIZE_QUANTUM)))))
    element.set(qn("w:space"), str(int(round(float(space_pt or 0.0)))))
    element.set(qn("w:color"), color_hex.lstrip("#").upper())
    p_bdr.append(element)


def _control_paragraph(
    paragraph: Any,
    token: StyleToken | None,
    *,
    space_before_pt: float | None = None,
    space_after_pt: float | None = None,
    alignment: str | None = None,
    keep_with_next: bool | None = None,
    left_indent_pt: float | None = None,
    right_indent_pt: float | None = None,
    first_line_indent_pt: float | None = None,
) -> None:
    """Explicitly control every renderer-relevant paragraph property.

    Built-in Word style defaults (Heading 1 / Normal / List Bullet spacing
    and indents) must never leak: spacing defaults to 0, line spacing comes
    from the measured token, and widow/orphan control is on."""
    paragraph_format = paragraph.paragraph_format
    # Word's spacing measures are unsigned twips at this boundary; the fitter
    # clamps at zero (a measured gap cannot go below 0pt).
    paragraph_format.space_before = Pt(round(max(float(space_before_pt or 0.0), 0.0), 3))
    paragraph_format.space_after = Pt(round(max(float(space_after_pt or 0.0), 0.0), 3))
    if token is not None and token.line_height_pt:
        # Measured line height as an EXACT line height (Word lineRule exact;
        # the renderer honors the measured points directly instead of
        # rescaling a multiple by its own substituted font metrics).
        paragraph_format.line_spacing = Pt(round(float(token.line_height_pt), 3))
    else:
        paragraph_format.line_spacing = 1.0
    paragraph_format.widow_control = True
    if alignment == "right":
        paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif alignment == "center":
        paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if keep_with_next is not None:
        paragraph_format.keep_with_next = keep_with_next
    paragraph_format.left_indent = Pt(round(float(left_indent_pt or 0.0), 3))
    paragraph_format.right_indent = Pt(round(float(right_indent_pt or 0.0), 3))
    if first_line_indent_pt is not None:
        paragraph_format.first_line_indent = Pt(round(float(first_line_indent_pt), 3))


def _no_table_borders(table: Any) -> None:
    """Borders genuinely absent: explicit ``none`` tblBorders (never inherited
    from a Word table style)."""
    tbl_pr = table._tbl.tblPr  # noqa: SLF001
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "none")
        element.set(qn("w:sz"), "0")
        element.set(qn("w:space"), "0")
        borders.append(element)
    tbl_pr.append(borders)


def _fixed_table_layout(table: Any, column_widths_pt: list[float], indent_pt: float) -> None:
    """Fixed layout with deliberate measured column widths and zero cell
    margins, so the rendered columns match the measured geometry.

    The tblGrid carries the explicit widths: with fixed layout the pinned
    renderer lays the table out from the grid (a bare tcW-only table is
    expanded to the writable page width, which measured wrong in the first
    corrective preview)."""
    tbl = table._tbl  # noqa: SLF001
    tbl_pr = tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    if abs(indent_pt) >= 0.01:
        indent = OxmlElement("w:tblInd")
        indent.set(qn("w:w"), str(int(round(indent_pt * 20))))
        indent.set(qn("w:type"), "dxa")
        tbl_pr.append(indent)
    margins = OxmlElement("w:tblCellMar")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:w"), "0")
        element.set(qn("w:type"), "dxa")
        margins.append(element)
    tbl_pr.append(margins)
    width = OxmlElement("w:tblW")
    width.set(qn("w:w"), str(int(round(sum(column_widths_pt) * 20))))
    width.set(qn("w:type"), "dxa")
    tbl_pr.append(width)
    old_grid = tbl.find(qn("w:tblGrid"))
    if old_grid is not None:
        tbl.remove(old_grid)
    grid = OxmlElement("w:tblGrid")
    for column_width in column_widths_pt:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(int(round(column_width * 20))))
        grid.append(column)
    tbl_pr.addnext(grid)
    for index, width in enumerate(column_widths_pt):
        for cell in table.columns[index].cells:
            cell.width = Pt(round(float(width), 3))


def _rows_cannot_split(table: Any) -> None:
    """Keep each entry row together across pages (keep-together)."""
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()  # noqa: SLF001
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))


def _style_of(state: C2LayoutState, style_id: str | None) -> StyleToken | None:
    if style_id is None:
        return None
    return next((style for style in state.styles if style.style_id == style_id), None)


def _entry_columns_of(state: C2LayoutState) -> dict[str, Any]:
    """Measured entry-row columns per section node id (left x0 + right x1)."""
    return {
        node.parent_id: node
        for node in state.nodes
        if node.kind == "entry_row" and node.parent_id
    }


def _header_row_text(fields: list[Any], separator: str | None) -> str:
    """One header row's paragraph text (same presentation contract as the
    C2-0b HTML renderer: spaces around the measured separator)."""
    joiner = f" {separator} " if separator else "   "
    return joiner.join(field.text for field in fields)


# ---------------------------------------------------------------------------
# DOCX renderer: plan -> native OOXML (deterministic; state stays authoritative)
# ---------------------------------------------------------------------------


def _write_text_paragraph(
    document: Any,
    text: str,
    token: StyleToken | None,
    *,
    written_font: str | None = None,
    alignment: str | None = None,
    space_before_pt: float | None = None,
    space_after_pt: float | None = None,
    keep_with_next: bool | None = None,
    left_indent_pt: float | None = None,
    first_line_indent_pt: float | None = None,
    style: str | None = None,
) -> Any:
    """One fully controlled paragraph in the document body."""
    paragraph = document.add_paragraph(style=style)
    run = paragraph.add_run(text)
    if token is not None:
        _apply_token(run, token, written_font)
    _control_paragraph(
        paragraph,
        token,
        space_before_pt=space_before_pt,
        space_after_pt=space_after_pt,
        alignment=alignment,
        keep_with_next=keep_with_next,
        left_indent_pt=left_indent_pt,
    )
    if first_line_indent_pt is not None:
        paragraph.paragraph_format.first_line_indent = Pt(round(float(first_line_indent_pt), 3))
    return paragraph


def _write_heading(
    document: Any,
    state: C2LayoutState,
    label: str,
    label_case: str | None,
    style_token: StyleToken | None,
    rule: dict[str, Any] | None,
    heading_gaps: tuple[float | None, float | None],
    written_font: str | None = None,
    adjustments: FitAdjustments | None = None,
    section_plan: Any = None,
) -> Any:
    """Native Word heading paragraph (outline-level editable) with every
    visual property explicitly controlled from the measured state; the
    measured rule renders as a paragraph border with the measured x-extent
    consumed as indents."""
    del label_case  # the state label already carries the measured casing
    adjustments = adjustments or FitAdjustments()
    fit = adjustments.section(section_plan.node_id if section_plan is not None else "")
    paragraph = document.add_paragraph(style="Heading 1")
    run = paragraph.add_run(label)
    if style_token is not None:
        _apply_token(run, style_token, written_font)
    gap_above, gap_below = heading_gaps
    left_indent = right_indent = 0.0
    border_space = 0.0
    if rule is not None:
        page = state.page
        if rule.get("x0_pt") is not None:
            left_indent = round(float(rule["x0_pt"]) - float(page.margin_left_pt), 3)
        if rule.get("x1_pt") is not None:
            right_indent = round(
                float(page.width_pt) - float(page.margin_right_pt) - float(rule["x1_pt"]), 3
            )
        if rule.get("placement") == "below_heading":
            # gap_above = heading text -> rule; gap_below = rule -> content.
            border_space = float(rule.get("gap_above_pt") or 0.0)
            space_after = float(rule.get("gap_below_pt") or gap_below or 0.0)
        else:
            # gap_above = content above -> rule; gap_below = rule -> heading.
            border_space = float(rule.get("gap_below_pt") or 0.0)
            space_after = float(gap_below or 0.0)
        border_space = max(0.0, border_space + fit.heading_border_space_pt)
        if rule.get("placement") == "below_heading":
            _paragraph_border(
                paragraph, "bottom", float(rule["stroke_pt"]), rule["color_hex"], border_space
            )
        else:
            _paragraph_border(
                paragraph, "top", float(rule["stroke_pt"]), rule["color_hex"], border_space
            )
        space_after += fit.heading_space_after_pt
    else:
        space_after = float(gap_below or 0.0) + fit.heading_space_after_pt
    left_indent = round(left_indent + fit.rule_left_indent_pt, 3)
    right_indent = round(right_indent + fit.rule_right_indent_pt, 3)
    # Headings never orphan: keep with the first content line. Explicit
    # spacing (0 when unmeasured) blocks the built-in Heading 1 defaults.
    _control_paragraph(
        paragraph,
        style_token,
        space_before_pt=(gap_above or 0.0) + fit.heading_space_before_pt,
        space_after_pt=space_after,
        alignment="left",
        keep_with_next=True,
        left_indent_pt=left_indent,
        right_indent_pt=right_indent,
    )
    return paragraph


def _write_entry_table(
    document: Any,
    state: C2LayoutState,
    section_plan: Any,
    entry: Any,
    title_token: StyleToken | None,
    detail_token: StyleToken | None,
    meta_token: StyleToken | None,
    written_fonts: dict[str, str] | None = None,
    space_before_pt: float | None = None,
    adjustments: FitAdjustments | None = None,
) -> list[Any]:
    """One entry's left/right topology as a borderless fixed-layout two-column
    Word table: left column = company/role/degree/detail lines, right column =
    location/dates metadata (right-aligned), preserving the measured target
    topology with fully editable native content. The measured inter-entry
    rhythm renders as space AFTER the entry's last paragraph (returned so the
    caller can set it) — a per-cell space-before would push only the left
    column down and break the row's internal baseline alignment."""
    adjustments = adjustments or FitAdjustments()
    written_fonts = written_fonts or {}
    entry_node = _entry_columns_of(state).get(section_plan.node_id)
    page = state.page
    base_x0 = section_plan.base_x0_pt if section_plan.base_x0_pt is not None else float(page.margin_left_pt)
    right_x1 = None
    if entry_node is not None and entry_node.columns:
        right_x1 = next(
            (column.x1_pt for column in entry_node.columns if column.x1_pt is not None), None
        )
    text_width = (
        float(right_x1) - float(base_x0)
        if right_x1 is not None
        else float(page.width_pt) - float(page.margin_left_pt) - float(page.margin_right_pt)
    )
    left_width = round(text_width * ENTRY_LEFT_COLUMN_FRACTION, 3)
    fit = adjustments.section(section_plan.node_id)
    right_width = max(12.0, text_width - left_width + fit.entry_right_edge_pt)
    column_widths = [left_width, round(right_width, 3)]
    table = document.add_table(rows=1, cols=2)
    _no_table_borders(table)
    _fixed_table_layout(
        table, column_widths,
        indent_pt=round(float(base_x0) - float(page.margin_left_pt) + fit.entry_table_indent_pt, 3),
    )
    _rows_cannot_split(table)
    left_cell, right_cell = table.rows[0].cells
    paragraphs: list[Any] = []
    # Empty cell paragraphs (python-docx seeds one per cell) are reused so no
    # stray empty paragraph enters the reading order.
    first_left = True
    for line_index, line in enumerate(entry.title_lines):
        paragraph = left_cell.paragraphs[0] if first_left else left_cell.add_paragraph()
        first_left = False
        run = paragraph.add_run(line.text)
        token = title_token if line_index == 0 else detail_token
        _apply_token(run, token, written_fonts.get(token.style_id) if token else None)
        _control_paragraph(
            paragraph,
            token,
            space_before_pt=space_before_pt if line_index == 0 else None,
            keep_with_next=bool(entry.bullet_items or entry.text_lines) or None,
        )
        paragraphs.append(paragraph)
    if first_left:  # no title lines: drop the seeded empty paragraph text
        left_cell.paragraphs[0].paragraph_format.space_after = Pt(0)
    first_right = True
    for line in entry.meta_lines:
        paragraph = right_cell.paragraphs[0] if first_right else right_cell.add_paragraph()
        first_right = False
        run = paragraph.add_run(line.text)
        _apply_token(run, meta_token, written_fonts.get(meta_token.style_id) if meta_token else None)
        _control_paragraph(paragraph, meta_token, alignment="right")
        paragraphs.append(paragraph)
    if first_right:
        right_cell.paragraphs[0].paragraph_format.space_after = Pt(0)
    return paragraphs


def category_grid_table_geometry(
    state: C2LayoutState, section_plan: Any
) -> dict[str, Any] | None:
    """The expected authored table geometry of the section's measured
    category grid (C2-0cS): contiguous cell boundaries from the measured
    anchors — [previous split (or writable left edge), value_x0] per column
    plus the writable right edge. Shared by the DOCX emitter (authoring) and
    the compatibility claim (output verification)."""
    grid = next(
        (
            node.category_grid for node in state.nodes
            if node.node_id == section_plan.node_id and node.category_grid
        ),
        None,
    )
    if grid is None:
        return None
    page = state.page
    right_edge = round(float(page.width_pt) - float(page.margin_right_pt), 3)
    splits = [*grid.column_splits_x_pt, right_edge]
    boundaries: list[float] = []
    for column_index, column in enumerate(grid.columns):
        boundaries.append(splits[column_index - 1] if column_index else float(page.margin_left_pt))
        boundaries.append(column.value_x0_pt)
    boundaries.append(right_edge)
    return {
        "columns": 2 * len(grid.columns),
        "column_widths_pt": [
            round(boundaries[index + 1] - boundaries[index], 3)
            for index in range(len(boundaries) - 1)
        ],
        "rows": max(len({cell.row_index for cell in section_plan.category_grid_cells}), 1),
        "label_value_gaps_pt": [column.label_value_gap_pt for column in grid.columns],
    }


def _write_category_grid(
    document: Any,
    state: C2LayoutState,
    section_plan: Any,
    content_token: StyleToken | None,
    content_written: str | None,
    written_fonts: dict[str, str] | None = None,
    adjustments: FitAdjustments | None = None,
) -> list[Any]:
    """The measured category grid (C2-0cS) as a borderless fixed-layout Word
    table: one row per candidate grid row; per column a right-aligned label
    sub-cell and a left-aligned value sub-cell. The column boundaries are
    the MEASURED anchors (label right edge, value x0, documented column
    split, writable right edge); the label sub-cell's right paragraph indent
    is the measured label→value gap, so the right-aligned label ends exactly
    at the measured label edge. Fully editable native table content — the
    fragments concatenate to each leaf's verbatim text (accounting)."""
    adjustments = adjustments or FitAdjustments()
    written_fonts = written_fonts or {}
    geometry = category_grid_table_geometry(state, section_plan)
    grid = next(
        (
            node.category_grid for node in state.nodes
            if node.node_id == section_plan.node_id and node.category_grid
        ),
        None,
    )
    if grid is None or geometry is None:
        return []
    label_token = _style_of(state, "style.body")
    column_widths = geometry["column_widths_pt"]
    table = document.add_table(rows=geometry["rows"], cols=geometry["columns"])
    _no_table_borders(table)
    _fixed_table_layout(table, column_widths, indent_pt=0.0)
    _rows_cannot_split(table)
    paragraphs: list[Any] = []
    by_row: dict[int, list[Any]] = {}
    for cell in section_plan.category_grid_cells:
        by_row.setdefault(cell.row_index, []).append(cell)
    controlled_cells: set[tuple[int, int]] = set()
    for row_index in sorted(by_row):
        row = table.rows[row_index]
        for cell in sorted(by_row[row_index], key=lambda c: c.column_index):
            column = grid.columns[cell.column_index]
            label_cell = row.cells[2 * cell.column_index]
            value_cell = row.cells[2 * cell.column_index + 1]
            controlled_cells.update(
                {
                    (row_index, 2 * cell.column_index),
                    (row_index, 2 * cell.column_index + 1),
                }
            )
            label_paragraph = label_cell.paragraphs[0]
            if cell.label_text:
                label_run = label_paragraph.add_run(cell.label_text)
                _apply_token(
                    label_run, label_token,
                    written_fonts.get(cell.label_style_id) if cell.label_style_id else None,
                )
            _control_paragraph(
                label_paragraph, content_token,
                alignment="right",
                right_indent_pt=round(float(column.label_value_gap_pt), 3),
            )
            paragraphs.append(label_paragraph)
            value_paragraph = value_cell.paragraphs[0]
            if cell.value_text:
                value_run = value_paragraph.add_run(cell.value_text)
                _apply_token(value_run, content_token, content_written)
            _control_paragraph(value_paragraph, content_token)
            paragraphs.append(value_paragraph)
    # C2-0eB truthfulness fix: sub-cell positions NOT owned by a bound grid
    # cell (fewer candidate rows than the measured grid, e.g. a 2-row
    # candidate grid in a 3-row measured topology, or a missing label/value
    # split) would otherwise carry python-docx's uncontrolled default
    # paragraphs, breaking the output-verified "every paragraph explicitly
    # controlled" claim. They are controlled empty paragraphs — never target
    # text, never invented content.
    for row_index in range(len(table.rows)):
        for position in range(len(table.rows[row_index].cells)):
            if (row_index, position) in controlled_cells:
                continue
            _control_paragraph(table.rows[row_index].cells[position].paragraphs[0], content_token)
    return paragraphs


def build_document(
    state: C2LayoutState, plan: Any, adjustments: FitAdjustments | None = None
) -> Document:
    """Compile the render plan into a native DOCX. Deterministic; the state is
    read-only input; no HTML route; fail closed on a failed plan. Optional
    bounded fitting corrections (see ``FitAdjustments``) shift declared
    spacing/indent/width values onto the target-measured geometry."""
    if plan.failures:
        raise RuntimeError(f"refusing to render a failed plan: {plan.failures[:3]}")
    adjustments = adjustments or FitAdjustments()
    written_fonts = {style_id: value["written"] for style_id, value in resolve_written_fonts(state).items()}
    document = Document()
    # Documented renderer rule: built-in Normal spacing/line defaults can
    # never leak (every paragraph also carries explicit control).
    normal = document.styles["Normal"]
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0

    # Measured page geometry -> Word section properties.
    section = document.sections[0]
    page = state.page
    section.page_width = Pt(round(float(page.width_pt), 3))
    section.page_height = Pt(round(float(page.height_pt), 3))
    section.top_margin = Pt(round(float(page.margin_top_pt), 3))
    section.bottom_margin = Pt(round(float(page.margin_bottom_pt), 3))
    section.left_margin = Pt(round(float(page.margin_left_pt), 3))
    section.right_margin = Pt(round(float(page.margin_right_pt), 3))

    rules_by_id = {rule.rule_id: rule.model_dump() for rule in state.rules}
    # Documented renderer rule (like unmeasured-color->black): a candidate-only
    # overflow heading has no measured gap of its own, so it consumes the
    # state's MEASURED section rhythm — the median heading gap of the mapped
    # sections — instead of cramming at 0pt. Measured state value, no per-pair
    # constant.
    measured_heading_gaps = [
        float(section.heading_gap_above_pt)
        for section in [*plan.sections, *plan.appended_sections]
        if not section.empty and section.heading_gap_above_pt is not None
    ]
    appended_heading_gap = (
        _median(measured_heading_gaps) if measured_heading_gaps else None
    )

    # Header region: one paragraph per measured row, fields in measured order.
    for row in plan.header_rows:
        if not row.fields:
            continue  # unfilled target slot row: dropped, recorded on the plan
        token = _style_of(state, row.style_id)
        paragraph = document.add_paragraph()
        for index, field in enumerate(row.fields):
            if index:
                separator_run = paragraph.add_run(f" {row.separator} " if row.separator else "   ")
                # Separator runs carry the row token too: an unformatted run
                # renders at the built-in default size and breaks the row.
                if token is not None:
                    _apply_token(separator_run, token, written_fonts.get(row.style_id))
            run = paragraph.add_run(field.text)
            if token is not None:
                _apply_token(run, token, written_fonts.get(row.style_id))
        _control_paragraph(
            paragraph, token,
            space_before_pt=row.gap_above_pt,
            alignment=row.alignment if row.alignment in {"left", "center", "right"} else "left",
        )
    if plan.header_overflow is not None:
        token = _style_of(state, plan.header_overflow.style_id)
        paragraph = document.add_paragraph()
        for index, field in enumerate(plan.header_overflow.fields):
            if index:
                paragraph.add_run("   ")
            run = paragraph.add_run(field.text)
            if token is not None:
                _apply_token(run, token, written_fonts.get(plan.header_overflow.style_id))
        _control_paragraph(
            paragraph, token, space_before_pt=plan.header_overflow.gap_above_pt
        )

    def emit_section(section_plan: Any) -> None:
        if section_plan.empty:
            return
        if (
            section_plan.candidate_only
            and section_plan.heading_gap_above_pt is None
            and appended_heading_gap is not None
        ):
            section_plan = section_plan.model_copy(
                update={"heading_gap_above_pt": appended_heading_gap}
            )
        rule = rules_by_id.get(section_plan.rule_id) if section_plan.rule_id else None
        heading_token = _style_of(state, section_plan.style_id)
        _write_heading(
            document,
            state,
            section_plan.label,
            section_plan.label_case,
            heading_token,
            rule,
            (section_plan.heading_gap_above_pt, section_plan.heading_gap_below_pt),
            written_font=written_fonts.get(section_plan.style_id) if section_plan.style_id else None,
            adjustments=adjustments,
            section_plan=section_plan,
        )
        content_token = _style_of(state, section_plan.content_style_id) or _style_of(state, "style.body")
        content_written = (
            written_fonts.get(section_plan.content_style_id)
            if section_plan.content_style_id
            else written_fonts.get("style.body")
        )
        if section_plan.content_kind == "paragraph":
            for line in section_plan.paragraph_lines:
                _write_text_paragraph(document, line.text, content_token, written_font=content_written)
            for styled_line in section_plan.styled_lines:
                # C2-0cC capability proof: one paragraph as ORDERED native
                # editable runs, each carrying its template-owned token's
                # measured formatting (w:color etc.). Text is candidate-owned
                # and rendered verbatim; ordered concatenation equals the
                # leaf text (accounting gate verifies it).
                paragraph = document.add_paragraph()
                for run in styled_line.runs:
                    run_token = _style_of(state, run.style_id) if run.style_id else None
                    run_object = paragraph.add_run(run.text)
                    _apply_token(run_object, run_token, written_fonts.get(run.style_id) if run.style_id else None)
                _control_paragraph(paragraph, content_token)
        # C2-0cM: a composite section renders its entries sub-content AND its
        # item sub-content under the ONE measured target heading/rule.
        if section_plan.content_kind in {"entries", "composite"}:
            title_token = _style_of(state, section_plan.title_style_id)
            detail_token = _style_of(state, section_plan.detail_style_id) or title_token
            meta_token = _style_of(state, section_plan.meta_style_id) or detail_token
            last_paragraph: Any = None
            for entry_index, entry in enumerate(section_plan.entries):
                inter_entry_gap = (
                    section_plan.inter_entry_gap_above_pt if entry_index else None
                )
                entry_paragraphs: list[Any] = []
                if entry.meta_lines:
                    # Left/right topology: borderless two-column table; the
                    # measured inter-entry rhythm renders as space AFTER the
                    # entry's last paragraph (a space-before inside the first
                    # cell would push only the left column down).
                    entry_paragraphs = _write_entry_table(
                        document, state, section_plan, entry,
                        title_token, detail_token, meta_token,
                        written_fonts=written_fonts,
                        adjustments=adjustments,
                    )
                else:
                    # No metadata: plain stacked paragraphs (nothing to split),
                    # still indented onto the measured entry-column x0 so the
                    # rendered left column matches the target topology.
                    entry_fit = adjustments.section(section_plan.node_id)
                    entry_left_indent = round(
                        float(section_plan.base_x0_pt) - float(state.page.margin_left_pt)
                        + entry_fit.entry_table_indent_pt,
                        3,
                    ) if section_plan.base_x0_pt is not None else None
                    for line_index, line in enumerate(entry.title_lines):
                        entry_paragraphs.append(
                            _write_text_paragraph(
                                document, line.text,
                                title_token if line_index == 0 else detail_token,
                                written_font=written_fonts.get(
                                    (section_plan.title_style_id if line_index == 0 else section_plan.detail_style_id)
                                    or section_plan.title_style_id
                                ) if (section_plan.title_style_id or section_plan.detail_style_id) else None,
                                left_indent_pt=entry_left_indent,
                                keep_with_next=bool(entry.bullet_items or entry.text_lines),
                            )
                        )
                if inter_entry_gap is not None and last_paragraph is not None:
                    last_paragraph.paragraph_format.space_after = Pt(
                        round(float(inter_entry_gap) + adjustments.section(section_plan.node_id).inter_entry_pt, 3)
                    )
                for item in entry.bullet_items:
                    entry_paragraphs.append(
                        _write_native_bullet(
                            document, state, section_plan, item.text, content_token,
                            written_font=content_written, adjustments=adjustments,
                        )
                    )
                for line in entry.text_lines:
                    # C2-0cR root-cause fix (general, not pair-specific): entry
                    # child detail/bullet lines align with their entry's
                    # content column (the same measured entry-column x0 the
                    # entry rows render on), not with the page margin. The
                    # source glyphs (e.g. a leading "•") stay verbatim TEXT —
                    # never converted into presentation bullets; the fitter's
                    # entry_child_text_indent_pt control shifts the whole
                    # paragraph onto a target-measured child anchor when the
                    # target measures a different child x.
                    fit = adjustments.section(section_plan.node_id)
                    child_left_indent = (
                        round(
                            float(section_plan.base_x0_pt)
                            - float(state.page.margin_left_pt)
                            + fit.entry_child_text_indent_pt,
                            3,
                        )
                        if section_plan.base_x0_pt is not None
                        else (fit.entry_child_text_indent_pt or None)
                    )
                    entry_paragraphs.append(
                        _write_text_paragraph(
                            document, line.text, content_token, written_font=content_written,
                            left_indent_pt=child_left_indent,
                        )
                    )
                # C2 nested-entry spike: titled sub-groups with their own
                # bullets, in candidate document order, after the entry's own
                # title/meta and any entry-level bullets. The title tier style
                # is the measured state declaration (subgroup_title_style_id);
                # subgroup bullets consume the section's measured bullet tiers.
                for subgroup in entry.subgroups:
                    subgroup_token = _style_of(state, section_plan.subgroup_title_style_id) or (
                        detail_token or title_token
                    )
                    subgroup_indent = (
                        round(
                            float(section_plan.base_x0_pt) - float(state.page.margin_left_pt),
                            3,
                        )
                        if section_plan.base_x0_pt is not None
                        else None
                    )
                    entry_paragraphs.append(
                        _write_text_paragraph(
                            document, subgroup.title.text, subgroup_token,
                            written_font=(
                                written_fonts.get(section_plan.subgroup_title_style_id)
                                if section_plan.subgroup_title_style_id
                                else None
                            ),
                            left_indent_pt=subgroup_indent,
                            keep_with_next=bool(subgroup.bullet_items or subgroup.text_lines) or None,
                        )
                    )
                    for item in subgroup.bullet_items:
                        entry_paragraphs.append(
                            _write_native_bullet(
                                document, state, section_plan, item.text, content_token,
                                written_font=content_written, adjustments=adjustments,
                            )
                        )
                    for line in subgroup.text_lines:
                        subgroup_fit = adjustments.section(section_plan.node_id)
                        subgroup_child_indent = (
                            round(
                                float(section_plan.base_x0_pt)
                                - float(state.page.margin_left_pt)
                                + subgroup_fit.entry_child_text_indent_pt,
                                3,
                            )
                            if section_plan.base_x0_pt is not None
                            else (subgroup_fit.entry_child_text_indent_pt or None)
                        )
                        entry_paragraphs.append(
                            _write_text_paragraph(
                                document, line.text, content_token, written_font=content_written,
                                left_indent_pt=subgroup_child_indent,
                            )
                        )
                last_paragraph = entry_paragraphs[-1] if entry_paragraphs else last_paragraph
        if section_plan.category_grid_cells:
            _write_category_grid(
                document, state, section_plan, content_token, content_written,
                written_fonts=written_fonts, adjustments=adjustments,
            )
        if section_plan.content_kind in {"item_list", "inline_items", "composite"}:
            for item in section_plan.items:
                if section_plan.bullet_marker == "bullet":
                    _write_native_bullet(
                        document, state, section_plan, item.text, content_token,
                        written_font=content_written, adjustments=adjustments,
                    )
                else:
                    _write_text_paragraph(
                        document, item.text, content_token, written_font=content_written,
                        left_indent_pt=adjustments.section(section_plan.node_id).item_left_indent_pt,
                    )

    for section_plan in plan.sections:
        emit_section(section_plan)
    for section_plan in plan.appended_sections:
        emit_section(section_plan)
    return document


def _write_native_bullet(
    document: Any,
    state: C2LayoutState,
    section_plan: Any,
    text: str,
    token: StyleToken | None,
    written_font: str | None = None,
    adjustments: FitAdjustments | None = None,
) -> Any:
    """A real Word list paragraph with measured hanging-indent geometry.

    A confirmed leading presentation marker (bullet glyph or dash) is
    converted into the native bullet; arrows and all other content stay
    verbatim (see ``strip_presentation_marker``). Measured absolute bullet
    tiers become paragraph indents relative to the page margin."""
    adjustments = adjustments or FitAdjustments()
    rendered = strip_presentation_marker(text, native_bullet=True)
    paragraph = document.add_paragraph(style="List Bullet")
    run = paragraph.add_run(rendered)
    if token is not None:
        _apply_token(run, token, written_font)
    left_indent, first_line = _bullet_indents(state, section_plan, adjustments)
    _control_paragraph(
        paragraph, token,
        left_indent_pt=left_indent,
        first_line_indent_pt=first_line,
        keep_with_next=False,
    )
    return paragraph


def _bullet_indents(
    state: C2LayoutState, section_plan: Any, adjustments: FitAdjustments | None = None
) -> tuple[float | None, float | None]:
    """Measured absolute bullet tiers -> paragraph indents (relative to the
    page margin); fitter corrections shift them onto the target-measured
    anchors when the state declares no tiers. Returns
    (left_indent_pt, first_line_indent_pt)."""
    adjustments = adjustments or FitAdjustments()
    fit = adjustments.section(section_plan.node_id)
    if (
        section_plan.bullet_text_x0_pt is None
        or section_plan.base_x0_pt is None
    ):
        if fit.bullet_tiers is None:
            return None, None
        left = round(float(fit.bullet_tiers.text_x0_pt) - float(state.page.margin_left_pt), 3)
        first_line = round(float(fit.bullet_tiers.marker_x0_pt) - float(fit.bullet_tiers.text_x0_pt), 3)
        return left, first_line
    left = round(
        float(section_plan.bullet_text_x0_pt) - float(state.page.margin_left_pt)
        + fit.bullet_text_correction_pt, 3
    )
    first_line = (
        round(
            float(section_plan.bullet_dot_x0_pt) - float(section_plan.bullet_text_x0_pt)
            + fit.bullet_marker_correction_pt, 3
        )
        if section_plan.bullet_dot_x0_pt is not None
        else None
    )
    return left, first_line


def deterministic_docx_bytes(document: Document) -> bytes:
    """Serialize the package with normalized zip metadata so identical
    documents produce identical bytes (known volatile package metadata only)."""
    buffer = BytesIO()
    document.save(buffer)
    with zipfile.ZipFile(BytesIO(buffer.getvalue())) as source:
        entries = [(info.filename, source.read(info.filename)) for info in source.infolist()]
    output = BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for name, data in entries:
            info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(info, data)
    return output.getvalue()


# ---------------------------------------------------------------------------
# Inspection: from the OOXML/python-docx, never inferred from source code
# ---------------------------------------------------------------------------


def _iter_block_paragraphs(document: Document):
    """Every paragraph in ACTUAL document order — including paragraphs inside
    table cells (``document.paragraphs`` silently omits table content, which
    would let accounting miss rendered leaves). Yields (in_table, paragraph)."""
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield False, Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            table = Table(child, document)
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        yield True, paragraph


def inspect_docx(path: Path) -> dict[str, Any]:
    """Structural inspection of the written package (opens it again)."""
    if not zipfile.is_zipfile(path):
        return {"valid_package": False, "reason": "not a zip container"}
    with zipfile.ZipFile(path) as package:
        names = set(package.namelist())
        if "[Content_Types].xml" not in names or "word/document.xml" not in names:
            return {"valid_package": False, "reason": "missing OOXML package parts"}
        document_xml = package.read("word/document.xml").decode("utf-8")
        styles_xml = package.read("word/styles.xml").decode("utf-8")
    try:
        document = Document(str(path))
    except Exception as error:  # noqa: BLE001 — inspection reports honestly
        return {"valid_package": False, "reason": f"python-docx open failed: {error}"}
    paragraphs = []
    for in_table, paragraph in _iter_block_paragraphs(document):
        p_pr = paragraph._p.pPr  # noqa: SLF001
        paragraphs.append(
            {
                "style": paragraph.style.name if paragraph.style is not None else None,
                "text": paragraph.text,
                "alignment": str(paragraph.alignment) if paragraph.alignment is not None else None,
                "explicit_spacing": p_pr is not None and p_pr.find(qn("w:spacing")) is not None,
                "in_table": in_table,
                "explicit_font_runs": sum(
                    1 for run in paragraph.runs if run.font.name is not None
                ),
                # Native w:color run values as written (C2-0cC): the authored
                # color evidence per paragraph, inspected from the package.
                "run_colors": [
                    (
                        str(run.font.color.rgb)
                        if run.font.color is not None and run.font.color.rgb is not None
                        else None
                    )
                    for run in paragraph.runs
                ],
            }
        )
    # Mark table paragraphs (evidence that table content is inspected).
    table_records: list[dict[str, Any]] = []
    for child in document.element.body.iterchildren():
        if child.tag != qn("w:tbl"):
            continue
        table = Table(child, document)
        widths = [
            round(cell.width.pt, 3)
            for column in table.columns
            for cell in column.cells[:1]
        ]
        tbl_xml = table._tbl.xml  # noqa: SLF001
        table_records.append(
            {
                "columns": len(table.columns),
                "rows": len(table.rows),
                "column_widths_pt": widths,
                "borders_none": 'w:val="none"' in tbl_xml and 'w:val="single"' not in tbl_xml,
                "rows_cannot_split": "<w:cantSplit/>" in tbl_xml,
            }
        )
    section = document.sections[0]
    geometry = {
        "page_width_pt": round(section.page_width.pt, 3),
        "page_height_pt": round(section.page_height.pt, 3),
        "margin_top_pt": round(section.top_margin.pt, 3),
        "margin_bottom_pt": round(section.bottom_margin.pt, 3),
        "margin_left_pt": round(section.left_margin.pt, 3),
        "margin_right_pt": round(section.right_margin.pt, 3),
    }
    rendered_paragraphs = [record for record in paragraphs if _norm(record["text"])]
    return {
        "valid_package": True,
        "package_parts": len(names),
        "paragraph_count": len(paragraphs),
        "paragraphs": paragraphs,
        "rendered_paragraph_count": len(rendered_paragraphs),
        "explicit_spacing_count": sum(1 for record in paragraphs if record["explicit_spacing"]),
        "heading_paragraphs": [
            record["text"] for record in paragraphs if (record["style"] or "").startswith("Heading")
        ],
        "list_paragraphs": [
            record["text"] for record in paragraphs if (record["style"] or "").startswith("List ")
        ],
        "list_style_numbered_in_styles_xml": "numPr" in styles_xml,
        "paragraph_borders_in_document_xml": document_xml.count("<w:pBdr>"),
        "explicit_font_run_count": sum(
            record["explicit_font_runs"] for record in paragraphs
        ),
        "character_spacing_runs_in_document_xml": document_xml.count('<w:spacing w:val='),
        "tables": table_records,
        "section_geometry": geometry,
    }


def expected_paragraphs(plan: Any) -> list[dict[str, Any]]:
    """The plan's deterministic reading order as the exact expected paragraph
    sequence, each entry carrying the candidate leaves it must render and the
    RENDERED text (confirmed presentation markers become native bullets).
    Measurement metadata (``section_node_id``/``style_id``/``tier``) rides on
    every entry for the rendered-geometry mapping."""
    paragraphs: list[dict[str, Any]] = []
    for row in plan.header_rows:
        if row.fields:
            paragraphs.append(
                {
                    "kind": "header_row",
                    "text": _header_row_text(row.fields, row.separator),
                    "leaf_ids": [field.leaf_id for field in row.fields],
                    "marker_conversions": {},
                    "native_bullet": False,
                    "section_node_id": None,
                    "entry_index": None,
                    "style_id": row.style_id,
                    "tier": "header",
                    "alignment": row.alignment,
                }
            )
    if plan.header_overflow is not None and plan.header_overflow.fields:
        paragraphs.append(
            {
                "kind": "header_overflow",
                "text": _header_row_text(plan.header_overflow.fields, None),
                "leaf_ids": [field.leaf_id for field in plan.header_overflow.fields],
                "marker_conversions": {},
                "native_bullet": False,
                "section_node_id": None,
                "entry_index": None,
                "style_id": plan.header_overflow.style_id,
                "tier": "header",
                "alignment": "left",
            }
        )
    for section_plan in [*plan.sections, *plan.appended_sections]:
        if section_plan.empty:
            continue
        paragraphs.append(
            {
                "kind": "heading",
                "text": section_plan.label,
                "leaf_ids": [],
                "marker_conversions": {},
                "native_bullet": False,
                "section_node_id": section_plan.node_id,
                "entry_index": None,
                "style_id": section_plan.style_id,
                "tier": "heading",
                "alignment": "left",
            }
        )
        for line in section_plan.paragraph_lines:
            paragraphs.append(
                {
                    "kind": "paragraph",
                    "text": line.text,
                    "leaf_ids": [line.leaf_id],
                    "marker_conversions": {},
                    "native_bullet": False,
                    "section_node_id": section_plan.node_id,
                    "entry_index": None,
                    "style_id": section_plan.content_style_id or "style.body",
                    "tier": "content",
                    "alignment": "left",
                }
            )
        for styled_line in section_plan.styled_lines:
            # One paragraph rendered as ordered styled runs (C2-0cC); the
            # paragraph's text is the ordered run concatenation (the leaf's
            # verbatim text) — accounting verifies it exactly.
            paragraphs.append(
                {
                    "kind": "styled_line",
                    "text": styled_line.text,
                    "leaf_ids": [styled_line.leaf_id],
                    "marker_conversions": {},
                    "native_bullet": False,
                    "section_node_id": section_plan.node_id,
                    "entry_index": None,
                    "style_id": section_plan.content_style_id or "style.body",
                    "tier": "content",
                    "alignment": "left",
                }
            )
        for entry_index, entry in enumerate(section_plan.entries):
            title_measured = bool(section_plan.title_style_id)
            meta_measured = bool(
                section_plan.meta_style_id or section_plan.detail_style_id or section_plan.title_style_id
            )
            for line_index, line in enumerate(entry.title_lines):
                paragraphs.append(
                    {
                        "kind": "entry_title",
                        "text": line.text,
                        "leaf_ids": [line.leaf_id],
                        "marker_conversions": {},
                        "native_bullet": False,
                        "measured_token": title_measured,
                        "section_node_id": section_plan.node_id,
                        "entry_index": entry_index,
                        "style_id": (
                            section_plan.title_style_id if line_index == 0
                            else section_plan.detail_style_id or section_plan.title_style_id
                        ),
                        "tier": "title" if line_index == 0 else "detail",
                        "alignment": "left",
                    }
                )
            for line in entry.meta_lines:
                paragraphs.append(
                    {
                        "kind": "entry_meta",
                        "text": line.text,
                        "leaf_ids": [line.leaf_id],
                        "marker_conversions": {},
                        "native_bullet": False,
                        "measured_token": meta_measured,
                        "section_node_id": section_plan.node_id,
                        "entry_index": entry_index,
                        "style_id": section_plan.meta_style_id or section_plan.detail_style_id or section_plan.title_style_id,
                        "tier": "meta",
                        "alignment": "right",
                    }
                )
            for line in entry.bullet_items:
                rendered = strip_presentation_marker(line.text, native_bullet=True)
                paragraphs.append(
                    {
                        "kind": "bullet",
                        "text": rendered,
                        "leaf_ids": [line.leaf_id],
                        "marker_conversions": {line.leaf_id: line.text != rendered},
                        "native_bullet": True,
                        "section_node_id": section_plan.node_id,
                        "entry_index": entry_index,
                        "style_id": section_plan.content_style_id or "style.body",
                        "tier": "bullet",
                        "alignment": "left",
                    }
                )
            for line in entry.text_lines:
                paragraphs.append(
                    {
                        "kind": "textline",
                        "text": line.text,
                        "leaf_ids": [line.leaf_id],
                        "marker_conversions": {},
                        "native_bullet": False,
                        "section_node_id": section_plan.node_id,
                        "entry_index": entry_index,
                        "style_id": section_plan.content_style_id or "style.body",
                        "tier": "content",
                        "alignment": "left",
                    }
                )
            # C2 nested-entry spike: titled sub-groups after the entry's own
            # title/meta and any entry-level bullets (document order; the plan
            # compiler fails closed on any other order).
            for subgroup in entry.subgroups:
                paragraphs.append(
                    {
                        "kind": "subgroup_title",
                        "text": subgroup.title.text,
                        "leaf_ids": [subgroup.title.leaf_id],
                        "marker_conversions": {},
                        "native_bullet": False,
                        "measured_token": bool(section_plan.subgroup_title_style_id),
                        "section_node_id": section_plan.node_id,
                        "entry_index": entry_index,
                        "style_id": (
                            section_plan.subgroup_title_style_id
                            or section_plan.title_style_id
                            or section_plan.content_style_id
                            or "style.body"
                        ),
                        "tier": "subgroup_title",
                        "alignment": "left",
                    }
                )
                for line in subgroup.bullet_items:
                    rendered = strip_presentation_marker(line.text, native_bullet=True)
                    paragraphs.append(
                        {
                            "kind": "bullet",
                            "text": rendered,
                            "leaf_ids": [line.leaf_id],
                            "marker_conversions": {line.leaf_id: line.text != rendered},
                            "native_bullet": True,
                            "section_node_id": section_plan.node_id,
                            "entry_index": entry_index,
                            "style_id": section_plan.content_style_id or "style.body",
                            "tier": "bullet",
                            "alignment": "left",
                        }
                    )
                for line in subgroup.text_lines:
                    paragraphs.append(
                        {
                            "kind": "textline",
                            "text": line.text,
                            "leaf_ids": [line.leaf_id],
                            "marker_conversions": {},
                            "native_bullet": False,
                            "section_node_id": section_plan.node_id,
                            "entry_index": entry_index,
                            "style_id": section_plan.content_style_id or "style.body",
                            "tier": "content",
                            "alignment": "left",
                        }
                    )
        for line in section_plan.items:
            native_bullet = section_plan.bullet_marker == "bullet"
            rendered = strip_presentation_marker(line.text, native_bullet=native_bullet)
            paragraphs.append(
                {
                    "kind": "item",
                    "text": rendered,
                    "leaf_ids": [line.leaf_id],
                    "marker_conversions": {line.leaf_id: line.text != rendered},
                    "native_bullet": native_bullet,
                    "section_node_id": section_plan.node_id,
                    "entry_index": None,
                    "style_id": section_plan.content_style_id or "style.body",
                    "tier": "item",
                    "alignment": "left",
                }
            )
        # C2-0cS: measured category-grid cells — one paragraph per CELL
        # FRAGMENT in document order (row-major, label then value per cell);
        # both fragments carry the same leaf id and their ordered
        # concatenation equals the leaf text (accounting verifies it).
        for cell in section_plan.category_grid_cells:
            cell_id = f"{section_plan.node_id}.category.r{cell.row_index}c{cell.column_index}"
            paragraphs.append(
                {
                    "kind": "grid_label",
                    "text": cell.label_text,
                    "leaf_ids": [cell.leaf_id],
                    "marker_conversions": {},
                    "native_bullet": False,
                    "section_node_id": section_plan.node_id,
                    "entry_index": None,
                    "style_id": cell.label_style_id or "style.body",
                    "tier": "content",
                    "alignment": "right",
                    "grid_row": cell.row_index,
                    "grid_column": cell.column_index,
                    "grid_cell_id": cell_id,
                }
            )
            paragraphs.append(
                {
                    "kind": "grid_value",
                    "text": cell.value_text,
                    "leaf_ids": [cell.leaf_id],
                    "marker_conversions": {},
                    "native_bullet": False,
                    "section_node_id": section_plan.node_id,
                    "entry_index": None,
                    "style_id": cell.value_style_id or section_plan.content_style_id or "style.body",
                    "tier": "content",
                    "alignment": "left",
                    "grid_row": cell.row_index,
                    "grid_column": cell.column_index,
                    "grid_cell_id": cell_id,
                }
            )
    return paragraphs


def expected_visual_rows(plan: Any) -> list[dict[str, Any]]:
    """The plan's reading order as expected VISUAL rows.

    A table entry renders title_lines[i] (left column) and meta_lines[i]
    (right column) on ONE visual baseline, so the rendered-geometry mapping
    consumes them as a single row carrying both tiers. Bullets, text lines,
    headings, and header rows map one row each, in document order."""
    rows: list[dict[str, Any]] = []
    plan_sections = [*plan.sections, *plan.appended_sections]
    paired_meta_leaf_ids: set[str] = set()
    for paragraph in expected_paragraphs(plan):
        if paragraph["section_node_id"] is None:  # header rows / overflow
            rows.append(paragraph)
    for section_plan in plan_sections:
        if section_plan.empty:
            continue
        section_paragraphs = [
            paragraph for paragraph in expected_paragraphs(plan)
            if paragraph["section_node_id"] == section_plan.node_id
        ]
        for paragraph in section_paragraphs:
            if paragraph["kind"] in {"grid_label", "grid_value"}:
                continue  # folded into their grid visual row below
            if paragraph["kind"] == "entry_meta":
                # Meta lines ride on their title's visual row when paired.
                if not (set(paragraph["leaf_ids"]) & paired_meta_leaf_ids):
                    rows.append(paragraph)
                continue
            if paragraph["kind"] != "entry_title":
                rows.append(paragraph)
            elif not _pair_with_meta(section_plan, paragraph, rows, paired_meta_leaf_ids):
                rows.append(paragraph)
        # C2-0cS: a grid row's label/value fragments share ONE visual
        # baseline per table row; the mapping consumes them as one row,
        # AFTER the section heading (document order).
        grid_fragments = [
            paragraph for paragraph in section_paragraphs
            if paragraph["kind"] in {"grid_label", "grid_value"}
        ]
        if grid_fragments:
            by_row = {}
            for fragment in grid_fragments:
                by_row.setdefault(fragment["grid_row"], []).append(fragment)
            for row_index in sorted(by_row):
                fragments = sorted(by_row[row_index], key=lambda p: p["grid_column"])
                rows.append(
                    {
                        "kind": "grid_row",
                        "text": " ".join(fragment["text"] for fragment in fragments).strip(),
                        "leaf_ids": sorted(
                            {leaf_id for fragment in fragments for leaf_id in fragment["leaf_ids"]}
                        ),
                        "marker_conversions": {},
                        "native_bullet": False,
                        "section_node_id": section_plan.node_id,
                        "entry_index": None,
                        "style_id": fragments[0]["style_id"],
                        "tier": "content",
                        "alignment": "left",
                        "grid_cells": [
                            {
                                "leaf_id": fragment["leaf_ids"][0],
                                "kind": fragment["kind"],
                                "text": fragment["text"],
                                "grid_row": fragment["grid_row"],
                                "grid_column": fragment["grid_column"],
                                "style_id": fragment["style_id"],
                            }
                            for fragment in fragments
                        ],
                    }
                )
    return rows


def _pair_with_meta(
    section_plan: Any,
    title_paragraph: dict[str, Any],
    rows: list[dict[str, Any]],
    paired_meta_leaf_ids: set[str],
) -> bool:
    """Fold one title line and its same-index meta line into a single visual
    row appended to ``rows`` (False when the entry has no meta on that row)."""
    entry = section_plan.entries[title_paragraph["entry_index"]]
    title_index = next(
        (
            index
            for index, line in enumerate(entry.title_lines)
            if line.text == title_paragraph["text"]
        ),
        None,
    )
    if title_index is None or title_index >= len(entry.meta_lines):
        return False
    meta_line = entry.meta_lines[title_index]
    paired_meta_leaf_ids.add(meta_line.leaf_id)
    rows.append(
        {
            **title_paragraph,
            "kind": "entry_row",
            "text": f"{title_paragraph['text']} {meta_line.text}",
            "meta_text": meta_line.text,
            "meta_style_id": (
                section_plan.meta_style_id or section_plan.detail_style_id or section_plan.title_style_id
            ),
            "meta_leaf_ids": [meta_line.leaf_id],
            "leaf_ids": [*title_paragraph["leaf_ids"], meta_line.leaf_id],
        }
    )
    return True


def expected_reading_order(plan: Any) -> list[str]:
    """The plan's deterministic reading order as flat paragraph text (empty
    fragments — e.g. an unused category-grid label cell — never enter the
    reading order; the document's own empty paragraphs are filtered alike)."""
    return [
        paragraph["text"] for paragraph in expected_paragraphs(plan) if _norm(paragraph["text"])
    ]


def content_accounting(plan: Any, inspection: dict[str, Any]) -> dict[str, Any]:
    """Exact candidate-content accounting from the written DOCX.

    The expected paragraph sequence (one entry per rendered paragraph, each
    carrying its candidate leaves and their rendered text) is aligned
    position-for-position against the document paragraphs traversed in actual
    document order (table content included); every ledger leaf must be
    carried by exactly one expected paragraph and occur exactly once within
    it. Omissions must never appear. Header-field values are matched as
    substrings of their row paragraph; every body leaf owns its whole
    paragraph; a confirmed presentation marker converted into a native bullet
    is recorded per leaf (the substantive text is verified verbatim).
    """
    expected = expected_paragraphs(plan)
    document_texts = [_norm(paragraph["text"]) for paragraph in inspection["paragraphs"]]
    nonempty_expected = [paragraph for paragraph in expected if _norm(paragraph["text"])]
    nonempty_document = [text for text in document_texts if text]
    aligned = (
        len(nonempty_expected) == len(nonempty_document)
        and all(
            _norm(paragraph["text"]) == text
            for paragraph, text in zip(nonempty_expected, nonempty_document)
        )
    )
    records: dict[str, Any] = {}
    missing: list[str] = []
    duplicated: list[str] = []
    marker_conversions: dict[str, bool] = {}
    for leaf_id in plan.leaf_ledger:
        owners = [paragraph for paragraph in expected if leaf_id in paragraph["leaf_ids"]]
        # C2-0cS: a category-grid leaf renders EXACTLY ONCE as its ordered
        # cell fragments; the document-order fragment concatenation must
        # equal the leaf's verbatim text (the position-for-position
        # alignment above already places each fragment at its position).
        fragment_owners = [
            paragraph for paragraph in owners
            if paragraph["kind"] in {"grid_label", "grid_value"}
        ]
        if owners and len(owners) == len(fragment_owners):
            rendered = _norm(" ".join(paragraph["text"] for paragraph in fragment_owners))
            expected_text = _norm(_leaf_text(plan, leaf_id))
            ok = bool(aligned and rendered == expected_text and rendered)
            records[leaf_id] = {
                "occurrences": 1 if ok else 0,
                "rendered_exactly_once": ok,
                "owning_paragraphs": len(fragment_owners),
                "grid_fragments": True,
            }
            if not ok:
                missing.append(leaf_id)
            continue
        rendered_text = owners[0]["text"] if len(owners) == 1 else _leaf_text(plan, leaf_id)
        token = _norm(rendered_text)
        if len(owners) != 1 or not aligned:
            count = 0 if len(owners) != 1 else sum(text.count(token) for text in nonempty_document)
            records[leaf_id] = {
                "occurrences": count,
                "rendered_exactly_once": False,
                "owning_paragraphs": len(owners),
            }
            (duplicated if len(owners) > 1 else missing).append(leaf_id)
            continue
        count = _norm(owners[0]["text"]).count(token)
        converted = owners[0]["marker_conversions"].get(leaf_id, False)
        records[leaf_id] = {
            "occurrences": count,
            "rendered_exactly_once": count == 1,
            "owning_paragraphs": 1,
            "presentation_marker_converted": converted,
        }
        if converted:
            marker_conversions[leaf_id] = True
        if count != 1:
            (duplicated if count > 1 else missing).append(leaf_id)
    omission_leaks = [
        {"text": omission.text, "occurrences": sum(text.count(_norm(omission.text)) for text in document_texts)}
        for omission in plan.explicit_omissions
        if sum(text.count(_norm(omission.text)) for text in document_texts)
    ]
    passed = not missing and not duplicated and not omission_leaks and aligned
    return {
        "passed": passed,
        "reading_order_aligned": aligned,
        "expected_paragraph_count": len(nonempty_expected),
        "document_paragraph_count": len(nonempty_document),
        "rendered_leaves": len(plan.leaf_ledger),
        "presentation_marker_conversions": marker_conversions,
        "explicitly_omitted": [
            {"text": omission.text, "reason": omission.reason} for omission in plan.explicit_omissions
        ],
        "missing": missing,
        "duplicated": duplicated,
        "omission_leaks": omission_leaks,
        "leaf_records": records,
    }


def reading_order_gate(plan: Any, inspection: dict[str, Any]) -> dict[str, Any]:
    """The DOCX paragraph sequence (tables traversed in document order) must
    equal the plan's reading order."""
    document_texts = [
        _norm(paragraph["text"])
        for paragraph in inspection["paragraphs"]
        if _norm(paragraph["text"])
    ]
    expected = [_norm(text) for text in expected_reading_order(plan)]
    return {
        "passed": document_texts == expected,
        "expected_count": len(expected),
        "document_count": len(document_texts),
        "first_mismatch": next(
            (
                {"index": index, "expected": expected[index], "document": document_texts[index]}
                for index in range(min(len(expected), len(document_texts)))
                if expected[index] != document_texts[index]
            ),
            None,
        )
        if document_texts != expected
        else None,
    }


def _leaf_text(plan: Any, leaf_id: str) -> str:
    if plan.header_overflow is not None:
        for field in plan.header_overflow.fields:
            if field.leaf_id == leaf_id:
                return field.text
    for row in plan.header_rows:
        for field in row.fields:
            if field.leaf_id == leaf_id:
                return field.text
    for section in [*plan.sections, *plan.appended_sections]:
        for line in [*section.paragraph_lines, *section.items]:
            if line.leaf_id == leaf_id:
                return line.text
        for styled_line in section.styled_lines:
            if styled_line.leaf_id == leaf_id:
                return styled_line.text
        # C2-0cS: a category-grid leaf's verbatim text is its ordered cell
        # fragments' concatenation.
        for cell in section.category_grid_cells:
            if cell.leaf_id == leaf_id:
                return cell.label_text + cell.value_text
        for entry in section.entries:
            for line in [*entry.title_lines, *entry.meta_lines, *entry.bullet_items, *entry.text_lines]:
                if line.leaf_id == leaf_id:
                    return line.text
            for subgroup in entry.subgroups:
                for line in [subgroup.title, *subgroup.bullet_items, *subgroup.text_lines]:
                    if line.leaf_id == leaf_id:
                        return line.text
    return ""


# ---------------------------------------------------------------------------
# Rendered DOCX geometry: measure -> map -> compare (work order Parts 1-2)
# ---------------------------------------------------------------------------

# Source/confirmed presentation markers (see ``strip_presentation_marker``)
# plus the glyphs LibreOffice actually renders for Word List Bullet numbering
# (the Symbol-font private-use bullet U+F0B7). The rendered set is for
# measurement and mapping only — it never rewrites candidate content.
_MARKER_GLYPHS = frozenset((*BULLET_GLYPH_MARKERS, *DASH_MARKERS))
_RENDERED_MARKER_GLYPHS = frozenset((*_MARKER_GLYPHS, "\uf0b7"))


def _strip_leading_marker_glyphs(text: str) -> str:
    stripped = text.lstrip()
    while stripped and stripped[0] in _RENDERED_MARKER_GLYPHS:
        stripped = stripped[1:].lstrip()
    return stripped

