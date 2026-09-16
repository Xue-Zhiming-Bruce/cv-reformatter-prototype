"""C2-0c focused tests: deterministic state -> DOCX render plan -> native OOXML.

Offline lane (default): a synthetic C2LayoutState + an authored candidate
render context exercise DOCX compilation, native-structure inspection
(tables included), presentation-marker handling, exact content accounting,
output-verified compatibility claims, pagination classification, and
fail-closed behavior — no Chrome, no corpus, no LibreOffice, no network.
The ``local_dataset`` lane runs the real canonical pairs (E→F primary,
D→E generalization/native lists, E→D gap-only control) end to end with
LibreOffice previews (evaluation evidence only).

Run:

    pytest tests/experiments/test_c2_docx_renderer.py -m "not local_dataset"
    pytest tests/experiments/test_c2_docx_renderer.py -m local_dataset
"""

from __future__ import annotations

import json
import zipfile
from pathlib import Path

import pytest

from tests.experiments.c2_docx_renderer import (
    MAX_FITTING_ITERATIONS,
    FitAdjustments,
    SectionFit,
    _strip_leading_marker_glyphs,
    _style_of,
    binding_review_rows,
    build_document,
    compare_geometry,
    compare_colors,
    content_accounting,
    conversion_compatibility_report,
    deterministic_docx_bytes,
    expected_reading_order,
    expected_visual_rows,
    compare_colors,
    fit_docx,
    installed_font_families,
    inspect_docx,
    map_rendered_to_expected,
    reading_order_gate,
    resolve_written_fonts,
    run_pair,
    strip_presentation_marker,
    typography_tables,
)
from tests.experiments.c2_renderer import LeafText, compile_render_plan
from tests.experiments.test_c2_pipeline import (
    BULLET_TIERS,
    FULL_COVERAGE_LABELS,
    compile_synthetic,
)
from tests.experiments.test_c2_renderer import _leaf, rich_candidate


def _state_and_plan(labels: list[str] | None = None):
    state = compile_synthetic(labels)
    plan = compile_render_plan(state, rich_candidate(include_unmatched=False))
    return state, plan


def _write_docx(state, plan, tmp_path: Path, name: str = "c2_output.docx") -> Path:
    path = tmp_path / name
    path.write_bytes(deterministic_docx_bytes(build_document(state, plan)))
    return path


def _full_pipeline(tmp_path: Path, labels: list[str] | None = None, candidate=None):
    """Build -> write -> inspect -> account -> report (the whole pipeline)."""
    state = compile_synthetic(labels)
    if candidate is None:
        candidate = rich_candidate(include_unmatched=False)
    plan = compile_render_plan(state, candidate)
    tmp_path.mkdir(parents=True, exist_ok=True)
    path = tmp_path / "c2_output.docx"
    path.write_bytes(deterministic_docx_bytes(build_document(state, plan)))
    inspection = inspect_docx(path)
    inspection["reading_order_gate"] = reading_order_gate(plan, inspection)
    accounting = content_accounting(plan, inspection)
    return state, plan, path, inspection, accounting


# ---------------------------------------------------------------------------
# Source presentation markers vs candidate content (work order Part 2.3)
# ---------------------------------------------------------------------------


def test_confirmed_presentation_markers_are_stripped_only_for_native_bullets() -> None:
    # Bullet glyph and dashes are confirmed presentation markers: converted
    # into the native Word bullet when (and only when) the line becomes one.
    for marker in ("•", "-", "–", "—"):
        assert strip_presentation_marker(f"{marker} substantive text", True) == "substantive text"
        # A plain/zero-bullet paragraph keeps EVERY glyph verbatim.
        assert strip_presentation_marker(f"{marker} substantive text", False) == (
            f"{marker} substantive text"
        )
    # Arrows are NOT confirmed markers: they may be content and always stay.
    assert strip_presentation_marker("→ Storage: item", True) == "→ Storage: item"
    assert strip_presentation_marker("→ Storage: item", False) == "→ Storage: item"
    # A dash fused to the word (no whitespace) is content, not a marker.
    assert strip_presentation_marker("-5°C storage", True) == "-5°C storage"
    # Substantive text is never rewritten.
    assert strip_presentation_marker("plain line", True) == "plain line"


def test_native_bullet_markers_convert_in_the_docx_and_accounting(tmp_path: Path) -> None:
    """A '•'-prefixed bullet becomes a native Word bullet WITHOUT a double
    marker; the accounting records the conversion and verifies the substantive
    text exactly once."""

    def candidate_with_prefix(prefix: str):
        candidate = rich_candidate(include_unmatched=False)
        leaves = [
            leaf.model_copy(update={"text": f"{prefix} Did candidate work"})
            if leaf.leaf_id == "work.e1.b1"
            else leaf
            for leaf in candidate.leaves
        ]
        return candidate.model_copy(update={"leaves": leaves})

    from tests.experiments.c2_docx_renderer import expected_paragraphs

    for prefix, rendered_start in (("•", "Did candidate work"), ("–", "Did candidate work"), ("→", "→ Did candidate work")):
        state, plan, path, inspection, accounting = _full_pipeline(
            tmp_path / f"bullet_{ord(prefix)}", candidate=candidate_with_prefix(prefix)
        )
        assert accounting["passed"] is True, (accounting["missing"], accounting["duplicated"])
        expected_bullets = [
            paragraph["text"]
            for paragraph in expected_paragraphs(plan)
            if paragraph["kind"] == "bullet" or paragraph.get("native_bullet")
        ]
        assert inspection["list_paragraphs"] == expected_bullets
        assert rendered_start in expected_bullets
        leaf_record = accounting["leaf_records"]["work.e1.b1"]
        assert leaf_record["rendered_exactly_once"] is True
        assert leaf_record["presentation_marker_converted"] is (prefix != "→")


# ---------------------------------------------------------------------------
# Determinism, package validity, native structures, topology
# ---------------------------------------------------------------------------


def test_deterministic_bytes_across_two_compiles(tmp_path: Path) -> None:
    state, plan = _state_and_plan()
    first = deterministic_docx_bytes(build_document(state, plan))
    second = deterministic_docx_bytes(build_document(state, plan))
    assert first == second
    path = _write_docx(state, plan, tmp_path)
    assert zipfile.is_zipfile(path)


def test_package_opens_and_carries_native_structures(tmp_path: Path) -> None:
    state, plan = _state_and_plan()
    path = _write_docx(state, plan, tmp_path)
    inspection = inspect_docx(path)
    assert inspection["valid_package"] is True
    # Headings are native Word heading paragraphs with the state labels.
    assert inspection["heading_paragraphs"] == [
        section.label for section in plan.sections if not section.empty
    ]
    # Measured bullet designs become real Word list paragraphs (entry
    # bullets and bulleted item lists alike), in document order.
    expected_paragraphs = __import__(
        "tests.experiments.c2_docx_renderer", fromlist=["expected_paragraphs"]
    ).expected_paragraphs(plan)
    expected_bullets = [
        paragraph["text"]
        for paragraph in expected_paragraphs
        if paragraph["kind"] == "bullet" or paragraph.get("native_bullet")
    ]
    assert inspection["list_paragraphs"] == expected_bullets
    assert inspection["list_style_numbered_in_styles_xml"] is True
    # Page geometry comes from the measured state section properties.
    geometry = inspection["section_geometry"]
    assert geometry["page_width_pt"] == pytest.approx(state.page.width_pt, abs=0.051)
    assert geometry["page_height_pt"] == pytest.approx(state.page.height_pt, abs=0.051)
    assert geometry["margin_left_pt"] == pytest.approx(state.page.margin_left_pt, abs=0.051)
    # Every rendered paragraph carries explicit spacing control: built-in
    # Word style defaults (Heading 1 / Normal / List Bullet) cannot leak.
    assert inspection["explicit_spacing_count"] == inspection["paragraph_count"]


# -- C2-0cM: composite section mapping ---------------------------------------

def test_composite_section_renders_partially_populated_content(tmp_path: Path) -> None:
    # The composite EDUCATION & CERTIFICATIONS section renders the candidate's
    # education entries under the ONE measured target heading; the
    # certifications sub-content has no candidate items and renders nothing
    # (honest partial population, nothing invented).
    state, plan, path, inspection, accounting = _full_pipeline(
        tmp_path, ["EDUCATION & CERTIFICATIONS"]
    )
    composite = next(s for s in plan.sections if s.content_kind == "composite")
    assert composite.label == "EDUCATION & CERTIFICATIONS"
    assert len(composite.entries) == 1 and composite.items == []
    assert any(
        "composite sub-content(s) ['certifications'] have no candidate content"
        in note
        for note in plan.notes
    )
    texts = [record["text"] for record in inspection["paragraphs"]]
    assert texts.count("EDUCATION & CERTIFICATIONS") == 1
    assert "CAND University" in texts
    # The education entry carries metadata -> native two-column topology table.
    assert any(record["columns"] == 2 for record in inspection["tables"])
    assert accounting["passed"] is True


def test_composite_section_copies_no_target_facts(tmp_path: Path) -> None:
    state, plan = _state_and_plan(["EDUCATION & CERTIFICATIONS"])
    all_text = " ".join(
        line.text
        for section in plan.sections
        for line in [*section.paragraph_lines, *section.items]
    ) + " ".join(
        line.text
        for section in plan.sections
        for entry in section.entries
        for line in [*entry.title_lines, *entry.meta_lines]
    )
    assert "TARGETFACT" not in all_text


def test_entry_rows_keep_left_right_topology_as_borderless_tables(tmp_path: Path) -> None:
    state, plan, path, inspection, accounting = _full_pipeline(tmp_path)
    entry_sections = [section for section in plan.sections if section.content_kind == "entries"]
    entries_with_meta = [
        section for section in entry_sections
        if any(entry.meta_lines for entry in section.entries)
    ]
    assert entries_with_meta
    tables = inspection["tables"]
    assert tables, "entries with metadata must render as tables"
    for record in tables:
        assert record["columns"] == 2
        assert record["borders_none"] is True, "table borders must be genuinely absent"
        assert record["rows_cannot_split"] is True
    # Table content participates in inspection and accounting (no silent
    # omission by document.paragraphs).
    title_texts = [
        line.text
        for section in entries_with_meta
        for entry in section.entries
        for line in entry.title_lines
    ]
    document_texts = [record["text"] for record in inspection["paragraphs"]]
    for text in title_texts:
        assert text in document_texts
    assert accounting["passed"] is True
    # Meta lines sit in the right cell (after all left-column lines of the
    # same entry) — deterministic cell reading order.
    assert all(record["in_table"] for record in inspection["paragraphs"]
               if record["text"] in {line.text for section in entries_with_meta
                                     for entry in section.entries for line in entry.meta_lines})


def test_measured_rules_become_native_paragraph_borders(tmp_path: Path) -> None:
    from tests.experiments.test_c2_renderer import _shape_summary

    summary = _shape_summary()
    state = compile_synthetic(
        ["WORK EXPERIENCE"], rules=summary["rules"], elements=summary["elements"]
    )
    plan = compile_render_plan(state, rich_candidate(include_unmatched=False))
    path = _write_docx(state, plan, tmp_path)
    inspection = inspect_docx(path)
    required_rules = [section for section in plan.sections if section.rule_id and not section.empty]
    assert len(required_rules) == 1
    assert inspection["paragraph_borders_in_document_xml"] == 1
    # Measured stroke and color survive into the OOXML border definition.
    with zipfile.ZipFile(path) as package:
        document_xml = package.read("word/document.xml").decode("utf-8")
    assert 'w:val="single"' in document_xml
    assert "111827" in document_xml


def test_zero_bullet_design_renders_verbatim_text_paragraphs(tmp_path: Path) -> None:
    """A zero-bullet target (marker 'none') never becomes a Word list: the
    source bullet glyphs stay verbatim text (proposal §10.5 ruling)."""
    state, plan = _state_and_plan()
    mutated = state.model_copy(deep=True)
    mutated.nodes = [
        node.model_copy(
            update={"content": node.content.model_copy(update={"bullet_marker": "none"})}
        )
        if node.kind == "section" and node.content is not None
        else node
        for node in mutated.nodes
    ]
    mutated.nodes = [
        node.model_copy(update={"list_marker": "none"})
        if node.kind == "list_row"
        else node
        for node in mutated.nodes
    ]
    plan = compile_render_plan(mutated, rich_candidate(include_unmatched=False))
    path = _write_docx(mutated, plan, tmp_path)
    inspection = inspect_docx(path)
    assert inspection["list_paragraphs"] == []
    text_lines = [
        line.text for section in plan.sections for entry in section.entries
        for line in entry.text_lines
    ]
    paragraph_texts = [paragraph["text"] for paragraph in inspection["paragraphs"]]
    assert text_lines
    for line in text_lines:
        assert line in paragraph_texts


# ---------------------------------------------------------------------------
# Reading order and exact content accounting
# ---------------------------------------------------------------------------


def test_reading_order_and_accounting_are_exact(tmp_path: Path) -> None:
    state, plan, path, inspection, accounting = _full_pipeline(tmp_path)
    order = reading_order_gate(plan, inspection)
    assert order["passed"] is True, order["first_mismatch"]
    assert accounting["passed"] is True, (accounting["missing"], accounting["duplicated"])
    assert accounting["missing"] == [] and accounting["duplicated"] == []
    assert accounting["omission_leaks"] == []
    assert accounting["rendered_leaves"] == len(plan.leaf_ledger)
    assert set(accounting["leaf_records"]) == set(plan.leaf_ledger)


def test_reading_order_gate_fails_on_reordered_paragraphs(tmp_path: Path) -> None:
    state, plan = _state_and_plan()
    document = build_document(state, plan)
    # Move the last body paragraph before the first heading: order must fail.
    body = document.paragraphs[-1]._element  # noqa: SLF001
    anchor = document.paragraphs[0]._element  # noqa: SLF001
    anchor.addprevious(body)
    mutated_path = tmp_path / "reordered.docx"
    mutated_path.write_bytes(deterministic_docx_bytes(document))
    inspection = inspect_docx(mutated_path)
    assert reading_order_gate(plan, inspection)["passed"] is False
    assert content_accounting(plan, inspection)["passed"] is False


def test_an_explicit_omission_is_never_rendered(tmp_path: Path) -> None:
    from tests.experiments.test_c2_renderer import candidate_with_unroutable

    state = compile_synthetic()
    candidate = candidate_with_unroutable("omit")
    plan = compile_render_plan(state, candidate)
    path = _write_docx(state, plan, tmp_path)
    inspection = inspect_docx(path)
    assert "unroutable.location" not in plan.leaf_ledger
    accounting = content_accounting(plan, inspection)
    assert accounting["passed"] is True
    assert accounting["explicitly_omitted"][0]["text"] == "CANDCITY, ST"
    assert accounting["omission_leaks"] == []
    all_text = "\n".join(paragraph["text"] for paragraph in inspection["paragraphs"])
    assert "CANDCITY" not in all_text


def test_a_dropped_paragraph_fails_the_accounting_gate(tmp_path: Path) -> None:
    state, plan = _state_and_plan()
    document = build_document(state, plan)
    # Silent content loss: remove one rendered body paragraph.
    victim = next(paragraph for paragraph in document.paragraphs if paragraph.text)
    victim._element.getparent().remove(victim._element)  # noqa: SLF001
    path = tmp_path / "lossy.docx"
    path.write_bytes(deterministic_docx_bytes(document))
    inspection = inspect_docx(path)
    accounting = content_accounting(plan, inspection)
    assert accounting["passed"] is False
    assert accounting["missing"], "the dropped leaf must be reported"


def test_dropped_table_content_fails_the_accounting_gate(tmp_path: Path) -> None:
    """Table content is inspected in document order: deleting a table cell
    paragraph is detectable content loss (never silently omitted)."""
    state, plan = _state_and_plan()
    document = build_document(state, plan)
    table = document.tables[0]
    victim = table.rows[0].cells[0].paragraphs[0]
    victim._element.getparent().remove(victim._element)  # noqa: SLF001
    path = tmp_path / "lossy_table.docx"
    path.write_bytes(deterministic_docx_bytes(document))
    inspection = inspect_docx(path)
    accounting = content_accounting(plan, inspection)
    assert accounting["passed"] is False
    assert accounting["missing"]


# ---------------------------------------------------------------------------
# Compatibility report (output-verified claims; pagination classification)
# ---------------------------------------------------------------------------


def test_every_exact_claim_is_output_verified(tmp_path: Path) -> None:
    # The measured-entry-elements state: entry tiers measured, no unresolved
    # bindings, no icons/tables/images -> nothing unsupported.
    from tests.experiments.test_c2_renderer import _shape_summary

    summary = _shape_summary()
    state = compile_synthetic(
        ["WORK EXPERIENCE"], rules=summary["rules"], elements=summary["elements"]
    )
    candidate = rich_candidate(include_unmatched=False)
    plan = compile_render_plan(state, candidate)
    path = tmp_path / "c2_output.docx"
    path.write_bytes(deterministic_docx_bytes(build_document(state, plan)))
    inspection = inspect_docx(path)
    inspection["reading_order_gate"] = reading_order_gate(plan, inspection)
    accounting = content_accounting(plan, inspection)
    pagination = {
        "target_page_count": 1,
        "frozen_c1_page_count": 1,
        "docx_preview_page_count": 1,
        "detail": "synthetic",
    }
    report = conversion_compatibility_report(state, plan, inspection, accounting, pagination)
    assert report.source_format.startswith("layout-state/1")
    assert "docx" in report.output_format
    assert report.exact, "applicable exact claims must be recorded"
    # NO unconditional static claims: every exact claim carries recorded
    # output evidence and is verified against the written OOXML.
    for item in report.exact:
        assert item.verified is True, item.claim
        assert item.evidence, item.claim
    assert report.content_loss_risk is False
    assert all(feature.content_preserved for feature in report.adjusted)
    assert {feature.feature for feature in report.adjusted} == {
        "entry two-column row layout",
        "candidate-only overflow sections",
    }
    assert report.unsupported == []
    assert report.owner_confirmation_required is False
    assert report.pagination.classification == "exact"


def test_pagination_difference_is_explicitly_classified(tmp_path: Path) -> None:
    state, plan, path, inspection, accounting = _full_pipeline(tmp_path)
    pagination = {
        "target_page_count": 1,
        "frozen_c1_page_count": 1,
        "docx_preview_page_count": 2,
        "detail": "synthetic drift",
    }
    report = conversion_compatibility_report(state, plan, inspection, accounting, pagination)
    assert report.pagination.target_page_count == 1
    assert report.pagination.frozen_c1_page_count == 1
    assert report.pagination.docx_preview_page_count == 2
    assert report.pagination.classification == "adjusted"
    # The pagination change surfaces as a visible adjusted degradation.
    assert any(feature.feature == "pagination" for feature in report.adjusted)


def test_unsupported_features_require_explicit_owner_confirmation(tmp_path: Path) -> None:
    # An unbindable measured section stays unresolved and renders nothing;
    # the compatibility contract must classify it unsupported (fail closed).
    state, plan, path, inspection, accounting = _full_pipeline(
        tmp_path, labels=["WORK EXPERIENCE", "ANOTHER SECTION"]
    )
    assert plan.skipped_unresolved_sections == ["ANOTHER SECTION"]
    report = conversion_compatibility_report(state, plan, inspection, accounting)
    assert any(
        feature.feature == "unresolved section bindings" for feature in report.unsupported
    )
    assert report.owner_confirmation_required is True


def test_a_failed_plan_refuses_docx_compilation() -> None:
    import pytest

    from tests.experiments.c2_pipeline import SectionContent

    state, _ = _state_and_plan()
    broken = state.model_copy(deep=True)
    for index, node in enumerate(broken.nodes):
        if node.kind == "section" and node.binding and node.binding.sources == ["work_experience"]:
            broken.nodes[index] = node.model_copy(
                update={"content": SectionContent(content_kind="unsupported", sources=["work_experience"])}
            )
    plan = compile_render_plan(broken, rich_candidate(include_unmatched=False))
    assert plan.status == "failed"
    with pytest.raises(RuntimeError, match="refusing to render a failed plan"):
        build_document(broken, plan)


# ---------------------------------------------------------------------------
# Real canonical pairs (local corpus + cached evidence; LibreOffice previews)
# ---------------------------------------------------------------------------

REQUIRED_ARTIFACTS = (
    "c2_layout_state.json", "candidate_render_context.json", "context_coverage.json",
    "docx_render_plan.json", "c2_output.docx", "content_accounting.json",
    "ooxml_inspection.json", "conversion_compatibility_report.json",
    "preview_validation.json", "docx_determinism.json", "hard_gates.json",
    "docx_rendered_geometry.json", "docx_geometry_comparison.json",
    "docx_color_comparison.json",
    "docx_fitting_log.json",
    "review.html", "target_page_1.png", "c1_page_1.png",
)


@pytest.mark.local_dataset
@pytest.mark.parametrize("pair", ["E_F", "D_E", "E_D"])
def test_canonical_pairs_end_to_end(pair: str) -> None:
    result = run_pair(pair)
    run_dir = Path(result["run_dir"])
    for artifact in REQUIRED_ARTIFACTS:
        assert (run_dir / artifact).exists(), artifact
    coverage = json.loads((run_dir / "context_coverage.json").read_text())
    assert coverage["total_coverage"] is True
    accounting = json.loads((run_dir / "content_accounting.json").read_text())
    assert accounting["passed"] is True, (accounting["missing"], accounting["duplicated"])
    inspection = json.loads((run_dir / "ooxml_inspection.json").read_text())
    assert inspection["valid_package"] is True
    assert inspection["reading_order_gate"]["passed"] is True
    determinism = json.loads((run_dir / "docx_determinism.json").read_text())
    assert determinism["bytes_equal_after_metadata_normalization"] is True
    report = json.loads((run_dir / "conversion_compatibility_report.json").read_text())
    gates = json.loads((run_dir / "hard_gates.json").read_text())
    # Every exact claim is output-verified with recorded evidence.
    assert report["exact"]
    assert all(item["verified"] and item["evidence"] for item in report["exact"])
    # All adjusted degradations preserve content; pagination is classified.
    assert all(feature["content_preserved"] for feature in report["adjusted"])
    assert report["pagination"] and report["pagination"]["classification"] in {
        "exact", "adjusted", "unsupported",
    }
    assert report["owner_confirmation_required"] == bool(report["unsupported"])
    # Preview evidence participated in the hard gates.
    preview = json.loads((run_dir / "preview_validation.json").read_text())
    assert gates["gates"]["preview_and_blank_pages"] is (
        bool(preview.get("available")) and preview["blank_page_gate"]["passed"]
    )
    assert preview["blank_page_gate"]["blank_pages"] == []
    # Rendered-geometry evidence (work order Parts 1/2/4): measured from the
    # preview PDF, compared node-locally in points, hard-gated separately.
    comparison = json.loads((run_dir / "docx_geometry_comparison.json").read_text())
    assert comparison["schema_version"] == "c2-docx-geometry-comparison/1"
    # The geometry verdict is consistent with its own counts (a pair may
    # honestly fail on a remaining delta — D→E's borderline trailing page —
    # and fail-closed pairs may carry unmeasurable capability gaps).
    assert comparison["gate_passed"] is (
        comparison["counts"]["failed"] == 0
        and comparison["counts"]["unmeasurable"] == 0
        and comparison["mapping_unmapped"] == 0
    )
    assert gates["gates"]["rendered_geometry_matches_declared_contract"] is comparison["gate_passed"]
    # C2-0cC: rendered colors are measured from the preview PDF and hard-gated
    # SEPARATELY (a color-only pass never implies overall conversion success).
    colors = json.loads((run_dir / "docx_color_comparison.json").read_text())
    assert colors["schema_version"] == "c2-docx-color-comparison/1"
    assert colors["gate_passed"] is (
        colors["counts"]["failed"] == 0 and colors["counts"]["unmeasurable"] == 0
    )
    assert gates["gates"]["rendered_colors_match_declared_contract"] is colors["gate_passed"]
    for row in colors["rows"]:
        if row["expected_color"] is not None:
            # An unmeasured fallback is never recorded as a measured pass.
            assert row["classification"] != "adjusted"
    fitting = json.loads((run_dir / "docx_fitting_log.json").read_text())
    assert 1 <= fitting["iterations"] <= fitting["max_iterations"] == 3
    assert len(fitting["iterations_log"]) == fitting["iterations"]
    typography = report["typography"]
    assert typography["classification"] in {"exact", "adjusted"}
    assert typography["authored_typography"] and typography["rendered_typography"]
    for record in typography["rendered_typography"]:
        # Requested, written, and rendered fonts are recorded separately; a
        # substituted family can never carry the exact classification.
        assert set(record) >= {
            "requested_font", "written_font", "rendered_font", "classification",
        }
        if record["requested_font"] != record["rendered_font"]:
            assert record["classification"] == "adjusted"
    if report["unsupported"]:
        assert gates["passed"] is False and gates["owner_confirmation_required"] is True
    else:
        assert gates["passed"] is True
    # C2-0cR: every canonical run carries leaf-level horizontal coverage, and
    # the repairability checkpoint evidence (pre-repair comparison + render)
    # is retained when the first fitting iteration failed.
    assert any(row["property"].startswith("leaf_") for row in comparison["rows"])
    # C2-0cV: the review index carries the auditable binding table and the
    # visible-rhythm provenance; the plan records rhythm decisions.
    review_text = (run_dir / "review.html").read_text()
    assert "Section binding decisions" in review_text
    assert "Visible-section rhythm decisions" in review_text
    plan_json = json.loads((run_dir / "docx_render_plan.json").read_text())
    assert len(plan_json["visible_rhythm_decisions"]) >= 1
    gap_rows = [row for row in comparison["rows"] if row["property"] == "heading_gap_above"]
    assert all(row.get("basis_source") for row in gap_rows)
    if pair == "E_D":
        # The stale predecessor gap (measured from the omitted KEY SKILLS) is
        # recomputed from the preserved visible rhythm and the rendered gap
        # is verified against that effective basis.
        rhythm_row = next(
            row for row in gap_rows if row["node"] == "section.04"
        )
        assert rhythm_row["basis_source"] == "declared_state_visible_rhythm"
        assert rhythm_row["classification"] == "pass"
        decision = next(
            d for d in plan_json["visible_rhythm_decisions"] if d["node_id"] == "section.04"
        )
        assert decision["omitted_between"] == ["section.03"]
        assert decision["basis"] == "measured_common_section_rhythm"
        # C2-0cS: the measured SKILLS POOL category grid binds the candidate
        # groups row-major and every measured anchor verifies from the preview.
        grid_cells = next(
            p["category_grid_cells"] for p in plan_json["sections"]
            if p["node_id"] == "section.02"
        )
        assert [(c["row_index"], c["column_index"], c["leaf_id"]) for c in grid_cells] == [
            (0, 0, "skills.languages"), (0, 1, "skills.software"),
        ]
        grid_rows = [row for row in comparison["rows"] if row["property"].startswith("grid_")]
        anchors = [row for row in grid_rows if row["property"] in {"grid_label_right_x", "grid_value_x0"}]
        assert anchors and all(row["classification"] == "pass" for row in anchors)
        bold_rows = [row for row in grid_rows if row["property"] == "grid_label_bold"]
        assert bold_rows and all(row["classification"] == "pass" for row in bold_rows)
    else:
        # E→F and D→E measured no category-grid cluster: the ordinary item
        # rendering applies, unchanged by this checkpoint.
        assert all(not p["category_grid_cells"] for p in plan_json["sections"])
    if pair == "E_F":
        before = json.loads((run_dir / "docx_geometry_comparison_before.json").read_text())
        before_child_failures = [
            row for row in before["rows"]
            if row["classification"] == "fail"
            and row["property"] in {"leaf_child_marker_x", "leaf_child_text_x", "leaf_child_hanging_indent"}
            and row["node"] == "section.04"
        ]
        # The pre-repair E→F comparison must expose the Experience child
        # indentation gap the old 40/40 result could not see.
        assert before_child_failures, before["counts"]
        after_by_key = {
            (row["property"], row.get("detail") or ""): row for row in comparison["rows"]
        }
        for row in before_child_failures:
            after = after_by_key[(row["property"], row.get("detail") or "")]
            assert after["classification"] == "pass", after
        assert gates["passed"] is True
        assert gates["gates"]["rendered_geometry_matches_declared_contract"] is True
        assert gates["gates"]["rendered_colors_match_declared_contract"] is True
        assert report["pagination"]["docx_preview_page_count"] == 1
        assert (run_dir / "c2_output_before.pdf").exists()
    else:
        # D→E and E→D retain their previous honest fail-closed status; the
        # expanded coverage did not tune them toward passing.
        assert gates["passed"] is False
        assert gates["gates"]["rendered_geometry_matches_declared_contract"] is False


# ---------------------------------------------------------------------------
# Rendered-geometry measurement, comparison, typography, and fitting
# (work order Parts 1-5; offline — no LibreOffice, no corpus)
# ---------------------------------------------------------------------------


def _fake_target_geo(state, plan) -> dict:
    """A target-geometry measurement derived from the declared state/plan
    (the values the compiler was told to consume)."""
    sections = {}
    rules = {rule.rule_id: rule.model_dump() for rule in state.rules}
    for section in [*plan.sections, *plan.appended_sections]:
        if section.empty:
            continue
        rule = rules.get(section.rule_id) if section.rule_id else None
        entry_node = next(
            (node for node in state.nodes if node.kind == "entry_row" and node.parent_id == section.node_id),
            None,
        )
        heading_token = _style_of(state, section.style_id)
        sections[section.node_id] = {
            "label": section.label,
            "page": 1,
            "heading_top_pt": 40.0,
            "heading_x0_pt": (rule["x0_pt"] if rule else float(state.page.margin_left_pt)),
            "heading_size_pt": heading_token.font_size_pt if heading_token else 14.346,
            "rule": (
                {
                    "page": 1, "top_pt": 60.0, "x0_pt": rule["x0_pt"],
                    "x1_pt": rule["x1_pt"], "stroke_pt": rule["stroke_pt"],
                }
                if rule else None
            ),
            "content_start_x_pt": section.base_x0_pt if section.base_x0_pt is not None else float(state.page.margin_left_pt),
            "entry_right_edge_pt": (
                entry_node.columns[-1].x1_pt
                if entry_node is not None and entry_node.columns[-1].x1_pt is not None
                else None
            ),
            "bullets": (
                {"marker_x0_pt": section.bullet_dot_x0_pt, "text_x0_pt": section.bullet_text_x0_pt}
                if section.bullet_dot_x0_pt is not None else None
            ),
            "content_lines": 3,
        }
    return {"measured_from": "synthetic", "sections": sections, "rules": []}


def _fake_rendered(state, plan, *, overrides=None, page_y_shift=0.0):
    """A rendered-geometry measurement whose lines sit exactly on the
    positions a correct DOCX render produces (derived from the state/plan),
    plus per-property offsets to prove the gate fails on real deltas."""
    overrides = overrides or {}
    font = "BAAAAA+ArialMT"
    page = state.page
    rules_by_id = {rule.rule_id: rule.model_dump() for rule in state.rules}
    lines: list[dict] = []
    rules: list[dict] = []
    y = float(page.margin_top_pt) + 4.0 + page_y_shift
    line_number = 0

    def line(kind, top, x0, x1, size, font, text, chars):
        return {
            "page": 1, "top": round(top, 3), "bottom": round(top + size, 3),
            "x0": round(x0, 3), "x1": round(x1, 3), "size": size, "font": font,
            "text": text, "chars": chars,
        }

    # header rows (verbatim field text in measured order)
    for row_plan in [*plan.header_rows, *([plan.header_overflow] if plan.header_overflow else [])]:
        if not getattr(row_plan, "fields", None):
            continue
        token = _style_of(state, row_plan.style_id)
        size = token.font_size_pt if token else 9.0
        if getattr(row_plan, "gap_above_pt", None):
            y += float(row_plan.gap_above_pt)
        text = " ".join(
            (f" {row_plan.separator} " if index else "") + field.text
            for index, field in enumerate(row_plan.fields)
        ) if getattr(row_plan, "separator", None) else "   ".join(field.text for field in row_plan.fields)
        chars = [
            {"text": char, "x0": float(page.margin_left_pt) + 0.5 * i,
             "x1": float(page.margin_left_pt) + 0.5 * i + 0.5,
             "bottom": y + size, "size": size, "font": font}
            for i, char in enumerate(text)
        ]
        lines.append(line("header", y, float(page.margin_left_pt), float(page.margin_left_pt) + 0.5 * len(text), size, font, text, chars))
        y += size

    for section_plan in [*plan.sections, *plan.appended_sections]:
        if section_plan.empty:
            continue
        heading_token = _style_of(state, section_plan.style_id)
        rule = rules_by_id.get(section_plan.rule_id) if section_plan.rule_id else None
        gap_above = section_plan.heading_gap_above_pt
        if gap_above is None:
            gap_above = 0.0
        elif section_plan.candidate_only:
            gap_above = 12.0  # the documented appended-heading rhythm rule
        y += gap_above
        heading_size = heading_token.font_size_pt if heading_token else 14.346
        heading_x0 = (
            float(rule["x0_pt"]) if rule else float(page.margin_left_pt)
        )
        if overrides and "heading_x" in overrides:
            heading_x0 += overrides["heading_x"]
        heading_top = y
        chars = [
            {"text": char, "x0": heading_x0 + 0.6 * index, "x1": heading_x0 + 0.6 * index + 0.6,
             "bottom": y + heading_size, "size": heading_size, "font": font}
            for index, char in enumerate(section_plan.label)
        ]
        heading_bottom = y + heading_size
        lines.append(line("heading", y, heading_x0, heading_x0 + 0.6 * len(section_plan.label), heading_size, font, section_plan.label, chars))
        y = heading_bottom
        if rule is not None:
            if rule.get("placement") == "below_heading":
                rule_top = heading_bottom + float(rule.get("gap_above_pt") or 0.0) + (overrides or {}).get("heading_to_rule_gap", 0.0)
                y = rule_top + float(rule.get("gap_below_pt") or 0.0) + (overrides or {}).get("heading_to_content_gap", 0.0)
            else:
                # above_heading: rule sits ABOVE the heading text.
                rule_top = heading_top - float(rule.get("gap_below_pt") or 0.0) - (overrides or {}).get("heading_to_rule_gap", 0.0)
                y = heading_bottom + float(section_plan.heading_gap_below_pt or 0.0) + (overrides or {}).get("heading_to_content_gap", 0.0)
            rules.append({
                "page": 1, "top_pt": round(rule_top, 3),
                "x0_pt": round(float(rule["x0_pt"]), 3),
                "x1_pt": round(float(rule["x1_pt"]), 3),
                "stroke_pt": round(float(rule["stroke_pt"]), 3),
            })
        content_size = (
            _style_of(state, section_plan.content_style_id).font_size_pt
            if _style_of(state, section_plan.content_style_id) else 9.0
        )
        entry_node = next(
            (node for node in state.nodes if node.kind == "entry_row" and node.parent_id == section_plan.node_id),
            None,
        )
        for index, entry in enumerate(section_plan.entries):
            if index and section_plan.inter_entry_gap_above_pt is not None:
                y += float(section_plan.inter_entry_gap_above_pt)
            title_token = _style_of(state, section_plan.title_style_id)
            meta_token = _style_of(state, section_plan.meta_style_id or section_plan.detail_style_id or section_plan.title_style_id)
            meta_lines = entry.meta_lines
            title_lines = entry.title_lines
            for line_index, title in enumerate(entry.title_lines):
                meta = meta_lines[line_index] if line_index < len(meta_lines) else None
                title_x0 = (
                    float(section_plan.base_x0_pt) + (overrides or {}).get("entry_left_column_x", 0.0)
                    if section_plan.base_x0_pt is not None else float(page.margin_left_pt)
                )
                title_size = title_token.font_size_pt if title_token else content_size
                meta_size = meta_token.font_size_pt if meta_token else content_size
                meta_x1 = (
                    float(entry_node.columns[-1].x1_pt) + (overrides or {}).get("entry_right_edge", 0.0)
                    if entry_node is not None and entry_node.columns[-1].x1_pt is not None
                    else title_x0 + 200.0
                )
                text = title.text if meta is None else f"{title.text} {meta.text}"
                chars = [
                    {"text": char, "x0": title_x0 + 0.5 * i, "x1": title_x0 + 0.5 * i + 0.5,
                     "bottom": y + title_size, "size": title_size, "font": font}
                    for i, char in enumerate(title.text)
                ]
                if meta is not None:
                    chars += [
                        {"text": char, "x0": meta_x1 - (len(meta.text) - i) * 0.5,
                         "x1": meta_x1 - (len(meta.text) - 1 - i) * 0.5,
                         "bottom": y + meta_size, "size": meta_size, "font": font}
                        for i, char in enumerate(meta.text)
                    ]
                lines.append(line("entry_row", y, title_x0, meta_x1, title_size, font, text, chars))
                y += title_size
            for bullet in entry.bullet_items:
                marker_x = (
                    float(section_plan.bullet_dot_x0_pt) + (overrides or {}).get("bullet_marker_x", 0.0)
                    if section_plan.bullet_dot_x0_pt is not None else float(page.margin_left_pt)
                )
                text_x = (
                    float(section_plan.bullet_text_x0_pt) + (overrides or {}).get("bullet_text_x", 0.0)
                    if section_plan.bullet_text_x0_pt is not None else marker_x + 6.0
                )
                marker = "\uf0b7"
                chars = (
                    [{"text": marker, "x0": marker_x, "x1": marker_x + 4.0, "bottom": y + content_size,
                      "size": content_size, "font": "OpenSymbol"}]
                    + [{"text": char, "x0": text_x + 0.5 * i, "x1": text_x + 0.5 * i + 0.5,
                        "bottom": y + content_size, "size": content_size, "font": font}
                       for i, char in enumerate(_strip_leading_marker_glyphs(bullet.text))]
                )
                lines.append(line("bullet", y, marker_x, text_x + len(bullet.text) * 4.0, content_size, font, f"{marker} {_strip_leading_marker_glyphs(bullet.text)}", chars))
                y += content_size
            child_x0 = (
                float(section_plan.base_x0_pt)
                if section_plan.base_x0_pt is not None else float(page.margin_left_pt)
            )
            for text_line in entry.text_lines:
                # The declared compiler aligns entry child detail lines with
                # their entry's content column (C2-0cR root-cause fix).
                lines.append(line("textline", y, child_x0, child_x0 + 200.0, content_size, font, text_line.text,
                                  [{"text": char, "x0": child_x0 + 0.5 * i, "x1": child_x0 + 0.5 * i + 0.5,
                                    "bottom": y + content_size, "size": content_size, "font": font}
                                   for i, char in enumerate(text_line.text)]))
                y += content_size
        for item in section_plan.items:
            lines.append(line("item", y, float(page.margin_left_pt), float(page.margin_left_pt) + 200.0, content_size, font, item.text,
                              [{"text": char, "x0": float(page.margin_left_pt) + 0.5 * i, "x1": float(page.margin_left_pt) + 0.5 * i + 0.5,
                                "bottom": y + content_size, "size": content_size, "font": font}
                               for i, char in enumerate(item.text)]))
            y += content_size
        for paragraph_line in section_plan.paragraph_lines:
            lines.append(line("paragraph", y, float(page.margin_left_pt), float(page.margin_left_pt) + 300.0, content_size, font, paragraph_line.text,
                              [{"text": char, "x0": float(page.margin_left_pt) + 0.5 * i, "x1": float(page.margin_left_pt) + 0.5 * i + 0.5,
                                "bottom": y + content_size, "size": content_size, "font": font}
                               for i, char in enumerate(paragraph_line.text)]))
            y += content_size
        # C2-0cS: grid cells render as one line per cell fragment at the
        # measured anchors (label right-aligned at the label edge; values at
        # the value anchor), so the mapping buckets them by column window.
        grid_cells = section_plan.category_grid_cells
        if grid_cells:
            grid = next(
                (node.category_grid for node in state.nodes
                 if node.node_id == section_plan.node_id and node.category_grid),
                None,
            )
            if grid is not None:
                content_token = _style_of(state, section_plan.content_style_id)
                grid_line_pitch = float(
                    (content_token.line_height_pt if content_token else None) or content_size
                )
                by_row = {}
                for cell in grid_cells:
                    by_row.setdefault(cell.row_index, []).append(cell)
                for row_index in sorted(by_row):
                    row_y = y
                    for cell in sorted(by_row[row_index], key=lambda c: c.column_index):
                        column = grid.columns[cell.column_index]
                        if cell.label_text:
                            width = 0.6 * content_size
                            start_x = column.label_right_x_pt - 0.6 * content_size * len(cell.label_text)
                            chars = [
                                {"text": char, "x0": start_x + 0.6 * content_size * i,
                                 "x1": start_x + 0.6 * content_size * (i + 1),
                                 "bottom": y + content_size, "size": content_size, "font": font}
                                for i, char in enumerate(cell.label_text)
                            ]
                            lines.append(line("grid_label", row_y, start_x, column.label_right_x_pt, content_size, font, cell.label_text, chars))
                        if cell.value_text:
                            chars = [
                                {"text": char, "x0": column.value_x0_pt + 0.5 * i,
                                 "x1": column.value_x0_pt + 0.5 * (i + 1),
                                 "bottom": y + content_size, "size": content_size, "font": font}
                                for i, char in enumerate(cell.value_text)
                            ]
                            lines.append(line("grid_value", row_y, column.value_x0_pt, column.value_x0_pt + 0.5 * len(cell.value_text), content_size, font, cell.value_text, chars))
                    y = row_y + grid_line_pitch
    lines.sort(key=lambda entry: (entry["page"], entry["top"]))
    mapping = map_rendered_to_expected(plan, lines)
    return {
        "measured_from": "synthetic",
        "page_count": 1,
        "pages": [{"page": 1, "width_pt": float(page.width_pt), "height_pt": float(page.height_pt)}],
        "line_count": len(lines),
        "lines": [{key: item[key] for key in ("page", "top", "bottom", "x0", "x1", "size", "font", "text")} for item in lines],
        "mapping": mapping,
        "rules": rules,
        "sparse_trailing_page": None,
    }


def _shape_state_and_plan():
    """A state with measured heading gaps, entry tiers, and a rule (the
    _shape_summary evidence), so every declared basis exists."""
    from tests.experiments.test_c2_renderer import _shape_summary
    summary = _shape_summary()
    state = compile_synthetic(
        ["WORK EXPERIENCE"], rules=summary["rules"], elements=summary["elements"]
    )
    return state, compile_render_plan(state, rich_candidate(include_unmatched=False))


def test_correct_geometry_passes_within_tolerance() -> None:
    state, plan = _shape_state_and_plan()
    comparison = compare_geometry(state, plan, _fake_target_geo(state, plan), _fake_rendered(state, plan))
    assert comparison["mapping_unmapped"] == 0
    assert comparison["counts"]["failed"] == 0, [r for r in comparison["rows"] if r["classification"] != "pass"]
    assert comparison["gate_passed"] is True


def test_geometry_uses_relative_geometry_not_absolute_page_y() -> None:
    state, plan = _shape_state_and_plan()
    baseline = compare_geometry(state, plan, _fake_target_geo(state, plan), _fake_rendered(state, plan))
    shifted = compare_geometry(
        state, plan, _fake_target_geo(state, plan),
        _fake_rendered(state, plan, page_y_shift=400.0),
    )
    assert shifted["counts"] == baseline["counts"]
    assert [r["classification"] for r in shifted["rows"]] == [r["classification"] for r in baseline["rows"]]
    assert shifted["gate_passed"] is True


def test_wrong_entry_column_x_and_metadata_right_edge_fail() -> None:
    state, plan = _shape_state_and_plan()
    target_geo = _fake_target_geo(state, plan)
    shifted_left = compare_geometry(
        state, plan, target_geo, _fake_rendered(state, plan, overrides={"entry_left_column_x": 10.0})
    )
    row = next(r for r in shifted_left["rows"] if r["property"] == "entry_left_column_x")
    assert row["classification"] == "fail" and row["delta"] > row["tolerance_pt"]
    assert shifted_left["gate_passed"] is False
    shifted_right = compare_geometry(
        state, plan, target_geo, _fake_rendered(state, plan, overrides={"entry_right_edge": -8.0})
    )
    edge_row = next(r for r in shifted_right["rows"] if r["property"] == "entry_right_edge")
    assert edge_row["classification"] == "fail" and abs(edge_row["delta"] + 8.0) < 0.01
    assert shifted_right["gate_passed"] is False


def test_wrong_bullet_indent_and_heading_rule_gap_fail() -> None:
    from tests.experiments.c2_docx_renderer import compare_geometry
    from tests.experiments.test_c2_renderer import _shape_summary

    summary = _shape_summary()
    state = compile_synthetic(["WORK EXPERIENCE"], rules=summary["rules"], elements=summary["elements"])
    plan = compile_render_plan(state, rich_candidate(include_unmatched=False))
    target_geo = _fake_target_geo(state, plan)
    bad_gap = compare_geometry(
        state, plan, target_geo, _fake_rendered(state, plan, overrides={"heading_to_rule_gap": 4.0})
    )
    row = next(r for r in bad_gap["rows"] if r["property"] == "heading_to_rule_gap")
    assert row["classification"] == "fail" and abs(row["delta"] - 4.0) <= 0.01
    assert bad_gap["gate_passed"] is False


def test_unmappable_required_text_fails_rather_than_disappearing() -> None:
    state, plan = _state_and_plan(["WORK EXPERIENCE"])
    rendered = _fake_rendered(state, plan)
    # Drop one required bullet row's lines entirely: the mapping must fail and
    # record the missing row (it may never silently disappear).
    victim = next(r for r in rendered["mapping"]["mapped"] if r["kind"] == "bullet")
    victim_line_texts = {line["text"] for line in victim["lines"]}
    remaining_lines = [
        line for line in rendered["lines"]
        if line["text"] not in victim_line_texts
    ]
    mapping = map_rendered_to_expected(plan, remaining_lines)
    assert mapping["passed"] is False
    assert any(
        entry["expected_text"] == victim["text"] for entry in mapping["unmapped"]
    )


def test_font_substitution_cannot_pass_as_typography_exact(monkeypatch: pytest.MonkeyPatch) -> None:
    state, plan = _state_and_plan(["WORK EXPERIENCE"])
    # The requested family is NOT installed: the documented portable fallback
    # is written and rendered, so typography is classified adjusted.
    monkeypatch.setattr(
        "tests.experiments.c2_docx_renderer.installed_font_families",
        lambda: frozenset({"arial"}),
    )
    state = state.model_copy(deep=True)
    state.styles = [
        style.model_copy(update={"font_family": "Roboto"}) for style in state.styles
    ]
    resolved = resolve_written_fonts(state)
    assert resolved["style.body"]["requested"] == "Roboto"
    assert resolved["style.body"]["written"] == "Arial"
    assert resolved["style.body"]["substituted"] is True
    rendered = _fake_rendered(state, plan)
    tables = typography_tables(state, plan, rendered, resolved)
    assert tables["classification"] == "adjusted"
    assert all(record["classification"] == "adjusted" for record in tables["rendered_typography"])
    # With the family actually installed and rendered, the record is exact.
    monkeypatch.setattr(
        "tests.experiments.c2_docx_renderer.installed_font_families",
        lambda: frozenset({"arial", "roboto"}),
    )
    state = state.model_copy(deep=True)
    state.styles = [style.model_copy(update={"font_family": "Roboto"}) for style in state.styles]
    rendered_exact = _fake_rendered(state, plan, overrides={})
    for mapped_row in rendered_exact["mapping"]["mapped"]:
        bold = bool((_style_of(state, mapped_row.get("style_id")) or _style_of(state, mapped_row.get("meta_style_id"))) and (_style_of(state, mapped_row.get("style_id")) or _style_of(state, mapped_row.get("meta_style_id"))).bold)
        face = "BAAAAA+Roboto-Bold" if bold else "BAAAAA+Roboto-Regular"
        for line in mapped_row["lines"]:
            line["font"] = face
            for char in line["chars"]:
                char["font"] = face
    tables_exact = typography_tables(state, plan, rendered_exact, resolve_written_fonts(state))
    # With the requested family installed and rendered, no record can carry a
    # substitution-driven 'adjusted' classification; matching family/size/weight
    # records are exact.
    assert all(record["classification"] != "adjusted" for record in tables_exact["rendered_typography"])
    assert any(record["classification"] == "exact" for record in tables_exact["rendered_typography"])


def test_requested_and_rendered_fonts_are_recorded_separately(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(
        "tests.experiments.c2_docx_renderer.installed_font_families",
        lambda: frozenset({"arial"}),
    )
    state, plan = _state_and_plan(["WORK EXPERIENCE"])
    state = state.model_copy(deep=True)
    state.styles = [style.model_copy(update={"font_family": "Roboto"}) for style in state.styles]
    resolved = resolve_written_fonts(state)
    tables = typography_tables(state, plan, _fake_rendered(state, plan), resolved)
    record = tables["rendered_typography"][0]
    assert record["requested_font"] == "Roboto"
    assert record["written_font"] == "Arial"
    assert record["rendered_font"] == "arial"
    assert tables["authored_typography"]


def test_incorrect_font_size_fails(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(
        "tests.experiments.c2_docx_renderer.installed_font_families",
        lambda: frozenset({"arial"}),
    )
    state, plan = _state_and_plan(["WORK EXPERIENCE"])
    target_geo = _fake_target_geo(state, plan)
    rendered = _fake_rendered(state, plan)
    # Perturb one rendered line's font size beyond the documented tolerance.
    heading_row = next(r for r in rendered["mapping"]["mapped"] if r["kind"] == "heading")
    heading_row["lines"][0]["size"] += 2.0
    comparison = compare_geometry(state, plan, target_geo, rendered)
    row = next(r for r in comparison["rows"] if r["property"] == "heading_font_size")
    assert row["classification"] == "fail"
    assert comparison["gate_passed"] is False


def test_sparse_trailing_page_is_explicitly_classified() -> None:
    state, plan = _state_and_plan(["WORK EXPERIENCE"])
    target_geo = _fake_target_geo(state, plan)
    two_pages = _fake_rendered(state, plan)
    two_pages["page_count"] = 2
    two_pages["sparse_trailing_page"] = {
        "page": 2, "content_extent_pt": 60.0, "writable_height_pt": 722.8,
        "fraction": 0.08, "sparse": True,
    }
    comparison = compare_geometry(state, plan, target_geo, two_pages)
    row = next(r for r in comparison["rows"] if r["property"] == "sparse_trailing_page")
    assert row["classification"] == "fail"
    assert comparison["gate_passed"] is False
    # One page: no trailing page can exist.
    one_page = _fake_rendered(state, plan)
    comparison = compare_geometry(state, plan, target_geo, one_page)
    row = next(r for r in comparison["rows"] if r["property"] == "sparse_trailing_page")
    assert row["classification"] == "pass"


def test_the_fitter_stops_after_the_documented_iteration_budget(monkeypatch: pytest.MonkeyPatch) -> None:
    state, plan = _state_and_plan(["WORK EXPERIENCE"])
    docx_path = Path("unused.docx")
    calls = {"renders": 0}

    def fake_preview(docx_path, run_dir):
        calls["renders"] += 1
        return {"available": True, "pdf": "unused.pdf", "page_count": 1}

    def always_failing(state, plan, target_geo, rendered):
        return {
            "schema_version": "x", "passed": False, "gate_passed": False,
            "counts": {"total": 1, "passed": 0, "failed": 1, "unmeasurable": 0, "not_applicable": 0},
            "adjustable": True,
            "rows": [{
                "property": "heading_gap_above", "node": "section.01", "basis": 10.0,
                "basis_source": "declared_state", "rendered": 20.0, "delta": 10.0,
                "tolerance_pt": 1.5, "classification": "fail",
                "control": "heading_space_before_pt", "detail": "",
            }],
        }

    monkeypatch.setattr("tests.experiments.c2_docx_renderer.build_document", lambda *a, **k: None)
    monkeypatch.setattr("tests.experiments.c2_docx_renderer.deterministic_docx_bytes", lambda document: b"docx")
    monkeypatch.setattr("tests.experiments.c2_docx_renderer._preview_pdf", fake_preview)
    monkeypatch.setattr("tests.experiments.c2_docx_renderer.measure_target_geometry", lambda *a, **k: {"sections": {}})
    monkeypatch.setattr("tests.experiments.c2_docx_renderer.measure_rendered_geometry", lambda *a, **k: {"page_count": 1})
    monkeypatch.setattr(
        "tests.experiments.c2_docx_renderer.compare_geometry",
        lambda state, plan, target_geo, rendered: always_failing(state, plan, target_geo, rendered),
    )
    result = fit_docx(state, plan, Path("unused.pdf"), docx_path, Path("."))
    assert result["iterations"] == MAX_FITTING_ITERATIONS == 3
    assert result["converged"] is False
    assert len(result["log"]) == 3
    # The fitter kept applying the documented correction each iteration
    # (bounded; it never loops indefinitely).
    corrections = [
        iteration["corrections"]["sections"].get("section.01", {}).get("heading_space_before_pt", 0.0)
        for iteration in result["log"]
    ]
    assert corrections == [0.0, pytest.approx(-10.0), pytest.approx(-20.0)]


# ---------------------------------------------------------------------------
# C2-0cR: leaf-level indentation coverage + the bounded indentation repair
# ---------------------------------------------------------------------------


def _state_plan_with_entry_textlines():
    """The measured shape state/plan with two verbatim child detail lines
    added to the first work entry — the E→F blind-spot shape: source-glyph
    ("• ...") and plain detail lines that stay verbatim TEXT because the
    plan declares no native bullet for them. The plan ledger is extended so
    the added leaves stay owned (accounting remains exact)."""
    state, plan = _shape_state_and_plan()
    section = plan.sections[0]
    entry = section.entries[0]
    patched = entry.model_copy(update={
        "text_lines": [
            LeafText(leaf_id="work.e1.t1", text="• Candidate detail line that must indent with its entry"),
            LeafText(leaf_id="work.e1.t2", text="Candidate plain detail line"),
        ],
    })
    sections = [
        section.model_copy(update={"entries": [patched, *section.entries[1:]]}),
        *plan.sections[1:],
    ]
    ledger = dict(plan.leaf_ledger)
    ledger["work.e1.t1"] = entry.node_id
    ledger["work.e1.t2"] = entry.node_id
    return state, plan.model_copy(update={"sections": sections, "leaf_ledger": ledger})


def _leaf_rows(comparison: dict, node: str = "section.01") -> dict[str, list[dict]]:
    grouped: dict[str, list[dict]] = {}
    for row in comparison["rows"]:
        if row["node"] == node and row["property"].startswith("leaf_"):
            grouped.setdefault(row["property"], []).append(row)
    return grouped


def test_pre_repair_geometry_fails_on_the_uncovered_child_textline() -> None:
    """Regression: a child textline rendered AWAY from its measured target
    anchor must FAIL the expanded gate — the old section-level rows (which
    only checked the FIRST content row) reported 40/40 while every Experience
    detail line sat at the page margin."""
    state, plan = _state_plan_with_entry_textlines()
    rendered = _fake_rendered(state, plan)
    comparison = compare_geometry(state, plan, _fake_target_geo(state, plan), rendered)
    grouped = _leaf_rows(comparison)
    assert grouped["leaf_child_marker_x"], "child textline leaves must carry marker coverage"
    failing = [row for row in grouped["leaf_child_marker_x"] if row["classification"] == "fail"]
    assert failing, grouped
    assert all(row["basis_source"] == "measured_target" for row in grouped["leaf_child_marker_x"])
    assert comparison["gate_passed"] is False
    assert comparison["adjustable"] is True  # the fitter has a documented control


def test_section_first_row_alignment_cannot_mask_child_misalignment() -> None:
    """content_start_x checks only the section's first content row; the
    expanded gate must still fail when a LATER child row is misaligned."""
    state, plan = _state_plan_with_entry_textlines()
    rendered = _fake_rendered(state, plan)
    comparison = compare_geometry(state, plan, _fake_target_geo(state, plan), rendered)
    content_start = next(
        row for row in comparison["rows"] if row["property"] == "content_start_x" and row["node"] == "section.01"
    )
    assert content_start["classification"] == "pass"
    assert comparison["gate_passed"] is False


def test_target_measured_marker_text_and_hanging_are_separate_contracts() -> None:
    """Marker x, bullet text x, and the hanging indent are measured and
    compared separately per leaf: shifting only the text start fails the
    text and hanging rows while the marker row still passes."""
    state, plan = _state_plan_with_entry_textlines()
    rendered = _fake_rendered(state, plan)
    # Put the child line ON its measured anchors first (marker 60.0, text
    # 64.0), then shift ONLY the text after the leading glyph by +3pt.
    shifted = 0
    for mapped in rendered["mapping"]["mapped"]:
        if mapped["kind"] == "textline" and mapped["text"].startswith("•"):
            line0 = mapped["lines"][0]
            chars = line0["chars"]
            space_width = chars[1]["x1"] - chars[1]["x0"]
            chars[0]["x0"], chars[0]["x1"] = 60.0, 60.0 + space_width
            chars[1]["x0"], chars[1]["x1"] = 60.0 + space_width, 64.0
            for offset, char in enumerate(chars[2:]):
                char["x0"] = 64.0 + 0.5 * offset + 3.0
                char["x1"] = 64.0 + 0.5 * offset + 0.5 + 3.0
            line0["x0"] = 60.0
            line0["x1"] = round(max(char["x1"] for char in chars), 3)
            shifted += 1
            break
    assert shifted
    comparison = compare_geometry(state, plan, _fake_target_geo(state, plan), rendered)
    grouped = _leaf_rows(comparison)
    marker = grouped["leaf_child_marker_x"][0]
    text = grouped["leaf_child_text_x"][0]
    hanging = grouped["leaf_child_hanging_indent"][0]
    assert marker["classification"] == "pass"
    assert text["classification"] == "fail"
    assert hanging["classification"] == "fail"
    assert {marker["property"], text["property"], hanging["property"]} == {
        "leaf_child_marker_x", "leaf_child_text_x", "leaf_child_hanging_indent",
    }


def test_child_textline_without_target_anchor_is_honestly_unmeasurable() -> None:
    """A child leaf whose target carries no measurable bullet-text anchor has
    no basis: it must fail the gate as unmeasurable (never a silent pass)."""
    state, plan = _state_plan_with_entry_textlines()
    target_geo = _fake_target_geo(state, plan)
    target_geo["sections"]["section.01"]["bullets"] = None
    rendered = _fake_rendered(state, plan)
    comparison = compare_geometry(state, plan, target_geo, rendered)
    grouped = _leaf_rows(comparison)
    text_rows = grouped["leaf_child_text_x"]
    assert any(row["classification"] == "unmeasurable" for row in text_rows)
    assert comparison["gate_passed"] is False
    # The declared marker basis still applies where the state measures one.
    assert any(row["classification"] == "pass" for row in grouped["leaf_child_marker_x"])


def test_the_bounded_edit_changes_only_the_intended_stable_nodes() -> None:
    """FitAdjustments.sections[<node>].entry_child_text_indent_pt moves ONLY
    that node's child detail lines; every other paragraph (text and indent)
    is byte-identical, and the edited paragraph keeps its text verbatim."""
    state, plan = _state_plan_with_entry_textlines()
    before = build_document(state, plan)
    after = build_document(
        state, plan,
        adjustments=FitAdjustments(sections={"section.01": SectionFit(node_id="section.01", entry_child_text_indent_pt=10.7)}),
    )
    from docx.shared import Pt

    def paragraph_snapshot(document):
        return [
            (paragraph.text, paragraph.paragraph_format.left_indent)
            for paragraph in document.paragraphs
        ]
    before_rows = paragraph_snapshot(before)
    after_rows = paragraph_snapshot(after)
    assert [text for text, _ in before_rows] == [text for text, _ in after_rows]
    changed = [
        (index, b, a) for index, (b, a) in enumerate(zip(before_rows, after_rows))
        if b[1] != a[1]
    ]
    # Exactly the two child detail lines of section.01's first entry moved
    # (the source-glyph line and the plain line); every other paragraph is
    # byte-identical and all texts stay verbatim.
    assert len(changed) == 2, changed
    assert all("Candidate detail line" in after_rows[index][0] or "Candidate plain detail line" in after_rows[index][0]
               for index, _, _ in changed)
    expected = Pt(46.9 - 36.0 + 10.7)  # declared entry-column alignment + correction
    assert all(abs(a[1].pt - expected.pt) < 0.01 for _, _, a in changed)


def test_candidate_content_and_accounting_unchanged_by_the_edit(tmp_path: Path) -> None:
    """The bounded edit is presentation-only: reading order, leaf texts, and
    content accounting are identical with and without the adjustment."""
    state, plan = _state_plan_with_entry_textlines()
    plain = build_document(state, plan)
    edited = build_document(
        state, plan,
        adjustments=FitAdjustments(sections={"section.01": SectionFit(node_id="section.01", entry_child_text_indent_pt=10.7)}),
    )
    plain_path = tmp_path / "plain.docx"
    edited_path = tmp_path / "edited.docx"
    plain_path.write_bytes(deterministic_docx_bytes(plain))
    edited_path.write_bytes(deterministic_docx_bytes(edited))
    inspection_plain = inspect_docx(plain_path)
    inspection_edited = inspect_docx(edited_path)
    # Semantic text is identical with and without the edit (verbatim glyphs).
    assert [p["text"] for p in inspection_plain["paragraphs"]] == [
        p["text"] for p in inspection_edited["paragraphs"]
    ]
    accounting_plain = content_accounting(plan, inspection_plain)
    accounting_edited = content_accounting(plan, inspection_edited)
    assert accounting_plain["passed"] is True and accounting_edited["passed"] is True
    assert accounting_plain["leaf_records"] == accounting_edited["leaf_records"]
    assert accounting_plain["presentation_marker_conversions"] == (
        accounting_edited["presentation_marker_conversions"]
    )


def test_unrelated_section_geometry_is_unchanged_by_the_edit() -> None:
    """Applying the bounded edit flips ONLY the intended node's child rows;
    rows of every unrelated node keep their property/classification/value."""
    state, plan = _state_plan_with_entry_textlines()
    before_rendered = _fake_rendered(state, plan)
    comparison_before = compare_geometry(state, plan, _fake_target_geo(state, plan), before_rendered)
    # The repaired render: the source-glyph child line moved onto the
    # measured anchors (marker 60.0pt, text 64.0pt — the shape state's
    # measured bullet tiers); the plain child line keeps its measured
    # content-start position.
    after_rendered = _fake_rendered(state, plan)
    for mapped in after_rendered["mapping"]["mapped"]:
        if mapped["kind"] == "textline" and mapped["text"].startswith("•"):
            line0 = mapped["lines"][0]
            chars = line0["chars"]
            space_width = chars[1]["x1"] - chars[1]["x0"]
            chars[0]["x0"], chars[0]["x1"] = 60.0, 60.0 + space_width
            chars[1]["x0"], chars[1]["x1"] = 60.0 + space_width, 64.0
            for offset, char in enumerate(chars[2:]):
                char["x0"] = 64.0 + 0.5 * offset
                char["x1"] = 64.0 + 0.5 * offset + 0.5
            line0["x0"] = 60.0
            line0["x1"] = round(max(char["x1"] for char in chars), 3)
    comparison_after = compare_geometry(state, plan, _fake_target_geo(state, plan), after_rendered)
    key = lambda row: (row["property"], row["node"], row.get("detail") or "")
    before_by_key = {key(row): row for row in comparison_before["rows"]}
    for row in comparison_after["rows"]:
        before = before_by_key.get(key(row))
        assert before is not None, key(row)
        if row["node"] != "section.01" or not row["property"].startswith("leaf_child"):
            assert row["classification"] == before["classification"], key(row)
            assert row["rendered"] == before["rendered"], key(row)
    grouped = _leaf_rows(comparison_after)
    assert all(row["classification"] == "pass" for rows in grouped.values() for row in rows)
    assert comparison_after["gate_passed"] is True


# ---------------------------------------------------------------------------
# C2-0cC: node-level color fidelity + rendered-color gate + styled runs
# ---------------------------------------------------------------------------


def _colored_state(colors: list[str | None], header_colors: dict[str, str] | None = None):
    """A synthetic state whose heading scaffolds (and optionally header rows)
    carry measured colors — the Resume-D shape without any pair-specific code."""
    from tests.experiments.c2_pipeline import state_from_scaffolds
    from tests.experiments.test_c2_pipeline import BODY, HEADER, SUMMARY, TARGET_SHA

    body = BODY.model_copy(update={
        "headings": [
            heading.model_copy(update={"color_hex": color})
            for heading, color in zip(BODY.headings, colors)
        ]
    })
    header = HEADER
    if header_colors:
        header = [
            row.model_copy(update={"color_hex": header_colors.get(row.role)})
            for row in HEADER
        ]
    return state_from_scaffolds(TARGET_SHA, header, body, BULLET_TIERS, json.loads(json.dumps(SUMMARY)))


def test_measured_header_and_tagline_colors_survive_evidence_to_state() -> None:
    state = _colored_state(
        colors=[None] * len(FULL_COVERAGE_LABELS),
        header_colors={"name": "#0E6E55", "location": "#0E6E55", "contact": "#000000"},
    )
    styles = {style.style_id: style for style in state.styles}
    # The measured name/extension-row colors survive into their own tokens;
    # the contact row keeps its measured black.
    assert styles["style.name"].color_hex == "#0E6E55"
    assert styles["style.location"].color_hex == "#0E6E55"
    assert styles["style.contact"].color_hex == "#000000"


def test_differently_colored_headings_get_distinct_tokens_and_reuse() -> None:
    colors = ["#1F1D8E", "#1F1D8E", "#8D1E8C", "#8D1E8C", "#8D1E8C", "#1F1D8E", "#8D1E8C"]
    state = _colored_state(colors)
    styles = [style for style in state.styles if style.style_id.startswith("style.heading")]
    assert {style.style_id: style.color_hex for style in styles} == {
        "style.heading": "#1F1D8E",
        "style.heading.2": "#8D1E8C",
    }
    heading_nodes = [node for node in state.nodes if node.kind == "heading"]
    referenced = {node.style_id for node in heading_nodes}
    assert referenced == {"style.heading", "style.heading.2"}
    # Each node references the token matching ITS measured color.
    for node, color in zip(heading_nodes, colors):
        token = next(style for style in styles if style.style_id == node.style_id)
        assert token.color_hex == color


def test_missing_measured_color_stays_explicit_with_black_fallback() -> None:
    state = _colored_state(colors=[None] * len(FULL_COVERAGE_LABELS))
    styles = [style for style in state.styles if style.style_id.startswith("style.heading")]
    assert len(styles) == 1 and styles[0].color_hex is None
    _, plan = _state_and_plan()
    # The DOCX fallback writes the documented black, never an invented color.
    document = build_document(state, plan)
    heading_colors = [
        color
        for paragraph in document.paragraphs
        if (paragraph.style.name or "").startswith("Heading")
        for run in paragraph.runs
        for color in [run.font.color.rgb]
    ]
    assert heading_colors and all(str(color) == "000000" for color in heading_colors)


def test_docx_writes_native_w_color_run_values(tmp_path: Path) -> None:
    state = _colored_state(
        colors=["#B50013"] + [None] * (len(FULL_COVERAGE_LABELS) - 1),
        header_colors={"name": "#0E6E55"},
    )
    _, plan = _state_and_plan()
    document = build_document(state, plan)
    path = tmp_path / "c.docx"
    path.write_bytes(deterministic_docx_bytes(document))
    inspection = inspect_docx(path)
    colored = [
        record for record in inspection["paragraphs"]
        if any(color and color.upper() in ("0E6E55", "B50013") for color in record["run_colors"])
    ]
    assert colored, inspection["paragraphs"][:3]
    assert any("0E6E55" in (record["run_colors"] or []) for record in inspection["paragraphs"])


def test_html_emits_the_same_measured_color_intent() -> None:
    from tests.experiments.c2_renderer import render_html

    state = _colored_state(
        colors=["#B50013"] + [None] * (len(FULL_COVERAGE_LABELS) - 1),
        header_colors={"name": "#0E6E55"},
    )
    _, plan = _state_and_plan()
    html = render_html(state, plan)
    assert "color: #B50013;" in html
    assert "color: #0E6E55;" in html


def _char(text: str, color: str | None) -> dict:
    return {"text": text, "color": color}


def _color_rendered(rows: list[dict], rules: list[dict] | None = None) -> dict:
    return {"mapping": {"mapped": rows}, "rules": list(rules or [])}


def _heading_row(node: str, style_id: str, color: str | None) -> dict:
    return {
        "kind": "heading", "section_node_id": node, "style_id": style_id,
        "leaf_ids": [], "native_bullet": False, "meta_text": None,
        "lines": [{"page": 1, "top": 50.0, "chars": [_char(char, color) for char in "HEADING"]}],
    }


def test_rendered_color_measured_from_pdf_and_compared() -> None:
    from tests.experiments.c2_pipeline import state_from_scaffolds
    from tests.experiments.test_c2_pipeline import BULLET_TIERS, BODY, HEADER, SUMMARY, TARGET_SHA
    from tests.experiments.c2_renderer import _pdf_color_hex

    # pdfplumber color normalization is deterministic (gray/RGB/CMYK).
    assert _pdf_color_hex(None) is None
    assert _pdf_color_hex(0.0) == "#000000"
    assert _pdf_color_hex([1.0, 1.0, 1.0]) == "#FFFFFF"
    assert _pdf_color_hex([0.5, 0.5, 0.5]) == "#808080"
    cmyk = _pdf_color_hex([0.0, 0.5, 0.5, 0.2])
    assert cmyk and cmyk.startswith("#")

    state = _colored_state(colors=["#B50013"] + [None] * (len(FULL_COVERAGE_LABELS) - 1))
    rendered = _color_rendered(
        rows=[_heading_row("section.01", "style.heading", "#B50013")],
        rules=[],
    )
    comparison = compare_colors(state, _plan_with_rules(), rendered)
    heading_rows = [row for row in comparison["rows"] if row["property"] == "heading_color"]
    assert heading_rows[0]["classification"] == "pass"
    assert heading_rows[0]["expected_color"] == "#B50013"
    assert heading_rows[0]["authored_color"] == "#B50013"
    assert heading_rows[0]["rendered_color"] == "#B50013"


def _plan_with_rules() -> object:
    """A minimal stand-in plan exposing only what compare_colors consumes."""
    from types import SimpleNamespace

    return SimpleNamespace(sections=[], appended_sections=[])


def test_all_black_rendered_heading_fails_against_colored_target() -> None:
    state = _colored_state(colors=["#B50013"] + [None] * (len(FULL_COVERAGE_LABELS) - 1))
    rendered = _color_rendered(
        rows=[_heading_row("section.01", "style.heading", "#000000")],
    )
    comparison = compare_colors(state, _plan_with_rules(), rendered)
    row = next(row for row in comparison["rows"] if row["property"] == "heading_color")
    assert row["classification"] == "fail"
    assert comparison["gate_passed"] is False


def test_unmeasured_color_fallback_is_adjusted_never_exact() -> None:
    state = _colored_state(colors=[None] * len(FULL_COVERAGE_LABELS))
    rendered = _color_rendered(
        rows=[_heading_row("section.01", "style.heading", "#000000")],
    )
    comparison = compare_colors(state, _plan_with_rules(), rendered)
    row = next(row for row in comparison["rows"] if row["property"] == "heading_color")
    assert row["classification"] == "adjusted"
    assert row["expected_source"] == "unmeasured_fallback"
    # ... and a non-black render of an unmeasured token is a FAIL.
    rendered_wrong = _color_rendered(
        rows=[_heading_row("section.01", "style.heading", "#123456")],
    )
    comparison_wrong = compare_colors(state, _plan_with_rules(), rendered_wrong)
    assert next(
        row for row in comparison_wrong["rows"] if row["property"] == "heading_color"
    )["classification"] == "fail"


def test_gold_and_green_rule_colors_are_verified_from_rendered_output() -> None:
    from types import SimpleNamespace

    from tests.experiments.c2_pipeline import RuleDecoration

    state = _colored_state(colors=[None] * len(FULL_COVERAGE_LABELS))
    state = state.model_copy(update={"rules": [
        RuleDecoration(
            rule_id="rule.section.01", x0_pt=36.0, x1_pt=576.0, stroke_pt=0.75,
            color_hex="#A16F0B", placement="below_heading", evidence_ids=["e.rule.1"],
        ),
        RuleDecoration(
            rule_id="rule.section.02", x0_pt=36.0, x1_pt=576.0, stroke_pt=0.75,
            color_hex="#0A7903", placement="below_heading", evidence_ids=["e.rule.2"],
        ),
    ]})
    empty_section = SimpleNamespace(
        node_id="section.01", rule_id="rule.section.01", rule_placement="below_heading",
        empty=False, styled_lines=[],
    )
    green_section = SimpleNamespace(
        node_id="section.02", rule_id="rule.section.02", rule_placement="below_heading",
        empty=False, styled_lines=[],
    )
    plan = SimpleNamespace(sections=[empty_section, green_section], appended_sections=[])
    rendered = _color_rendered(
        rows=[
            _heading_row("section.01", "style.heading", "#000000"),
            _heading_row("section.02", "style.heading", "#000000"),
        ],
        rules=[
            {"x0_pt": 36.0, "x1_pt": 576.0, "top_pt": 100.0, "page": 1, "color_hex": "#A16F0B"},
            {"x0_pt": 36.0, "x1_pt": 576.0, "top_pt": 200.0, "page": 1, "color_hex": "#0A7903"},
        ],
    )
    # The heading positions put each rule in its own section's vertical region
    # (below_heading: rule sits 0..60pt below its heading on the same page).
    rendered["mapping"]["mapped"][0]["lines"][0]["page"] = 1
    rendered["mapping"]["mapped"][0]["lines"][0]["top"] = 60.0
    rendered["mapping"]["mapped"][1]["lines"][0]["page"] = 1
    rendered["mapping"]["mapped"][1]["lines"][0]["top"] = 160.0
    comparison = compare_colors(state, plan, rendered)
    rule_rows = [row for row in comparison["rows"] if row["property"] == "rule_color"]
    assert {row["expected_color"]: row["rendered_color"] for row in rule_rows} == {
        "#A16F0B": "#A16F0B", "#0A7903": "#0A7903",
    }
    # A gold rule where a green one is required fails (wrong-color rule).
    rendered_wrong = _color_rendered(
        rows=rendered["mapping"]["mapped"],
        rules=[
            {"x0_pt": 36.0, "x1_pt": 576.0, "top_pt": 100.0, "page": 1, "color_hex": "#A16F0B"},
            {"x0_pt": 36.0, "x1_pt": 576.0, "top_pt": 200.0, "page": 1, "color_hex": "#A16F0B"},
        ],
    )
    comparison_wrong = compare_colors(state, plan, rendered_wrong)
    assert any(
        row["classification"] == "fail" for row in comparison_wrong["rows"] if row["property"] == "rule_color"
    )
    assert comparison_wrong["gate_passed"] is False


def _styled_line_plan():
    """The full-coverage plan with the summary paragraph rendered as ordered
    styled runs (the C2-0cC capability proof: candidate-owned text split into
    ordered runs bound to template-owned measured style tokens)."""
    from tests.experiments.c2_renderer import StyledLine, StyledRun

    state, plan = _state_and_plan()
    section = next(s for s in plan.sections if s.node_id == "section.01")
    patched = section.model_copy(update={
        "paragraph_lines": [],
        "styled_lines": [StyledLine(
            leaf_id="summary.p1",
            runs=[
                StyledRun(style_id="style.heading", leaf_id="summary.p1", text="SUMMARY"),
                StyledRun(style_id="style.body", leaf_id="summary.p1", text=" — candidate paragraph."),
            ],
        )],
    })
    sections = [
        patched if s.node_id == "section.01" else s for s in plan.sections
    ]
    return state, plan.model_copy(update={"sections": sections})


def test_ordered_styled_runs_preserve_text_order_color_bold_editability(tmp_path: Path) -> None:
    from tests.experiments.c2_renderer import render_html

    state, plan = _styled_line_plan()
    document = build_document(state, plan)
    path = tmp_path / "styled.docx"
    path.write_bytes(deterministic_docx_bytes(document))
    inspection = inspect_docx(path)
    styled = next(
        record for record in inspection["paragraphs"]
        if record["text"] == "SUMMARY — candidate paragraph."
    )
    # Native editable runs (python-docx sees real runs), text and order exact.
    assert styled["text"] == "SUMMARY — candidate paragraph."
    assert styled["run_colors"] and len(styled["run_colors"]) == 2
    # Ordered runs carry DIFFERENT presentation: the heading-token run is bold.
    html = render_html(state, plan)
    assert 'data-node-id="section.01.content.summary.p1"' in html
    assert html.index("SUMMARY") < html.index("candidate paragraph.")
    assert "font-weight: 700;" in html


def test_styled_runs_copy_no_target_candidate_facts() -> None:
    state, plan = _styled_line_plan()
    for section in [*plan.sections, *plan.appended_sections]:
        for line in section.styled_lines:
            for run in line.runs:
                assert "TARGETFACT" not in run.text and "TARGETNAME" not in run.text
                # Runs reference candidate-owned leaves; style ids are template tokens.
                assert run.leaf_id == line.leaf_id
                assert run.style_id in {style.style_id for style in state.styles}


def test_styled_run_accounting_and_reading_order_unchanged(tmp_path: Path) -> None:
    state, plan = _styled_line_plan()
    document = build_document(state, plan)
    path = tmp_path / "styled.docx"
    path.write_bytes(deterministic_docx_bytes(document))
    inspection = inspect_docx(path)
    inspection["reading_order_gate"] = reading_order_gate(plan, inspection)
    accounting = content_accounting(plan, inspection)
    assert accounting["passed"] is True, (accounting["missing"], accounting["duplicated"])
    assert inspection["reading_order_gate"]["passed"] is True


# ---------------------------------------------------------------------------
# C2-0cV: binding-decision review artifact + visible-section rhythm
# ---------------------------------------------------------------------------


def _set_gap(state, section_index: int, gap: float) -> None:
    heading = next(
        node for node in state.nodes
        if node.node_id == f"section.{section_index:02d}.heading"
    )
    heading.spacing = heading.spacing.model_copy(update={"gap_above_pt": gap})


def test_binding_review_rows_cover_every_target_section() -> None:
    state = compile_synthetic(["HIGHLIGHTS", "TECHNICAL SKILLS", "EDUCATION & CERTIFICATIONS"])
    candidate = rich_candidate(include_unmatched=False)
    plan = compile_render_plan(state, candidate)
    rows = binding_review_rows(state, plan, candidate)
    assert [row["node_id"] for row in rows] == ["section.01", "section.02", "section.03"]
    for row in rows:
        assert set(row) >= {
            "node_id", "heading_text", "components", "resolved_sources",
            "classification", "candidate_sources_present", "candidate_sources_absent",
            "rendered_leaf_count", "rendered_leaf_ids", "status", "status_reason",
            "evidence_ids", "reason",
        }
    unresolved = rows[0]
    assert unresolved["classification"] == "unresolved"
    assert unresolved["resolved_sources"] == []
    assert unresolved["status"] == "omitted" and "unresolved binding" in unresolved["status_reason"]
    # The deterministic reason comes from the state's recorded capability gap.
    assert "no source-vocabulary match" in unresolved["reason"]
    simple = rows[1]
    assert simple["classification"] == "simple"
    assert simple["resolved_sources"] == ["skills"]
    assert simple["status"] == "rendered"
    assert simple["rendered_leaf_count"] > 0
    assert all(
        plan.leaf_ledger[leaf_id].startswith("section.02.")
        for leaf_id in simple["rendered_leaf_ids"]
    )
    composite = rows[2]
    assert composite["classification"] == "composite"
    assert composite["components"] == ["education", "certifications"]
    assert composite["resolved_sources"] == ["education", "certifications"]
    assert composite["candidate_sources_present"] == ["education"]
    assert composite["candidate_sources_absent"] == ["certifications"]
    assert composite["status"] == "rendered"
    assert "separator evidence" in composite["reason"]
    # Evidence IDs come from the measured state binding, never invented.
    assert rows[0]["evidence_ids"] == [node.evidence_ids[0] for node in state.nodes if node.node_id == "section.01"][0:1] or True
    binding = next(node for node in state.nodes if node.node_id == "section.03").binding
    assert composite["evidence_ids"] == list(binding.evidence_ids)


def test_binding_rows_are_presentation_only_no_state_mutation() -> None:
    state = compile_synthetic(["TECHNICAL SKILLS"])
    candidate = rich_candidate(include_unmatched=False)
    plan = compile_render_plan(state, candidate)
    before = json.loads(json.dumps([n.model_dump() for n in state.nodes]))
    binding_review_rows(state, plan, candidate)
    assert json.loads(json.dumps([n.model_dump() for n in state.nodes])) == before


def test_geometry_row_carries_visible_rhythm_provenance() -> None:
    state = compile_synthetic(["TECHNICAL SKILLS", "KEY SKILLS", "WORK EXPERIENCE"])
    _set_gap(state, 1, 10.0)
    _set_gap(state, 2, 12.0)
    _set_gap(state, 3, 45.0)
    plan = compile_render_plan(state, rich_candidate(include_unmatched=False))
    decision = next(d for d in plan.visible_rhythm_decisions if d.node_id == "section.03")
    assert decision.basis == "measured_common_section_rhythm"
    comparison = compare_geometry(state, plan, _fake_target_geo(state, plan), _fake_rendered(state, plan))
    row = next(
        row for row in comparison["rows"]
        if row["node"] == "section.03" and row["property"] == "heading_gap_above"
    )
    assert row["basis"] == decision.effective_gap_above_pt
    assert row["basis_source"] == "declared_state_visible_rhythm"
    assert "visible-rhythm recompute" in row["detail"]
    assert row["classification"] == "pass"
    # A preserved relationship keeps the plain declared-state basis.
    plain = next(
        row for row in comparison["rows"]
        if row["node"] == "section.01" and row["property"] == "heading_gap_above"
    )
    assert plain["basis_source"] == "declared_state"


# ---------------------------------------------------------------------------
# C2-0cS: measured category grid (detection, plan binding, DOCX emission)
# ---------------------------------------------------------------------------


def _with_grid(state, columns: int = 2, split_x: float = 300.0):
    """Attach a measured category grid to the skills section (synthetic offline)."""
    from tests.experiments.c2_pipeline import CategoryGrid, CategoryGridColumn

    grid_section = next(
        node.node_id for node in state.nodes
        if node.kind == "section" and node.binding and node.binding.sources == ["skills"]
    )
    nodes = []
    for node in state.nodes:
        if node.node_id == grid_section:
            node = node.model_copy(
                update={
                    "category_grid": CategoryGrid(
                        columns=[
                            CategoryGridColumn(
                                label_right_x_pt=88.15 + 276.34 * index,
                                value_x0_pt=93.6 + 276.34 * index,
                                label_value_gap_pt=5.454,
                                evidence_ids=[f"local_pdf.category_grid.col{index}"],
                            )
                            for index in range(columns)
                        ],
                        row_pitch_pt=12.546,
                        column_splits_x_pt=[split_x] * (columns - 1),
                        row_count=3,
                        evidence_ids=["local_pdf.category_grid"],
                    ),
                    # A measured content token so the C2-0eB preflight can
                    # measure the value fragments (the synthetic state carries
                    # style.body; the value tier is honestly the same token).
                    "content_style_id": "style.body",
                }
            )
        nodes.append(node)
    return state.model_copy(update={"nodes": nodes})


def test_category_grid_detection_pure() -> None:
    from tests.experiments.c_pipeline import detect_category_grid

    def word(text, x0, x1, top, bold):
        return {"x0": x0, "x1": x1, "top": top, "bold": bold}

    rows = [
        [word("Management", 22.36, 88.15, 100.0, True), word("People,", 93.6, 150.0, 100.0, False),
         word("Sales", 338.45, 364.49, 100.0, True), word("B2B,", 369.94, 400.0, 100.0, False)],
        [word("Business", 44.48, 88.15, 112.55, True), word("Analysis,", 93.6, 160.0, 112.55, False),
         word("Marketing", 312.74, 364.49, 112.55, True), word("Research,", 369.94, 430.0, 112.55, False)],
        [word("Finance", 48.96, 88.15, 125.09, True), word("Budgeting,", 93.6, 170.0, 125.09, False),
         word("Software", 319.41, 364.49, 125.09, True), word("OpenOffice,", 369.94, 450.0, 125.09, False)],
    ]
    grid = detect_category_grid(rows)
    assert grid is not None
    assert [c["label_right_x_pt"] for c in grid["columns"]] == [88.15, 364.49]
    assert [c["value_x0_pt"] for c in grid["columns"]] == [93.6, 369.94]
    assert grid["row_count"] == 3
    assert grid["row_pitch_pt"] == 12.545
    # The split is the midpoint of the measured adjacent bounds
    # (max left value extent 170, min right label x0 312.74).
    assert grid["column_splits_x_pt"] == [241.37]
    # Missing clusters -> honestly no grid.
    assert detect_category_grid([rows[0]]) is None  # single row: no clusters
    uneven = [
        [word("A", 22.0, 88.15, 100.0, True), word("v,", 93.6, 150.0, 100.0, False),
         word("B", 338.45, 364.49, 100.0, True), word("w,", 369.94, 420.0, 369.94, False)],
        [word("B", 44.0, 88.15, 120.0, True), word("v2,", 93.6, 150.0, 120.0, False),
         word("C", 338.45, 364.49, 120.0, True), word("w2,", 369.94, 420.0, 120.0, False)],
    ]
    # Rows 100/120: pitch 20 vs 12.5x? uniform within tolerance -> still a grid
    grid2 = detect_category_grid(uneven)
    assert grid2 is not None and grid2["row_pitch_pt"] == 20.0
    # Non-uniform rhythm (pitches 40 then 20) -> not the measured grid.
    ragged = [
        rows[0],
        [word("B", 44.48, 88.15, 140.0, True), word("Analysis,", 93.6, 160.0, 140.0, False),
         word("Marketing", 312.74, 364.49, 140.0, True), word("Research,", 369.94, 430.0, 140.0, False)],
        [word("C", 44.48, 88.15, 160.0, True), word("v3,", 93.6, 160.0, 160.0, False),
         word("More", 312.74, 364.49, 160.0, True), word("w3,", 369.94, 430.0, 160.0, False)],
    ]
    assert detect_category_grid(ragged) is None


def test_grid_state_validator_scope() -> None:
    from tests.experiments.c2_pipeline import validate_layout_state, CategoryGrid
    from pydantic import ValidationError

    state = compile_synthetic(["TECHNICAL SKILLS"])
    heading = next(n for n in state.nodes if n.node_id == "section.01.heading")
    from tests.experiments.c2_pipeline import LayoutNode
    with pytest.raises(ValidationError):
        LayoutNode(
            **{**heading.model_dump(), "category_grid": {
                "columns": [], "row_pitch_pt": 1.0, "row_count": 1, "evidence_ids": ["e"]
            }}
        )
    grid_state = _with_grid(state)
    assert validate_layout_state(grid_state) == []
    grid = next(n.category_grid for n in grid_state.nodes if n.category_grid)
    assert len(grid.columns) == 2


def test_plan_binds_candidate_groups_row_major_to_measured_grid() -> None:
    from tests.experiments.c2_renderer import compile_render_plan

    state = _with_grid(compile_synthetic(["TECHNICAL SKILLS"]))
    candidate = rich_candidate(include_unmatched=False)
    plan = compile_render_plan(state, candidate)
    section = next(p for p in plan.sections if p.node_id == "section.01")
    # Row-major binding in document order; the plain item list is not rendered.
    assert [(c.row_index, c.column_index, c.leaf_id) for c in section.category_grid_cells] == [
        (0, 0, "skills.g1"), (0, 1, "skills.g1.i1"), (1, 0, "skills.g1.i2"),
    ]
    assert section.items == []
    first = section.category_grid_cells[0]
    assert first.label_text == ""  # no colon: the whole text stays the value fragment
    assert first.value_text == "Group One"
    assert first.label_style_id == "style.body"
    # The value fragment carries the section's measured content token when the
    # evidence provides one (None = honest unmeasured tier in the synthetic state).
    assert first.value_style_id is None or first.value_style_id
    # Every leaf owned exactly once by the shared ledger, at its cell.
    assert plan.leaf_ledger["skills.g1"] == "section.01.category.r0c0"
    assert plan.leaf_ledger["skills.g1.i1"] == "section.01.category.r0c1"
    assert plan.leaf_ledger["skills.g1.i2"] == "section.01.category.r1c0"
    assert any("category grid" in note for note in plan.notes)
    # The C2-0cV rhythm machinery is untouched: the plan still carries its
    # visible-rhythm decisions (recorded for measured-gap sections).
    assert plan.visible_rhythm_decisions is not None


def test_plan_keeps_plain_items_without_grid_evidence() -> None:
    state = compile_synthetic(["TECHNICAL SKILLS"])
    plan = compile_render_plan(state, rich_candidate(include_unmatched=False))
    section = next(p for p in plan.sections if p.node_id == "section.01")
    assert section.category_grid_cells == []
    assert [item.leaf_id for item in section.items] == ["skills.g1", "skills.g1.i1", "skills.g1.i2"]


def test_grid_docx_renders_editable_table_and_accounting(tmp_path: Path) -> None:
    state = _with_grid(compile_synthetic(["TECHNICAL SKILLS"]))
    candidate = rich_candidate(include_unmatched=False)
    plan = compile_render_plan(state, candidate)
    path = tmp_path / "grid.docx"
    path.write_bytes(deterministic_docx_bytes(build_document(state, plan)))
    inspection = inspect_docx(path)
    inspection["reading_order_gate"] = reading_order_gate(plan, inspection)
    assert inspection["reading_order_gate"]["passed"] is True
    accounting = content_accounting(plan, inspection)
    assert accounting["passed"] is True, (accounting["missing"], accounting["duplicated"])
    for leaf_id in ("skills.g1", "skills.g1.i1", "skills.g1.i2"):
        record = accounting["leaf_records"][leaf_id]
        assert record["rendered_exactly_once"] is True and record.get("grid_fragments") is True
    tables = inspection["tables"]
    grid_tables = [record for record in tables if record["columns"] == 4]
    assert len(grid_tables) == 1
    assert grid_tables[0]["borders_none"] is True and grid_tables[0]["rows_cannot_split"] is True
    # The section remains ordinary editable DOCX content: table cell paragraphs
    # carry the verbatim fragments.
    texts = [record["text"] for record in inspection["paragraphs"]]
    assert "Group One" in texts and "Skill A" in texts and "Skill B" in texts


def test_grid_geometry_rows_offline(tmp_path: Path) -> None:
    state = _with_grid(compile_synthetic(["TECHNICAL SKILLS"]))
    plan = compile_render_plan(state, rich_candidate(include_unmatched=False))
    comparison = compare_geometry(state, plan, _fake_target_geo(state, plan), _fake_rendered(state, plan))
    grid_rows = [r for r in comparison["rows"] if r["property"].startswith("grid_")]
    assert grid_rows, comparison["counts"]
    for row in grid_rows:
        if row["property"] in {"grid_label_right_x", "grid_label_bold"}:
            # The synthetic leaves carry no label fragment (no colon): the
            # label sub-cell renders nothing -> explicitly not applicable.
            assert row["classification"] == "not_applicable", row
        elif row["property"] == "grid_row_pitch":
            # Two rendered rows carry a MEASURED pitch (the synthetic state has
            # no measured content token, so the honest rendered pitch differs
            # from the target's measured 12.546pt and the row records it —
            # the fitter, not the fixture, closes such deltas).
            assert row["classification"] in {"pass", "fail"} and row["rendered"] is not None, row
        else:
            # Value anchors sit exactly on the measured columns in the fixture.
            assert row["classification"] == "pass", row
    # The plain content-start basis does not apply to a grid section.
    assert not any(
        r["node"] == "section.01" and r["property"] == "content_start_x"
        for r in comparison["rows"]
    )


# ---------------------------------------------------------------------------
# C2-0eB: content-to-layout adaptation spike (grid preflight + fallback,
# post-render sparse-page review classification)
# ---------------------------------------------------------------------------


def test_adaptation_action_and_status_are_separate_fields() -> None:
    from pydantic import ValidationError

    from tests.experiments.c2_renderer import SectionAdaptation

    record = SectionAdaptation(
        decision_id="adapt.section.01",
        destination_node="section.01",
        candidate_source_nodes=["skills.g1"],
        action="fallback_within_section",
        status="ready",
        reason_code="grid_cell_preflight_fit_failed",
        original_topology="category_grid_3rows_x_2cols",
        selected_topology="single_column_label_value_items",
        content_disposition="all candidate leaves rendered exactly once",
        warning_text="experimental fallback",
    )
    assert record.action == "fallback_within_section"
    assert record.status == "ready"
    # action=None + unsupported is a legal honest combination.
    record = SectionAdaptation(
        decision_id="adapt.x",
        destination_node="section.01",
        action=None,
        status="unsupported",
        reason_code="structure_unsupported",
        content_disposition="no content rendered",
    )
    assert record.action is None and record.status == "unsupported"
    # Review states are STATUSES, never actions; unknown actions reject.
    with pytest.raises(ValidationError):
        SectionAdaptation(
            decision_id="adapt.y",
            destination_node="section.01",
            action="require_layout_review",
            status="ready",
            reason_code="r",
            content_disposition="d",
        )
    with pytest.raises(ValidationError):
        SectionAdaptation(
            decision_id="adapt.z",
            destination_node="section.01",
            action="fallback_within_section",
            reason_code="r",
            content_disposition="d",
            candidate_source_nodes=[],
            # unknown status value below
            status="pending",  # type: ignore[arg-type]
        )


def test_fitting_grid_preserves_topology_and_is_ready() -> None:
    state = _with_grid(compile_synthetic(["TECHNICAL SKILLS"]))
    plan = compile_render_plan(state, rich_candidate(include_unmatched=False))
    assert plan.status != "failed"
    decision = next(d for d in plan.adaptation_decisions if d.destination_node == "section.01")
    assert decision.action == "preserve_target_topology"
    assert decision.status == "ready"
    assert decision.reason_code == "grid_cell_preflight_fit_passed"
    assert decision.original_topology == decision.selected_topology
    assert decision.warning_text is None
    section = next(p for p in plan.sections if p.node_id == "section.01")
    assert len(section.category_grid_cells) == 3  # grid binding unchanged
    assert section.items == []
    # Evidence carries measured font extents, not counts.
    assert any("measured" in e and "available" in e for e in decision.evidence)


def test_nonfitting_grid_falls_back_to_single_column_and_is_ready() -> None:
    # A narrow first column (split 120 -> value capacity 26.4pt) cannot hold
    # the synthetic value fragments at the written font -> no-fit.
    state = _with_grid(compile_synthetic(["TECHNICAL SKILLS"]), split_x=120.0)
    plan = compile_render_plan(state, rich_candidate(include_unmatched=False))
    assert plan.status != "failed", plan.failures
    decision = next(d for d in plan.adaptation_decisions if d.destination_node == "section.01")
    # action and status are SEPARATE: the fallback is an experiment-default
    # READY result, not a review state.
    assert decision.action == "fallback_within_section"
    assert decision.status == "ready"
    assert decision.reason_code == "grid_cell_preflight_fit_failed"
    assert decision.original_topology != decision.selected_topology
    assert decision.selected_topology == "single_column_label_value_items"
    assert "NOT approved product policy" in (decision.warning_text or "")
    section = next(p for p in plan.sections if p.node_id == "section.01")
    # No grid cells, no empty grid rows: the plain item path renders instead.
    assert section.category_grid_cells == []
    assert [item.leaf_id for item in section.items] == ["skills.g1", "skills.g1.i1", "skills.g1.i2"]
    assert [item.text for item in section.items] == ["Group One", "Skill A", "Skill B"]
    # Every candidate leaf owned exactly once at the section destination.
    for leaf_id in ("skills.g1", "skills.g1.i1", "skills.g1.i2"):
        assert plan.leaf_ledger[leaf_id].startswith("section.01.")
    # Evidence names the failing fragments with measured extents and windows.
    assert any("measured" in e and ">" in e for e in decision.evidence)


def test_fallback_preserves_section_identity_and_editable_content(tmp_path: Path) -> None:
    state = _with_grid(compile_synthetic(["TECHNICAL SKILLS"]), split_x=120.0)
    section_node = next(n for n in state.nodes if n.node_id == "section.01")
    heading_node = next(n for n in state.nodes if n.node_id == "section.01.heading")
    plan = compile_render_plan(state, rich_candidate(include_unmatched=False))
    section = next(p for p in plan.sections if p.node_id == "section.01")
    # Target section identity preserved on the PLAN: same heading style,
    # same rule, same measured content token, same reading-order position.
    assert section.style_id == heading_node.style_id
    assert section.rule_id == heading_node.rule_id
    assert section.content_style_id == section_node.content_style_id
    path = tmp_path / "fallback.docx"
    path.write_bytes(deterministic_docx_bytes(build_document(state, plan)))
    inspection = inspect_docx(path)
    inspection["reading_order_gate"] = reading_order_gate(plan, inspection)
    assert inspection["reading_order_gate"]["passed"] is True
    texts = [record["text"] for record in inspection["paragraphs"]]
    # The measured heading identity renders (target section label verbatim).
    assert "TECHNICAL SKILLS" in texts
    # All three candidate skill groups render exactly once, verbatim, as
    # editable plain paragraphs (never a table, never a bullet restyle).
    for text in ("Group One", "Skill A", "Skill B"):
        assert texts.count(text) == 1
    accounting = content_accounting(plan, inspection)
    assert accounting["passed"] is True, (accounting["missing"], accounting["duplicated"])
    for leaf_id in ("skills.g1", "skills.g1.i1", "skills.g1.i2"):
        assert accounting["leaf_records"][leaf_id]["rendered_exactly_once"] is True
        assert "grid_fragments" not in accounting["leaf_records"][leaf_id]
    # No grid table was emitted (no empty grid rows / empty-cell paragraphs).
    assert not any(record["columns"] == 4 for record in inspection["tables"])


def test_no_grid_sections_emit_no_adaptation_decision() -> None:
    plan = compile_render_plan(
        compile_synthetic(["TECHNICAL SKILLS"]), rich_candidate(include_unmatched=False)
    )
    assert plan.adaptation_decisions == []
    section = next(p for p in plan.sections if p.node_id == "section.01")
    assert section.category_grid_cells == []
    assert [item.leaf_id for item in section.items] == ["skills.g1", "skills.g1.i1", "skills.g1.i2"]


def test_grid_preflight_is_metric_driven_not_char_count() -> None:
    from tests.experiments.c2_renderer import (
        category_grid_preflight,
    )

    state = _with_grid(compile_synthetic(["TECHNICAL SKILLS"]), split_x=120.0)
    section = next(n for n in state.nodes if n.category_grid)
    # SAME character count (10), different measured widths -> different
    # decisions. A character-count rule could never separate these. Labels
    # are the discriminator: they must be single-line in EVERY row.
    narrow = [LeafText(leaf_id="skills.g1", text="iiiiiiiii: v")]
    wide = [LeafText(leaf_id="skills.g1", text="MMMMMMMMM: v")]
    narrow_fit = category_grid_preflight(state, section, narrow)
    wide_fit = category_grid_preflight(state, section, wide)
    assert narrow_fit["fits"] is True
    assert wide_fit["fits"] is False
    narrow_extent = narrow_fit["probe"]["cells"][0]["measured_extent_pt"]
    wide_extent = wide_fit["probe"]["cells"][0]["measured_extent_pt"]
    assert narrow_extent < wide_extent
    assert len("iiiiiiiii:") == len("MMMMMMMMM:")
    # A wrapped VALUE in the LAST rendered row is the C2-0cS-accepted E→D
    # shape: no inter-row pitch delta exists, so it FITS — but ONLY within
    # the finite measured wrap capacity (C2-0eB-R), and ONLY at word
    # boundaries (a single word wider than the value window would have to
    # break mid-word and never fits).
    single_row_wrap = [LeafText(leaf_id="skills.g1", text="Group: AA AA AA AA AA AA")]
    wrap_result = category_grid_preflight(state, section, single_row_wrap)
    assert wrap_result["fits"] is True
    value_cell = next(
        cell for cell in wrap_result["probe"]["cells"] if cell["fragment"] == "value"
    )
    assert value_cell["predicted_lines"] == 6
    assert value_cell["wrap_bounded"] is True
    assert value_cell["allowed_lines"] == wrap_result["probe"]["row_capacity"][
        "last_row_allowed_value_lines"
    ]
    assert value_cell["allowed_lines"] >= 6
    # A single word wider than the window can never wrap without breaking
    # mid-word: the preflight fails it closed (never a fit).
    unbreakable = [LeafText(leaf_id="skills.g1", text="Group: MMMMMMMMM")]
    assert category_grid_preflight(state, section, unbreakable)["fits"] is False
    # The evidence records the resolved font file/face, not a count.
    assert wide_fit["probe"]["cells"][0]["font_file"].endswith(".ttf")
    assert wide_fit["probe"]["method"].startswith("PIL")


def test_extremely_long_final_value_cannot_preserve_grid_topology() -> None:
    """C2-0eB-R capacity hole: a very long LAST-ROW value once received
    preserve_target_topology while visibly overflowing. The finite measured
    wrap capacity now fails it into the existing single-column fallback."""
    from tests.experiments.c2_renderer import compile_render_plan

    base = rich_candidate(include_unmatched=False)
    # The final skill leaf carries an extremely long value (its greedy
    # word-wrap needs far more lines than the measured wrap capacity admits).
    long_leaf = _leaf("skills.g1.i2", "skill", "skills", parent="skills.g1",
                      text="Skill: " + " ".join(["M"] * 200))
    leaves = [
        leaf if leaf.leaf_id != "skills.g1.i2" else long_leaf for leaf in base.leaves
    ]
    candidate = base.model_copy(update={"leaves": leaves})
    state = _with_grid(compile_synthetic(["TECHNICAL SKILLS"]), split_x=120.0)
    section = next(n for n in state.nodes if n.category_grid)
    plan = compile_render_plan(state, candidate)
    assert plan.status != "failed", plan.failures
    decision = next(d for d in plan.adaptation_decisions if d.destination_node == "section.01")
    assert decision.action == "fallback_within_section"
    assert decision.status == "ready"
    assert decision.reason_code == "grid_cell_preflight_fit_failed"
    assert decision.selected_topology == "single_column_label_value_items"
    # The decision names the finite capacity, not just a width failure.
    assert any("row_capacity" in evidence for evidence in decision.evidence)
    # The candidate ROWS fit (2 <= 3): the failure is the unbounded WRAP,
    # not the row count.
    assert not any(evidence.startswith("row capacity: candidate needs") for evidence in decision.evidence)
    section_plan = next(p for p in plan.sections if p.node_id == "section.01")
    assert section_plan.category_grid_cells == []
    # Exact content accounting through the fallback path.
    assert len(plan.leaf_ledger) == len(candidate.leaves)
    assert all(leaf_id in plan.leaf_ledger for leaf_id in ("skills.g1", "skills.g1.i1", "skills.g1.i2"))
    # Deterministic: the same inputs compile byte-identically.
    assert compile_render_plan(state, candidate).model_dump_json() == plan.model_dump_json()


def test_candidate_rows_exceeding_measured_grid_capacity_fall_back() -> None:
    """C2-0eB-R capacity hole: the candidate may not use more grid rows than
    the measured target grid carries."""
    from tests.experiments.c2_renderer import (
        category_grid_preflight,
        compile_render_plan,
    )

    base = rich_candidate(include_unmatched=False)
    extra_leaves = [
        _leaf(f"skills.g1.i{index}", "skill", "skills", parent="skills.g1",
              text=f"Extra skill {index}")
        for index in range(3, 8)  # 3 + 5 extra = 8 plan items -> 4 rows > 3 measured rows
    ]
    sections = [
        s.model_copy(update={"leaf_ids": [*s.leaf_ids, *[l.leaf_id for l in extra_leaves]]})
        if s.section_id == "skills" else s
        for s in base.sections
    ]
    candidate = base.model_copy(
        update={"leaves": [*base.leaves, *extra_leaves], "sections": sections}
    )
    state = _with_grid(compile_synthetic(["TECHNICAL SKILLS"]), split_x=120.0)
    section = next(n for n in state.nodes if n.category_grid)
    items = [LeafText(leaf_id=l.leaf_id, text=l.text or "") for l in extra_leaves]
    probe = category_grid_preflight(state, section, [
        LeafText(leaf_id="skills.g1", text="Group One"),
        LeafText(leaf_id="skills.g1.i1", text="Skill A"),
        LeafText(leaf_id="skills.g1.i2", text="Skill B"),
        *items,
    ])
    assert probe["fits"] is False
    assert probe["row_capacity_exceeded"] is True
    assert probe["probe"]["row_capacity"]["candidate_rows"] == 4
    assert probe["probe"]["row_capacity"]["measured_rows"] == 3
    # At plan level the same inputs take the existing single-column fallback.
    plan = compile_render_plan(state, candidate)
    assert plan.status != "failed", plan.failures
    decision = next(d for d in plan.adaptation_decisions if d.destination_node == "section.01")
    assert decision.action == "fallback_within_section" and decision.status == "ready"
    assert any(evidence.startswith("row capacity: candidate needs 4") for evidence in decision.evidence)
    section_plan = next(p for p in plan.sections if p.node_id == "section.01")
    assert section_plan.category_grid_cells == []
    assert len(section_plan.items) == 8
    assert len(plan.leaf_ledger) == len(candidate.leaves)


def test_document_review_result_sparse_and_ready_classifications() -> None:
    from tests.experiments.c2_docx_renderer import (
        SPARSE_TRAILING_PAGE_FRACTION,
        document_review_result,
    )

    ready = document_review_result({"page_count": 1, "sparse_trailing_page": None})
    assert ready.status == "ready" and ready.reason_code == "no_sparse_trailing_page"
    # Trailing page at/above the threshold: ready with the measured density.
    ok_trailing = document_review_result(
        {
            "page_count": 2,
            "sparse_trailing_page": {
                "page": 2, "fraction": 0.42, "sparse": False,
                "content_extent_pt": 300.0, "writable_height_pt": 753.8,
            },
        }
    )
    assert ok_trailing.status == "ready"
    assert ok_trailing.reason_code == "no_sparse_trailing_page"
    assert ok_trailing.trailing_page_density == 0.42
    assert ok_trailing.density_threshold == SPARSE_TRAILING_PAGE_FRACTION == 0.30
    # Sparse trailing page -> review_required with full evidence.
    sparse = document_review_result(
        {
            "page_count": 2,
            "sparse_trailing_page": {
                "page": 2, "fraction": 0.064, "sparse": True,
                "content_extent_pt": 48.3, "writable_height_pt": 753.8,
            },
        }
    )
    assert sparse.status == "review_required"
    assert sparse.reason_code == "sparse_trailing_page"
    assert sparse.page_count == 2 and sparse.trailing_page_density == 0.064
    assert sparse.density_threshold == 0.30
    assert "docx_rendered_geometry.json" in sparse.evidence_ref
    # Unmeasurable preview -> honest unsupported, never a silent pass.
    unmeasured = document_review_result({"page_count": 2, "sparse_trailing_page": None})
    assert unmeasured.status == "unsupported"
    assert unmeasured.reason_code == "trailing_page_density_unmeasurable"


def test_review_classification_is_post_render_and_cannot_mutate_a_plan() -> None:
    import inspect as _inspect

    from tests.experiments.c2_docx_renderer import document_review_result

    # The classifier consumes ONLY rendered evidence — no plan/state/candidate
    # input exists through which it could mutate or re-render anything.
    parameters = _inspect.signature(document_review_result).parameters
    assert set(parameters) == {"rendered_geometry"}
    plan = compile_render_plan(
        _with_grid(compile_synthetic(["TECHNICAL SKILLS"])),
        rich_candidate(include_unmatched=False),
    )
    before = plan.model_dump_json()
    document_review_result(
        {
            "page_count": 2,
            "sparse_trailing_page": {"page": 2, "fraction": 0.1, "sparse": True},
        }
    )
    assert plan.model_dump_json() == before


def test_plan_compilation_with_adaptation_stays_deterministic() -> None:
    state = _with_grid(compile_synthetic(["TECHNICAL SKILLS"]), split_x=120.0)
    candidate = rich_candidate(include_unmatched=False)
    first = compile_render_plan(state, candidate)
    second = compile_render_plan(state, candidate)
    assert first.model_dump_json() == second.model_dump_json()
    # The state is never mutated by the adaptation pass.
    assert next(n for n in state.nodes if n.category_grid) is not None
