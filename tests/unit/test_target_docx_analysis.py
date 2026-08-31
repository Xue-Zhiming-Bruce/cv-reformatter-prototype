from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from app.template_analysis.docx_style_analyzer import (
    TargetDocxAnalysisError,
    analyze_target_docx,
)


def test_target_docx_analysis_extracts_page_and_typography(tmp_path: Path) -> None:
    target_path = tmp_path / "target.docx"
    document = Document()
    section = document.sections[0]
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.7)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)
    heading = document.styles["Heading 2"]
    heading.font.name = "Times New Roman"
    heading.font.size = Pt(12)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor(0x12, 0x34, 0x56)
    document.save(target_path)

    spec = analyze_target_docx(target_path, tmp_path / "analysis")

    assert spec.source_type == "docx_analysis"
    assert spec.page.margin_left_pt == pytest.approx(43.2, abs=0.1)
    assert spec.page.margin_right_pt == pytest.approx(50.4, abs=0.1)
    assert spec.body.font_family == "Times New Roman"
    assert spec.body.font_size_pt == pytest.approx(10.5)
    assert spec.heading.font_size_pt == pytest.approx(12)
    assert spec.heading.bold is True
    assert spec.heading.color_hex == "#123456"
    assert spec.schema_version == "2.0"
    assert spec.structure_contract == "limited_capability"
    assert "limited_capability:docx_section_structure_not_measured" in spec.unsupported_features
    assert (tmp_path / "analysis" / "layout_template_spec.json").exists()
    assert not (tmp_path / "analysis" / "template_style_spec.json").exists()


def test_target_docx_analysis_rejects_corrupt_file(tmp_path: Path) -> None:
    target_path = tmp_path / "target.docx"
    target_path.write_bytes(b"not a docx")

    with pytest.raises(TargetDocxAnalysisError, match="corrupt"):
        analyze_target_docx(target_path, tmp_path / "analysis")


def _set_bottom_border(value, *, color: str, size_eighths_pt: int) -> None:
    properties = value._p.get_or_add_pPr() if hasattr(value, "_p") else value.element.get_or_add_pPr()
    borders = properties.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        properties.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:color"), color)
    bottom.set(qn("w:sz"), str(size_eighths_pt))
    borders.append(bottom)


def test_target_docx_analysis_measures_style_rules(tmp_path: Path) -> None:
    target_path = tmp_path / "ruled.docx"
    document = Document()
    document.styles["Normal"].font.name = "Times New Roman"
    _set_bottom_border(document.styles["Title"], color="111111", size_eighths_pt=12)
    _set_bottom_border(document.styles["Heading 2"], color="123456", size_eighths_pt=6)
    document.add_paragraph("Synthetic title", style="Title")
    document.add_paragraph("Synthetic heading", style="Heading 2")
    document.save(target_path)

    spec = analyze_target_docx(target_path, tmp_path / "analysis")

    assert spec.decoration.header_rule is True
    assert spec.decoration.heading_rule is True
    assert spec.decoration.rule_color_hex == "#123456"
    assert spec.decoration.rule_width_pt == pytest.approx(0.75)


def test_target_docx_analysis_does_not_globalize_partial_heading_rules(
    tmp_path: Path,
) -> None:
    target_path = tmp_path / "partial-rules.docx"
    document = Document()
    document.styles["Normal"].font.name = "Times New Roman"
    headings = [
        document.add_paragraph(f"Synthetic heading {index}", style="Heading 2")
        for index in range(3)
    ]
    for heading in (headings[0], headings[2]):
        _set_bottom_border(heading, color="123456", size_eighths_pt=6)
    document.save(target_path)

    spec = analyze_target_docx(target_path, tmp_path / "analysis")

    assert spec.decoration.heading_rule is False
    assert any("Non-uniform Heading 2" in warning for warning in spec.warnings)


def test_target_docx_analysis_emits_measured_v2_structure_contract(
    tmp_path: Path,
) -> None:
    target_path = tmp_path / "structured.docx"
    document = Document()
    document.styles["Normal"].font.name = "Times New Roman"
    heading = document.styles["Heading 2"]
    heading.font.name = "Times New Roman"
    heading.font.size = Pt(12)
    run_properties = heading.element.get_or_add_rPr()
    tracking = OxmlElement("w:spacing")
    tracking.set(qn("w:val"), "14")
    run_properties.append(tracking)
    document.styles["Heading 3"].font.name = "Times New Roman"
    document.styles["Heading 3"].font.size = Pt(11)
    document.add_paragraph("synthetic@example.test | +65 6123 4567")
    document.add_paragraph("PROFESSIONAL SUMMARY", style="Heading 2")
    document.add_paragraph("Synthetic summary")
    document.add_paragraph("WORK EXPERIENCE", style="Heading 2")
    document.add_paragraph("PROJECTS", style="Heading 2")
    document.save(target_path)

    spec = analyze_target_docx(target_path, tmp_path / "analysis")

    assert spec.structure_contract == "measured"
    assert spec.header.contact_fields == ["email", "phone"]
    assert [section.label for section in spec.sections] == [
        "PROFESSIONAL SUMMARY",
        "WORK EXPERIENCE",
        "PROJECTS",
    ]
    assert spec.sections[0].heading_style.character_spacing_pt == pytest.approx(0.7)
    assert spec.sections[1].entry_title_style.font_size_pt == pytest.approx(11)
    assert spec.sections[2].additional_section_heading == "PROJECTS"
