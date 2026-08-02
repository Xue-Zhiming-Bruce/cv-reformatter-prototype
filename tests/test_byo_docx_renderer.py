from pathlib import Path

from docx import Document

from app.extraction.candidate_schema import CandidateProfile, WorkExperience
from app.generation.byo_docx_renderer import (
    build_docxtpl_context,
    docx_has_placeholders,
    render_byo_docx,
)


def _write_template(tmp_path: Path, *paragraphs: str) -> Path:
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    path = tmp_path / "template.docx"
    doc.save(str(path))
    return path


def _docx_text(path: Path) -> str:
    return "\n".join(p.text for p in Document(path).paragraphs)


# ── docx_has_placeholders ──────────────────────────────────────────────────


def test_has_placeholders_true_for_jinja_tags(tmp_path: Path) -> None:
    tpl = _write_template(tmp_path, "Dear {{ full_name }},")
    assert docx_has_placeholders(tpl) is True


def test_has_placeholders_true_for_block_tags(tmp_path: Path) -> None:
    tpl = _write_template(tmp_path, "{% for job in experience %}{{ job.title }}{% endfor %}")
    assert docx_has_placeholders(tpl) is True


def test_has_placeholders_false_for_plain_docx(tmp_path: Path) -> None:
    tpl = _write_template(tmp_path, "Hello, this is a plain document with no tags.")
    assert docx_has_placeholders(tpl) is False


# ── render_byo_docx ────────────────────────────────────────────────────────


def test_render_fills_name_and_job_title(tmp_path: Path) -> None:
    tpl = _write_template(
        tmp_path,
        "{{ full_name }}",
        "{% for job in experience %}{{ job.title }}{% endfor %}",
    )
    profile = CandidateProfile(
        full_name="Jane Candidate",
        work_experience=[WorkExperience(title="Senior Engineer", company="Acme Corp")],
    )
    output = tmp_path / "output.docx"
    render_byo_docx(tpl, profile, output, blind=False)

    assert output.exists()
    text = _docx_text(output)
    assert "Jane Candidate" in text
    assert "Senior Engineer" in text


def test_render_blind_blanks_identity(tmp_path: Path) -> None:
    tpl = _write_template(tmp_path, "{{ full_name }} {{ email }} {{ phone }}")
    profile = CandidateProfile(
        full_name="Jane Candidate",
        email="jane@example.com",
        phone="+1 555 0100",
    )
    output = tmp_path / "blind.docx"
    render_byo_docx(tpl, profile, output, blind=True)

    text = _docx_text(output)
    assert "Jane Candidate" not in text
    assert "jane@example.com" not in text
    assert "555" not in text


# ── build_docxtpl_context ─────────────────────────────────────────────────


def test_context_none_values_become_empty_string() -> None:
    profile = CandidateProfile()
    ctx = build_docxtpl_context(profile, blind=False)
    assert ctx["full_name"] == ""
    assert ctx["summary"] == ""
    assert ctx["salary_expectation"] == ""
