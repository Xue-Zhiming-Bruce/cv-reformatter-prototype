"""Stateful local Mock API for frontend development.

Run with:
    uvicorn scripts.mock_api:app --reload --port 8000

This module deliberately performs no extraction, AI calls, or persistence.
"""

from __future__ import annotations

import hashlib
import html
import io
import json
import uuid
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Literal

from docx import Document
from fastapi import FastAPI, File, Header, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import HTMLResponse, Response
from pydantic import BaseModel, ConfigDict, Field

from app.extraction.candidate_schema import (
    CandidateProfile,
    DisplayRule,
    Education,
    Language,
    MissingField,
    SkillGroup,
    WorkExperience,
)


class StrictModel(BaseModel):
    model_config = ConfigDict(extra="forbid", str_strip_whitespace=True)


class ApprovalRequest(StrictModel):
    profile: CandidateProfile
    reviewer_note: str | None = None


class GenerateRequest(StrictModel):
    artifact_id: str
    approved_profile_version_id: str
    blind_profile: bool = False
    output_format: Literal["pdf", "docx"] | None = None


class FollowupRequest(StrictModel):
    profile: CandidateProfile
    language: str = "English"


class LayoutEditRequest(StrictModel):
    instruction: str = Field(min_length=1)
    selected_node_id: str | None = None
    scope: Literal["node", "section", "document"] = "node"
    base_layout_version_id: str


class LayoutEditDecision(StrictModel):
    decision: Literal["accept", "reject"]


@dataclass
class Artifact:
    artifact_id: str
    workflow_lane: Literal["pdf", "docx"]
    source_filename: str
    profile: CandidateProfile
    approved_profiles: dict[str, CandidateProfile] = field(default_factory=dict)
    target_filename: str | None = None
    active_layout_version_id: str = "layout_v1"
    layout_versions: set[str] = field(default_factory=lambda: {"layout_v1"})
    edits: dict[str, dict[str, Any]] = field(default_factory=dict)


app = FastAPI(
    title="CV Reformatter Mock API",
    version="0.1.0",
    description=(
        "A local, in-memory implementation of the frontend-facing API contract. "
        "Layout-edit routes are provisional product exploration endpoints."
    ),
)
app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173", "http://127.0.0.1:5173"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

ARTIFACTS: dict[str, Artifact] = {}


def _id(prefix: str) -> str:
    return f"{prefix}_{uuid.uuid4().hex[:12]}"


def _artifact(artifact_id: str) -> Artifact:
    try:
        return ARTIFACTS[artifact_id]
    except KeyError as exc:
        raise HTTPException(404, "Artifact not found.") from exc


def _lane(filename: str | None) -> Literal["pdf", "docx"]:
    suffix = Path(filename or "").suffix.lower()
    if suffix == ".pdf":
        return "pdf"
    if suffix == ".docx":
        return "docx"
    raise HTTPException(400, "Only .pdf and .docx files are supported.")


def _profile() -> CandidateProfile:
    return CandidateProfile(
        full_name="Alex Morgan",
        current_title="Senior Product Designer",
        email="alex.morgan@example.com",
        phone="+1 555 010 2048",
        location="Singapore",
        linkedin_url="https://www.linkedin.com/in/alex-morgan-example",
        professional_summary=(
            "Product designer with eight years of experience turning complex "
            "workflows into clear, accessible products."
        ),
        skills=["Product strategy", "Interaction design", "Design systems"],
        skill_groups=[
            SkillGroup(
                label="Design",
                skills=["Figma", "Prototyping", "User research"],
            )
        ],
        languages=[Language(name="English", proficiency="Fluent")],
        work_experience=[
            WorkExperience(
                company="Northstar Labs",
                title="Senior Product Designer",
                location="Singapore",
                start_date="2021",
                end_date="Present",
                description=[
                    "Led the redesign of a recruiter workflow used across three regions.",
                    "Built a shared design system that reduced delivery time by 30%.",
                ],
            )
        ],
        education=[
            Education(
                institution="Example University",
                degree="Bachelor of Design",
                end_date="2016",
            )
        ],
        missing_fields=[
            MissingField(
                field_name="notice_period",
                label="Notice period",
                reason="Not present in the source document.",
            )
        ],
        client_display_rules={"salary_expectation": DisplayRule.HIDE},
    )


def _ledger() -> dict[str, Any]:
    return {
        "source_count": 13,
        "reviewed_count": 12,
        "visible_count": 11,
        "hidden_count": 1,
        "pending_count": 1,
        "unresolved_count": 1,
    }


def _pdf_bytes(version: str) -> bytes:
    """Create a tiny visible PDF using only the PDF primitives needed here."""
    lines = [
        "Alex Morgan",
        "Senior Product Designer",
        f"Mock preview - {version}",
        "EXPERIENCE",
        "Northstar Labs | 2021 - Present",
        "- Led the redesign of a recruiter workflow.",
        "- Built a shared design system.",
    ]
    commands = ["BT", "/F1 22 Tf", "72 740 Td"]
    for index, line in enumerate(lines):
        safe = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        if index:
            commands.extend(["0 -34 Td", "/F1 11 Tf"])
        commands.append(f"({safe}) Tj")
    commands.append("ET")
    stream = "\n".join(commands).encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, obj in enumerate(objects, 1):
        offsets.append(len(output))
        output.extend(f"{number} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return bytes(output)


def _docx_bytes() -> bytes:
    document = Document()
    document.add_heading("Alex Morgan", 0)
    document.add_paragraph("Senior Product Designer")
    document.add_heading("Experience", level=1)
    document.add_paragraph("Northstar Labs | 2021 - Present")
    document.add_paragraph("Led the redesign of a recruiter workflow.", style="List Bullet")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _surface(artifact: Artifact, version: str) -> str:
    profile = artifact.profile
    return f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><title>Mock CV</title>
<style>body{{font:14px Arial;max-width:760px;margin:40px auto;color:#20242b}}
h1{{margin-bottom:4px}} h2{{border-bottom:1px solid #555;padding-bottom:5px}}
.meta{{color:#606873}} li{{margin:7px 0}}</style></head><body data-layout-version="{html.escape(version)}">
<header data-node-id="header"><h1 data-node-id="header.name">{html.escape(profile.full_name or "Candidate")}</h1>
<p>{html.escape(profile.current_title or "")}</p><p class="meta" data-node-id="header.contact_row">{html.escape(profile.email or "")}</p></header>
<section data-node-id="section.experience"><h2 data-node-id="section.experience.heading">Experience</h2>
<div data-node-id="section.experience.entries"><strong>Northstar Labs</strong><ul><li>Led the redesign of a recruiter workflow.</li><li>Built a shared design system.</li></ul></div></section>
</body></html>"""


@app.get("/health")
def health() -> dict[str, str]:
    return {"status": "ok", "service": "cv-reformatter-mock-api"}


@app.post("/__mock__/reset", include_in_schema=False)
def reset() -> dict[str, str]:
    ARTIFACTS.clear()
    return {"status": "reset"}


@app.post("/api/process")
async def process_candidate(
    file: UploadFile = File(...),
    x_mock_scenario: str = Header("happy", alias="X-Mock-Scenario"),
) -> dict[str, Any]:
    lane = _lane(file.filename)
    if x_mock_scenario == "processing_failed":
        raise HTTPException(422, "Mock extraction failed. Ask the user to upload a clearer file.")
    artifact_id = _id("artifact")
    profile = _profile()
    ARTIFACTS[artifact_id] = Artifact(artifact_id, lane, file.filename or "candidate", profile)
    return {
        "artifact_id": artifact_id,
        "workflow_lane": lane,
        "profile": profile,
        "original_filename": file.filename,
        "original_pdf_preview_url": f"/api/artifacts/{artifact_id}/previews/original.pdf",
        "ledger_summary": _ledger(),
        "warnings": ["Synthetic Mock API data; no source extraction was performed."],
    }


@app.post("/api/target-format")
async def upload_target_format(
    file: UploadFile = File(...), artifact_id: str | None = None
) -> dict[str, Any]:
    lane = _lane(file.filename)
    artifact = _artifact(artifact_id) if artifact_id else None
    if artifact and artifact.workflow_lane != lane:
        raise HTTPException(409, "Candidate and target files must use the same workflow lane.")
    if artifact:
        artifact.target_filename = file.filename
    return {
        "artifact_id": artifact_id,
        "workflow_lane": lane,
        "target_format": {
            "filename": file.filename,
            "document_type": lane,
            "page_count": 1,
            "analysis_status": "ready",
            "layout_template_spec_version": "mock-layout-spec-v1",
        },
        "warnings": ["Mock target analysis only; the uploaded bytes were not inspected."],
    }


@app.post("/api/artifacts/{artifact_id}/profiles/approve")
def approve_profile(artifact_id: str, request: ApprovalRequest) -> dict[str, Any]:
    artifact = _artifact(artifact_id)
    if any(section.review_state == "pending_review" for section in request.profile.additional_sections):
        raise HTTPException(409, "All additional sections must be reviewed before approval.")
    payload = request.profile.model_dump(mode="json")
    version_id = _id("profile")
    artifact.profile = request.profile
    artifact.approved_profiles[version_id] = request.profile
    return {
        "artifact_id": artifact_id,
        "approved_profile_version_id": version_id,
        "profile_checksum": hashlib.sha256(
            json.dumps(payload, sort_keys=True, separators=(",", ":")).encode()
        ).hexdigest(),
        "approved_at": "2026-09-12T12:00:00Z",
        "reviewer_note": request.reviewer_note,
    }


@app.post("/api/generate")
def generate(request: GenerateRequest) -> dict[str, Any]:
    artifact = _artifact(request.artifact_id)
    if request.approved_profile_version_id not in artifact.approved_profiles:
        raise HTTPException(409, "A valid approved profile version is required.")
    output_format = request.output_format or artifact.workflow_lane
    if output_format != artifact.workflow_lane:
        raise HTTPException(409, "Output format must match the artifact workflow lane.")
    base = f"/api/artifacts/{artifact.artifact_id}"
    response: dict[str, Any] = {
        "artifact_id": artifact.artifact_id,
        "workflow_lane": artifact.workflow_lane,
        "approved_profile_version_id": request.approved_profile_version_id,
        "layout_version_id": artifact.active_layout_version_id,
        "blind_profile": request.blind_profile,
        "warnings": [],
    }
    if output_format == "pdf":
        response.update(
            html_surface_url=f"{base}/surface.html",
            pdf_preview_url=f"{base}/previews/{artifact.active_layout_version_id}.pdf",
            pdf_download_url=f"{base}/generated.pdf",
        )
    else:
        response["docx_download_url"] = f"{base}/generated.docx"
    return response


@app.post("/api/followup")
def followup(request: FollowupRequest) -> dict[str, str]:
    name = request.profile.full_name or "the candidate"
    return {
        "message": (
            f"Please confirm the notice period and interview availability for {name}. "
            "The attached profile has been preserved for recruiter review."
        ),
        "language": request.language,
    }


@app.get("/api/artifacts/{artifact_id}/layout-nodes")
def layout_nodes(artifact_id: str) -> dict[str, Any]:
    artifact = _artifact(artifact_id)
    return {
        "artifact_id": artifact_id,
        "layout_version_id": artifact.active_layout_version_id,
        "nodes": [
            {"id": "header", "role": "section", "label": "Header"},
            {"id": "header.name", "role": "text", "label": "Candidate name"},
            {"id": "header.contact_row", "role": "contact_row", "label": "Contact row"},
            {"id": "section.experience", "role": "section", "label": "Experience"},
            {"id": "section.experience.heading", "role": "heading", "label": "Experience heading"},
            {"id": "section.experience.entries", "role": "list", "label": "Experience entries"},
        ],
    }


@app.post("/api/artifacts/{artifact_id}/layout-edits")
def create_layout_edit(
    artifact_id: str,
    request: LayoutEditRequest,
    x_mock_scenario: str = Header("happy", alias="X-Mock-Scenario"),
) -> dict[str, Any]:
    artifact = _artifact(artifact_id)
    if request.base_layout_version_id != artifact.active_layout_version_id:
        raise HTTPException(409, "The base layout version is stale.")
    edit_id = _id("edit")
    base = request.base_layout_version_id
    if x_mock_scenario == "needs_clarification":
        return {
            "edit_id": edit_id,
            "status": "needs_clarification",
            "base_layout_version_id": base,
            "layout_version_id": None,
            "operations": [],
            "changed_node_ids": [],
            "preview_url": None,
            "message": "Which bullet list should be changed? Select a section or node.",
            "warnings": [],
        }
    if x_mock_scenario == "unsupported":
        return {
            "edit_id": edit_id,
            "status": "unsupported",
            "base_layout_version_id": base,
            "layout_version_id": None,
            "operations": [],
            "changed_node_ids": [],
            "preview_url": None,
            "message": "That instruction is outside the supported layout-edit vocabulary.",
            "warnings": [],
        }
    if x_mock_scenario == "render_failed":
        return {
            "edit_id": edit_id,
            "status": "rejected",
            "base_layout_version_id": base,
            "layout_version_id": base,
            "operations": [],
            "changed_node_ids": [],
            "preview_url": f"/api/artifacts/{artifact_id}/previews/{base}.pdf",
            "message": "The preview failed validation; the previous layout remains active.",
            "warnings": ["Mock rollback completed."],
        }

    instruction = request.instruction.lower()
    target = request.selected_node_id or "document"
    if "bullet" in instruction:
        operation = {"type": "SetListStyle", "target": target, "changes": {"marker": "disc", "indent_pt": 10}}
    elif "line" in instruction or "rule" in instruction:
        operation = {"type": "SetHeadingRule", "target": target, "changes": {"placement": "inline_after", "gap_pt": 6}}
    else:
        operation = {"type": "SetSpacing", "target": target, "changes": {"after_pt": 8}}
    version = _id("layout")
    artifact.layout_versions.add(version)
    artifact.edits[edit_id] = {
        "status": "pending",
        "base_layout_version_id": base,
        "layout_version_id": version,
    }
    return {
        "edit_id": edit_id,
        "status": "applied",
        "base_layout_version_id": base,
        "layout_version_id": version,
        "operations": [operation],
        "changed_node_ids": [target],
        "preview_url": f"/api/artifacts/{artifact_id}/previews/{version}.pdf",
        "message": "Mock edit applied to a preview. Accept or reject it before generation.",
        "warnings": [],
    }


@app.post("/api/artifacts/{artifact_id}/layout-edits/{edit_id}/decision")
def decide_layout_edit(
    artifact_id: str, edit_id: str, request: LayoutEditDecision
) -> dict[str, Any]:
    artifact = _artifact(artifact_id)
    try:
        edit = artifact.edits[edit_id]
    except KeyError as exc:
        raise HTTPException(404, "Layout edit not found.") from exc
    if edit["status"] != "pending":
        raise HTTPException(409, "This layout edit already has a decision.")
    edit["status"] = "accepted" if request.decision == "accept" else "rejected"
    if request.decision == "accept":
        artifact.active_layout_version_id = edit["layout_version_id"]
    return {
        "artifact_id": artifact_id,
        "edit_id": edit_id,
        "status": edit["status"],
        "active_layout_version_id": artifact.active_layout_version_id,
        "preview_url": f"/api/artifacts/{artifact_id}/previews/{artifact.active_layout_version_id}.pdf",
    }


@app.get("/api/artifacts/{artifact_id}/metadata")
def artifact_metadata(artifact_id: str) -> dict[str, Any]:
    artifact = _artifact(artifact_id)
    return {
        "artifact_id": artifact_id,
        "workflow_lane": artifact.workflow_lane,
        "source_filename": artifact.source_filename,
        "target_filename": artifact.target_filename,
        "active_layout_version_id": artifact.active_layout_version_id,
        "approved_profile_version_ids": list(artifact.approved_profiles),
    }


@app.get("/api/artifacts/{artifact_id}/previews/{filename}")
def preview(artifact_id: str, filename: str) -> Response:
    artifact = _artifact(artifact_id)
    version = filename.removesuffix(".pdf")
    if version != "original" and version not in artifact.layout_versions:
        raise HTTPException(404, "Preview not found.")
    return Response(_pdf_bytes(version), media_type="application/pdf")


@app.get("/api/artifacts/{artifact_id}/surface.html", response_class=HTMLResponse)
def surface(artifact_id: str) -> str:
    artifact = _artifact(artifact_id)
    return _surface(artifact, artifact.active_layout_version_id)


@app.get("/api/artifacts/{artifact_id}/generated.pdf")
def download_pdf(artifact_id: str) -> Response:
    artifact = _artifact(artifact_id)
    return Response(
        _pdf_bytes(artifact.active_layout_version_id),
        media_type="application/pdf",
        headers={"Content-Disposition": 'attachment; filename="mock-candidate.pdf"'},
    )


@app.get("/api/artifacts/{artifact_id}/generated.docx")
def download_docx(artifact_id: str) -> Response:
    _artifact(artifact_id)
    return Response(
        _docx_bytes(),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": 'attachment; filename="mock-candidate.docx"'},
    )
